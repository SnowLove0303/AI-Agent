from __future__ import annotations
import argparse
from pathlib import Path
from docx import Document
from tds_common import FEATURE_HANGING, FEATURE_NUMBER_START, FEATURE_TEXT_START, ROOT, SECTION_HEADINGS, dump, feature_spacing_signature, norm, numbering_shape, sha256
ROWS_CN=['乳液外观','环氧当量（EEW）','固含量','PH值（25℃）','粘度（25℃）']
ROWS_EN=['Emulsion Appearance','Epoxy Equivalent Weight','Solids Content','pH Value (25°C)','Viscosity (25°C)']
def content_after_heading(doc, heading):
    index=next(i for i,p in enumerate(doc.paragraphs) if norm(p.text)==norm(heading))
    return next(i for i in range(index+1,len(doc.paragraphs)) if doc.paragraphs[i].text.strip())

def feature_layout(doc, language):
    heading=SECTION_HEADINGS['product.features'][language]
    start=next(i for i,p in enumerate(doc.paragraphs) if norm(p.text)==norm(heading))
    next_headings={norm(SECTION_HEADINGS[field][language]) for field in ('product.application','product.storage')}
    end=next((i for i in range(start+1,len(doc.paragraphs)) if norm(doc.paragraphs[i].text) in next_headings),len(doc.paragraphs))
    indices=[i for i in range(start+1,end) if doc.paragraphs[i].text.strip()][:2]
    if len(indices)!=2: raise ValueError(f'{doc.core_properties.title or "template"}: expected two feature paragraphs')
    separator=next((i for i in range(indices[-1]+1,end) if not doc.paragraphs[i].text.strip()),None)
    if separator is None: raise ValueError(f'{doc.core_properties.title or "template"}: feature separator paragraph missing')
    paragraphs=[doc.paragraphs[i] for i in indices]
    numbers=[numbering_shape(p) for p in paragraphs]
    if any(n is None or n.get('ilvl')!='0' for n in numbers) or len({n.get('numId') for n in numbers})!=1:
        raise ValueError(f'{doc.core_properties.title or "template"}: feature numbering must be enabled and consistent')
    if feature_spacing_signature(paragraphs[0]) != feature_spacing_signature(paragraphs[1]):
        raise ValueError(f'{doc.core_properties.title or "template"}: feature paragraph spacing is not equal')
    return indices,separator
def registry_for(path):
    d=Document(str(path)); en='_EN_' in path.name; rows=[]; labels=ROWS_EN if en else ROWS_CN
    section_names={'description':'【Characterization】' if en else '【产品描述】','supply':'【Supply Form】' if en else '【供应形式】','application':'【Application】' if en else '【应用】','storage':'【Storage】' if en else '【储存】'}
    title=next(i for i,p in enumerate(d.paragraphs[6:],6) if p.text.strip() and not p.text.strip().startswith('【'))
    description=content_after_heading(d,section_names['description']); supply=content_after_heading(d,section_names['supply']); application=content_after_heading(d,section_names['application']); storage=content_after_heading(d,section_names['storage'])
    feature_indices,separator_index=feature_layout(d,'en-US' if en else 'zh-CN')
    slots=[{'field_id':'product.title','kind':'paragraph','locator':{'paragraph_index':title}},{'field_id':'product.description','kind':'paragraph','locator':{'paragraph_index':description}},{'field_id':'product.supply_form','kind':'paragraph','locator':{'paragraph_index':supply}},{'field_id':'product.features','kind':'paragraph_list','locator':{'paragraph_indices':feature_indices}},{'field_id':'product.application','kind':'paragraph','locator':{'paragraph_index':application}},{'field_id':'product.storage','kind':'paragraph','locator':{'paragraph_index':storage}}]
    for idx,label in enumerate(labels,1):
        if norm(d.tables[0].rows[idx].cells[0].text)!=norm(label): raise ValueError(f'{path.name}: locked row {idx} mismatch')
        rows.append({'field_id':'performance.'+['appearance','eew','solid_content','ph_25c','viscosity_25c'][idx-1],'kind':'performance_row','label':label,'locator':{'table_index':0,'row_index':idx}})
    typography={'section_heading':{'first_before_twips':100,'before_twips':160,'after_twips':80},'feature_items':{'before_twips':0,'after_twips':60,'line_twips':280,'line_rule':'auto'},'body_paragraphs':{'before_twips':0,'after_twips':60,'line_twips':300,'line_rule':'auto'}}
    if not en:
        typography['product_title']={'alignment':'center','left_indent_twips':0}
    budget={'allowed':en,'threshold_total_items':12 if en else 16,'levels':[{'name':'level-1','max_total_items':16,'heading_before_twips':120,'heading_after_twips':60,'item_after_twips':30,'line_twips':280,'line_rule':'auto'},{'name':'level-2','max_total_items':None,'heading_before_twips':80,'heading_after_twips':40,'item_after_twips':20,'line_twips':260,'line_rule':'exact'}]}
    return {'template':str(path.relative_to(ROOT)),'template_sha256':sha256(path),'language':'en-US' if en else 'zh-CN','company':'guocai' if '国彩' in path.name else 'guanzhi','slots':slots+rows,'performance_table':{'header_row_index':0,'data_start_row_index':1,'source_led':True,'max_source_rows':100},'performance_extension':{'allowed':True,'row_template_index':5,'insert_after_row_index':5,'max_rows':100,'label_cell_allowed_for_new_rows':True},'feature_extension':{'allowed':True,'paragraph_template_index':feature_indices[-1],'separator_paragraph_index':separator_index,'insert_after_paragraph_index':feature_indices[-1],'max_items':100},'feature_format_contract':{'paragraph_indices':feature_indices,'numbering':'enabled','number_start_twips':FEATURE_NUMBER_START,'text_start_twips':FEATURE_TEXT_START,'hanging_twips':FEATURE_HANGING,'equal_paragraph_spacing':True,'equal_character_spacing':True,'empty_numbered_paragraphs':False,'items_contiguous':True},'tail_blank_trim':{'allowed':True},'body_blank_trim':{'allowed':True,'max_body_blank_paragraphs':0,'separator_paragraphs':0,'compact_line_twips':120},'inter_section_spacing':{'separator_paragraphs':0,'heading_before_twips':160,'content_after_twips':60,'tolerance_twips':0},'english_vertical_budget':budget,'standard_typography':typography}
def main():
    ap=argparse.ArgumentParser(); ap.add_argument('--output',type=Path,default=ROOT/'mapping'/'template_field_registry.json'); args=ap.parse_args(); variants={p.stem:registry_for(p) for p in sorted((ROOT/'templates'/'active').glob('*.docx'))}; dump(args.output,{'schema_version':'1.0.0','variants':variants}); print(f'registry={args.output} variants={len(variants)}')
if __name__=='__main__': main()
