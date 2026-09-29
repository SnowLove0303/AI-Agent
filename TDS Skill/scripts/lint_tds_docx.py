"""Preflight linter for one generated TDS DOCX.

The linter checks inherited layout and company assets; it never repairs a
document. A failed check must stop the DOCX-to-PDF stage.
"""
from __future__ import annotations

import argparse
from pathlib import Path

from docx import Document

from tds_common import (
    OUTPUT_HEADINGS,
    SECTION_HEADINGS,
    body_heading_texts,
    hidden_field_ids,
    load,
    paragraph_shape,
    source_output_fidelity_errors,
    template_registry_errors,
    variant_asset_errors,
)


BODY_FIELDS = ("product.description", "product.supply_form", "product.application", "product.storage")


def _section_paragraphs(doc, heading: str | set[str], headings: set[str]) -> list:
    accepted = {heading} if isinstance(heading, str) else heading
    start = next((i for i, p in enumerate(doc.paragraphs) if p.text.strip() in accepted), None)
    if start is None:
        return []
    end = next((i for i in range(start + 1, len(doc.paragraphs)) if doc.paragraphs[i].text.strip() in headings), len(doc.paragraphs))
    return [p for p in doc.paragraphs[start + 1:end] if p.text.strip()]


def _line_break_errors(doc: Document, variant_id: str) -> list[str]:
    for paragraph in doc.paragraphs:
        if "\n" in paragraph.text or "\r" in paragraph.text:
            return [f"intra_paragraph_line_break:{variant_id}"]
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    if "\n" in paragraph.text or "\r" in paragraph.text:
                        return [f"intra_cell_line_break:{variant_id}"]
    return []


def lint_variant(base_doc: Document, output_doc: Document, variant_id: str, variant: dict, mapping: dict) -> list[str]:
    errors = _line_break_errors(output_doc, variant_id)
    headings = body_heading_texts(variant) | {"【性能指标】", "Technical Data"}
    hidden = set(hidden_field_ids(mapping))
    slots = {item.get("field_id"): item for item in variant.get("slots", [])}
    for field_id in BODY_FIELDS:
        if field_id in hidden:
            continue
        slot = slots.get(field_id) or {}
        index = (slot.get("locator") or {}).get("paragraph_index")
        if index is None or index >= len(base_doc.paragraphs):
            errors.append(f"body_anchor_missing:{variant_id}:{field_id}")
            continue
        lang = variant.get("language", "zh-CN")
        heading = {SECTION_HEADINGS[field_id][lang], OUTPUT_HEADINGS[field_id][lang]}
        actual = _section_paragraphs(output_doc, heading, headings)
        if not actual:
            errors.append(f"body_content_missing:{variant_id}:{field_id}")
            continue
        expected = paragraph_shape(base_doc.paragraphs[index])
        for offset, paragraph in enumerate(actual):
            if paragraph_shape(paragraph) != expected:
                errors.append(f"body_layout:{variant_id}:{field_id}:{offset}")

    feature = slots.get("product.features") or {}
    feature_indices = (feature.get("locator") or {}).get("paragraph_indices", [])
    if feature_indices and "product.features" not in hidden:
        heading = SECTION_HEADINGS["product.features"][variant.get("language", "zh-CN")]
        actual = _section_paragraphs(output_doc, heading, headings)
        anchor_index = variant.get("feature_extension", {}).get("paragraph_template_index", feature_indices[-1])
        if actual and anchor_index < len(base_doc.paragraphs):
            expected = paragraph_shape(base_doc.paragraphs[anchor_index])
            for offset, paragraph in enumerate(actual):
                if paragraph_shape(paragraph) != expected:
                    errors.append(f"feature_layout:{variant_id}:{offset}")

    if variant.get("language") == "zh-CN":
        errors.extend(f"source_output_fidelity:{variant_id}:{item}" for item in source_output_fidelity_errors(output_doc, mapping, {"variants": {variant_id: variant}}, variant_id))
    return errors


def lint_docx(docx_path: Path, template_path: Path, variant_id: str, variant: dict, mapping: dict) -> list[str]:
    if not docx_path.is_file():
        return [f"missing_docx:{docx_path.name}"]
    if not template_path.is_file():
        return [f"missing_template:{variant_id}"]
    return lint_variant(Document(str(template_path)), Document(str(docx_path)), variant_id, variant, mapping)


def main() -> int:
    parser = argparse.ArgumentParser(description="Lint one generated TDS DOCX")
    parser.add_argument("--docx", type=Path, required=True)
    parser.add_argument("--template", type=Path, required=True)
    parser.add_argument("--registry", type=Path, required=True)
    parser.add_argument("--mapping", type=Path, required=True)
    parser.add_argument("--variant-id", required=True)
    args = parser.parse_args()
    registry = load(args.registry)
    mapping = load(args.mapping)
    errors = template_registry_errors(registry)
    variant = registry.get("variants", {}).get(args.variant_id)
    if not variant:
        errors.append(f"unknown_variant:{args.variant_id}")
    else:
        errors.extend(lint_docx(args.docx, args.template, args.variant_id, variant, mapping))
        errors.extend(variant_asset_errors(args.docx, variant, args.variant_id))
    if errors:
        for error in sorted(set(errors)):
            print(f"FAIL {error}")
        return 1
    print(f"PASS {args.variant_id} {args.docx}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
