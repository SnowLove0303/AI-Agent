# -*- coding: utf-8 -*-
"""
Generate TDS Model Database (TDS型号数据库.xlsx)

对 'F:\\APP Location\\Guanzhi Tong\\旧版\\03_数据库\\批量覆写\\TDS MSDS' 目录下的所有 TDS 文件进行全量统计，
严格参照 'MSDS 型号数据库.xlsx' 的结构、排版、样式规范与多维统计口径，
自动输出标准企业级数据库文件：'F:\\APP Location\\Guanzhi Tong\\旧版\\03_数据库\\批量覆写\\TDS型号数据库.xlsx'。
"""

import os
import re
import sys
from typing import Dict, List, Set, Tuple, Any
import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter

ROOT_DIR = r"F:\APP Location\Guanzhi Tong\旧版\03_数据库\批量覆写\TDS MSDS"
MSDS_EXCEL = r"F:\APP Location\Guanzhi Tong\旧版\03_数据库\批量覆写\MSDS 型号数据库.xlsx"
OUTPUT_EXCEL = r"F:\APP Location\Guanzhi Tong\旧版\03_数据库\批量覆写\TDS型号数据库.xlsx"


def load_msds_canonical_models(msds_path: str) -> Dict[str, str]:
    """读取 MSDS 型号数据库中已登记的标准型号名，用于归一化对齐。"""
    canonical_map = {}
    if not os.path.exists(msds_path):
        print(f"[WARN] MSDS reference excel not found at: {msds_path}")
        return canonical_map

    wb = openpyxl.load_workbook(msds_path, data_only=True)
    if "MSDS型号清单" in wb.sheetnames:
        ws = wb["MSDS型号清单"]
        for r in range(2, ws.max_row + 1):
            val = ws.cell(r, 2).value
            if val:
                model_str = str(val).strip()
                # 键：去除空格、破折号、下划线并转大写
                norm_key = re.sub(r"[\s\-_]", "", model_str).upper()
                canonical_map[norm_key] = model_str
    return canonical_map


def parse_model_name(filename: str, canonical_map: Dict[str, str]) -> str:
    """从 TDS 文件名中稳健抽取产品型号，并与 MSDS 标准型号库对齐。"""
    fname, _ = os.path.splitext(filename)
    fname = fname.strip()

    # 1. 冠志TDS 前缀的特殊命名（如 冠志TDS B760 CN.pdf、冠志TDS FB-805 CN.pdf）
    if fname.startswith("冠志TDS "):
        rest = fname[len("冠志TDS "):].strip()
        m = re.match(r"^([A-Za-z0-9\+\-]+)\s*(?:CN|EN)?", rest)
        raw_m = m.group(1) if m else rest
    elif fname.startswith("冠志TDS"):
        rest = fname[len("冠志TDS"):].strip()
        m = re.match(r"^([A-Za-z0-9\+\-]+)\s*(?:CN|EN)?", rest)
        raw_m = m.group(1) if m else rest
    elif "PA-3615TDS" in fname:
        raw_m = "PA-3615"
    else:
        # 2. 标准模式：<型号> <...tds...>
        parts = re.split(r"[\s_\-]tds[\s_\-\(\（]", fname, flags=re.IGNORECASE)
        if len(parts) > 1:
            raw_m = parts[0].strip(" _-")
        else:
            parts2 = re.split(r"tds", fname, flags=re.IGNORECASE)
            raw_m = parts2[0].strip(" _-")

    # 3. 查表归一化（对齐 MSDS 中的大小写与空格规范，如 OS-9547 PGDA80 -> OS-9547PGDA80）
    norm_key = re.sub(r"[\s\-_]", "", raw_m).upper()
    if norm_key in canonical_map:
        return canonical_map[norm_key]
    return raw_m


def parse_file_attributes(filename: str, rel_path: str, ext: str) -> Dict[str, Any]:
    """抽取单个文件的语言、公司、版本格式、一级分类等维度特征。"""
    fname, _ = os.path.splitext(filename)
    fname_upper = fname.upper()

    # 格式：PDF 或 WORD
    fmt = "PDF" if ext.lower() == ".pdf" else "WORD"

    # 分类与版本文件夹
    # rel_path 结构形如: TDS MSDS\产品 TDS MSDS -- <PDF/WORD>版本\<Category Folder>\<subfolders...>
    parts = rel_path.split(os.sep)
    ver_folder = parts[1] if len(parts) > 1 else ""
    cat_folder = parts[2] if len(parts) > 2 else "其他"

    # 语言判断
    is_en = False
    is_cn = False
    if (
        re.search(r"(?:[\s_\-\(]|^)EN(?:[\s_\-\.\)]|$)", fname_upper)
        or "TDS_EN" in fname_upper
        or "TDS-EN" in fname_upper
        or "TDS EN" in fname_upper
        or "TDS+EN" in fname_upper
        or "GUANZHI" in fname_upper
        or "GUOCAI" in fname_upper
        or "英文" in rel_path
    ):
        is_en = True

    if (
        re.search(r"(?:[\s_\-\(]|^)CN(?:[\s_\-\.\)]|$)", fname_upper)
        or "TDS_CN" in fname_upper
        or "TDS-CN" in fname_upper
        or "TDS CN" in fname_upper
        or "冠志" in fname
        or "国彩" in fname
        or "中文" in rel_path
    ):
        is_cn = True

    # 若未明确匹配，根据是否存在中文字符降级判断
    if not is_cn and not is_en:
        if any("\u4e00" <= ch <= "\u9fff" for ch in fname):
            is_cn = True
        else:
            is_en = True

    # 归属公司判断
    is_gz = False
    is_gc = False
    if "冠志" in fname or "GUANZHI" in fname_upper or fname.startswith("冠志TDS"):
        is_gz = True
    if "国彩" in fname or "GUOCAI" in fname_upper or re.search(r"guo\s*cai", fname, re.IGNORECASE):
        is_gc = True

    return {
        "format": fmt,
        "category": cat_folder,
        "is_cn": is_cn,
        "is_en": is_en,
        "is_guanzhi": is_gz,
        "is_guocai": is_gc,
    }


def pick_representative_file(files: List[Dict[str, Any]]) -> str:
    """
    优选该型号的代表性示例文件。
    优先级规则：
    1. 中文 Word 格式（.docx / .doc 含 CN 或 冠志/国彩）
    2. 任意 Word 格式（.docx / .doc）
    3. 中文 PDF 格式
    4. 任意其它格式
    """
    def score(item: Dict[str, Any]) -> int:
        f = item["filename"]
        ext = item["ext"].lower()
        is_word = ext in [".docx", ".doc"]
        is_cn = item["attrs"]["is_cn"]
        s = 0
        if is_word and is_cn:
            s += 100
        elif is_word:
            s += 80
        elif is_cn:
            s += 50
        else:
            s += 10
        if ext == ".docx":
            s += 5
        return s

    sorted_files = sorted(files, key=score, reverse=True)
    return sorted_files[0]["filename"]


def get_major_category(model: str) -> str:
    """提取型号的大类前缀（如 PU, PA, OS, BL 等；无前缀归为 '其他'）。"""
    if "-" in model:
        return model.split("-")[0].strip()
    m = re.match(r"^([A-Za-z]+)", model)
    if m:
        return m.group(1).upper()
    return "其他"


def generate_database():
    print("=" * 60)
    print("开始统计与生成 TDS 型号数据库...")
    print(f"源扫描目录: {ROOT_DIR}")
    print(f"参考模板库: {MSDS_EXCEL}")
    print(f"目标输出库: {OUTPUT_EXCEL}")
    print("=" * 60)

    canonical_map = load_msds_canonical_models(MSDS_EXCEL)
    print(f"已加载 MSDS 规范型号词条: {len(canonical_map)} 项")

    # 1. 遍历扫描所有文件
    valid_tds_records = []
    skipped_temp_files = 0

    for root, dirs, files in os.walk(ROOT_DIR):
        for f in files:
            # 严格过滤 Word 临时锁文件（~$ 开头）
            if f.startswith("~"):
                skipped_temp_files += 1
                continue
            ext = os.path.splitext(f)[1].lower()
            if ext in [".doc", ".docx", ".pdf"] and "tds" in f.lower():
                full_path = os.path.join(root, f)
                rel_path = os.path.relpath(full_path, ROOT_DIR)
                attrs = parse_file_attributes(f, rel_path, ext)
                model = parse_model_name(f, canonical_map)
                valid_tds_records.append({
                    "filename": f,
                    "rel_path": rel_path,
                    "full_path": full_path,
                    "ext": ext,
                    "model": model,
                    "attrs": attrs,
                })

    print(f"扫描完成: 过滤临时锁文件 {skipped_temp_files} 个，获取有效 TDS 文件 {len(valid_tds_records)} 个")

    # 2. 按型号聚合
    model_groups: Dict[str, List[Dict[str, Any]]] = {}
    for rec in valid_tds_records:
        model_groups.setdefault(rec["model"], []).append(rec)

    distinct_models = sorted(model_groups.keys(), key=lambda x: x.upper())
    total_models = len(distinct_models)
    print(f"成功归纳提取独立型号: {total_models} 个")

    # 3. 创建 Excel 工作簿与样式定义
    wb = openpyxl.Workbook()
    # 移除默认的空 sheet
    default_sheet = wb.active
    wb.remove(default_sheet)

    font_header = Font(name="宋体", size=11, bold=True)
    font_body = Font(name="宋体", size=11, bold=False)
    align_center = Alignment(horizontal="center", vertical="center")
    align_left = Alignment(horizontal="left", vertical="center")

    thin_border_side = Side(style="thin", color="D9D9D9")
    border_cell = Border(
        left=thin_border_side,
        right=thin_border_side,
        top=thin_border_side,
        bottom=thin_border_side,
    )
    fill_header = PatternFill(start_color="F2F2F2", end_color="F2F2F2", fill_type="solid")

    # =========================================================================
    # Sheet 1: TDS型号清单
    # =========================================================================
    ws1 = wb.create_sheet(title="TDS型号清单")
    headers1 = [
        "序号", "型号", "TDS文件数", "中文文件", "英文文件",
        "冠志", "国彩", "PDF", "WORD", "分类", "示例文件", "八格式覆写(V1.3.15)"
    ]
    ws1.append(headers1)

    for idx, model in enumerate(distinct_models, start=1):
        items = model_groups[model]
        total_files = len(items)
        cn_count = sum(1 for it in items if it["attrs"]["is_cn"])
        en_count = sum(1 for it in items if it["attrs"]["is_en"])
        gz_count = sum(1 for it in items if it["attrs"]["is_guanzhi"])
        gc_count = sum(1 for it in items if it["attrs"]["is_guocai"])
        pdf_count = sum(1 for it in items if it["attrs"]["format"] == "PDF")
        word_count = sum(1 for it in items if it["attrs"]["format"] == "WORD")

        # 取出现频次最高的分类目录作为主分类
        cats = [it["attrs"]["category"] for it in items if it["attrs"]["category"]]
        primary_cat = max(set(cats), key=cats.count) if cats else "其他"

        sample_file = pick_representative_file(items)
        rewrite_status = None  # 状态预留位

        ws1.append([
            idx, model, total_files, cn_count, en_count,
            gz_count, gc_count, pdf_count, word_count,
            primary_cat, sample_file, rewrite_status
        ])

    # 样式与列宽设置 - Sheet 1
    col_widths1 = {
        "A": 8.0, "B": 18.0, "C": 12.0, "D": 10.0, "E": 10.0,
        "F": 8.0, "G": 8.0, "H": 8.0, "I": 8.0, "J": 35.0,
        "K": 45.0, "L": 20.0
    }
    for col_letter, w in col_widths1.items():
        ws1.column_dimensions[col_letter].width = w

    for r_idx, row in enumerate(ws1.iter_rows(), start=1):
        ws1.row_dimensions[r_idx].height = 20.0
        for c_idx, cell in enumerate(row, start=1):
            cell.font = font_header if r_idx == 1 else font_body
            cell.border = border_cell
            if r_idx == 1:
                cell.fill = fill_header
                cell.alignment = align_center
            else:
                # B列型号、J列分类、K列示例文件左对齐，其余列居中
                if c_idx in [2, 10, 11]:
                    cell.alignment = align_left
                else:
                    cell.alignment = align_center

    # =========================================================================
    # Sheet 2: 大类统计
    # =========================================================================
    ws2 = wb.create_sheet(title="大类统计")
    headers2 = ["序号", "大类", "型号数", "占比", "型号列表"]
    ws2.append(headers2)

    cat_map: Dict[str, List[str]] = {}
    for m in distinct_models:
        major = get_major_category(m)
        cat_map.setdefault(major, []).append(m)

    # 按型号数降序排序
    sorted_cats = sorted(cat_map.items(), key=lambda x: len(x[1]), reverse=True)

    for idx, (cat_name, m_list) in enumerate(sorted_cats, start=1):
        m_list_sorted = sorted(m_list, key=lambda x: x.upper())
        m_count = len(m_list_sorted)
        ratio_str = f"{m_count / total_models * 100:.1f}%"
        model_list_str = "、".join(m_list_sorted)
        ws2.append([idx, cat_name, m_count, ratio_str, model_list_str])

    col_widths2 = {"A": 8.0, "B": 10.0, "C": 10.0, "D": 10.0, "E": 150.0}
    for col_letter, w in col_widths2.items():
        ws2.column_dimensions[col_letter].width = w

    for r_idx, row in enumerate(ws2.iter_rows(), start=1):
        ws2.row_dimensions[r_idx].height = 20.0
        for c_idx, cell in enumerate(row, start=1):
            cell.font = font_header if r_idx == 1 else font_body
            cell.border = border_cell
            if r_idx == 1:
                cell.fill = fill_header
                cell.alignment = align_center
            else:
                if c_idx == 5:
                    cell.alignment = align_left
                else:
                    cell.alignment = align_center

    # =========================================================================
    # Sheet 3: 文件夹统计
    # =========================================================================
    ws3 = wb.create_sheet(title="文件夹统计")
    headers3 = ["序号", "版本", "文件夹", "型号数", "型号列表"]
    ws3.append(headers3)

    # 聚合 (version, category) -> Set of models
    folder_agg: Dict[Tuple[str, str], Set[str]] = {}
    for rec in valid_tds_records:
        rel_p = rec["rel_path"]
        parts = rel_p.split(os.sep)
        ver = "PDF" if "PDF版本" in parts[1] else "WORD"
        cat = parts[2] if len(parts) > 2 else "其他"
        folder_agg.setdefault((ver, cat), set()).add(rec["model"])

    # 排序：PDF 优先，WORD 随后；内部按文件夹自然排序
    sorted_folders = sorted(
        folder_agg.items(),
        key=lambda x: (0 if x[0][0] == "PDF" else 1, x[0][1])
    )

    for idx, ((ver, cat), m_set) in enumerate(sorted_folders, start=1):
        m_list_sorted = sorted(m_set, key=lambda x: x.upper())
        m_count = len(m_list_sorted)
        model_list_str = "、".join(m_list_sorted)
        ws3.append([idx, ver, cat, m_count, model_list_str])

    col_widths3 = {"A": 8.0, "B": 10.0, "C": 44.0, "D": 10.0, "E": 150.0}
    for col_letter, w in col_widths3.items():
        ws3.column_dimensions[col_letter].width = w

    for r_idx, row in enumerate(ws3.iter_rows(), start=1):
        ws3.row_dimensions[r_idx].height = 20.0
        for c_idx, cell in enumerate(row, start=1):
            cell.font = font_header if r_idx == 1 else font_body
            cell.border = border_cell
            if r_idx == 1:
                cell.fill = fill_header
                cell.alignment = align_center
            else:
                if c_idx in [3, 5]:
                    cell.alignment = align_left
                else:
                    cell.alignment = align_center

    # 保存工作簿
    os.makedirs(os.path.dirname(OUTPUT_EXCEL), exist_ok=True)
    wb.save(OUTPUT_EXCEL)
    print(f"\n[SUCCESS] 成功保存 TDS型号数据库至: {OUTPUT_EXCEL}")

    # =========================================================================
    # 自检与闭环验证
    # =========================================================================
    print("\n--- 执行数据完整性校验 ---")
    wb_verify = openpyxl.load_workbook(OUTPUT_EXCEL, data_only=True)
    ws1_v = wb_verify["TDS型号清单"]
    ws2_v = wb_verify["大类统计"]
    ws3_v = wb_verify["文件夹统计"]

    total_files_sum = sum(ws1_v.cell(r, 3).value for r in range(2, ws1_v.max_row + 1))
    total_pdf_sum = sum(ws1_v.cell(r, 8).value for r in range(2, ws1_v.max_row + 1))
    total_word_sum = sum(ws1_v.cell(r, 9).value for r in range(2, ws1_v.max_row + 1))
    total_cat_models = sum(ws2_v.cell(r, 3).value for r in range(2, ws2_v.max_row + 1))

    print(f"清单总行数 (含表头): {ws1_v.max_row} 行 (有效型号数: {ws1_v.max_row - 1})")
    print(f"清单文件数求和: {total_files_sum} 份 (PDF: {total_pdf_sum}, WORD: {total_word_sum})")
    print(f"大类统计型号总数求和: {total_cat_models}")
    print(f"文件夹统计分组数: {ws3_v.max_row - 1} 组")

    assert total_files_sum == len(valid_tds_records), "文件求和与扫描总数不一致！"
    assert total_files_sum == (total_pdf_sum + total_word_sum), "PDF + WORD 与文件总数不闭环！"
    assert total_cat_models == total_models, "大类统计型号总和与清单型号数不一致！"
    print(">>> 自动化断言检验全部 100% 通过！")


if __name__ == "__main__":
    generate_database()
