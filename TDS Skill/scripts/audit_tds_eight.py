from __future__ import annotations
import argparse, hashlib, json, re, subprocess, zipfile
from pathlib import Path
from docx import Document
from tds_common import OUTPUT_HEADINGS, PERFORMANCE_TOPOLOGIES, ROOT, SECTION_HEADINGS, apply_vertical_budget_snapshot, body_blank_count, dump, hidden_block_indices, hidden_field_ids, inter_section_gap_errors, load, package_inventory, performance_topology, sha256, doc_snapshot, source_fidelity_errors, source_output_fidelity_errors, trim_body_blank_snapshot, typography_contract_errors

def shape(s):
    s=json.loads(json.dumps(s));
    for p in s['paragraphs']:
        p['text']=''
        p['shape']['runs']=[run for run in p['shape']['runs'] if any(run.get(key) is not None for key in ('bold','italic','underline','font','size_pt','color','character_spacing','kerning','position'))]
    for t in s['tables']:
        for row in t['rows']:
            for c in row['cells']:
                c['text']=''
                for p in c.get('shape',{}).get('paragraphs',[]):
                    p['text']=''
                    p['runs']=[run for run in p['runs'] if any(run.get(key) is not None for key in ('bold','italic','underline','font','size_pt','color','character_spacing','kerning','position'))]
    for sec in s['sections']:
        for p in sec['header']+sec['footer']: p['text']=''
    return s
def semantic_fields(mapping):
    model=mapping.get('normalized_model',{})
    normalized=model.get('fields',{})
    return normalized or mapping.get('mapped_fields',{})
def semantic_rows(mapping):
    rows=mapping.get('normalized_model',{}).get('performance_rows')
    return rows if rows is not None else mapping.get('performance_rows')
def semantic_values(item,lang):
    return (item.get('normalized_values',{}) if 'normalized_values' in item else item.get('values',{})).get(lang)
def semantic_labels(item,lang):
    return (item.get('normalized_label_values',{}) if 'normalized_label_values' in item else item.get('label_values',{})).get(lang)
def semantic_units(item,lang):
    return (item.get('normalized_unit_values',{}) if 'normalized_unit_values' in item else item.get('unit_values',{})).get(lang)
def semantic_methods(item,lang):
    return (item.get('normalized_test_method_values',{}) if 'normalized_test_method_values' in item else item.get('test_method_values',{})).get(lang)
def feature_contract_ok(base_doc, output_doc, variant, mapping=None):
    """Use the same template-clone shape gate as the full geometry audit."""
    return audit_shape(base_doc, output_doc, variant, mapping)
def hidden_paragraph_indices(base_doc, variant, mapping):
    """Pristine-template paragraph indices removed by whole-section hiding. None when the heading cannot be located (fail closed). Uses the single template-authority rule shared with overwrite."""
    hidden=hidden_field_ids(mapping); lang=variant['language']
    slots={x['field_id']:x for x in variant['slots']}
    return hidden_block_indices(list(base_doc.paragraphs),lang,slots,hidden)

def topology_template(table, variant, topology):
    spec=variant.get('performance_table',{}); widths=spec.get('topology_widths',{}).get(topology)
    if not widths or len(widths) != PERFORMANCE_TOPOLOGIES[topology]: return None
    target=len(widths); table=json.loads(json.dumps(table)); table['column_count']=target; table['grid_widths']=[str(width) for width in widths]
    for row in table['rows']:
        row['cells']=row['cells'][:target]
        for index,cell in enumerate(row['cells']): cell['shape']['props']['width']=str(widths[index])
    return table

def audit_shape(base_doc, output_doc, variant, mapping):
    base_snapshot=doc_snapshot(base_doc); output_snapshot=doc_snapshot(output_doc)
    left, right = shape(base_snapshot), shape(output_snapshot)
    left_table, right_table = left['tables'][0], right['tables'][0]
    topology=performance_topology(mapping,variant); spec=variant.get('performance_table',{}); widths=spec.get('topology_widths',{}).get(topology)
    if topology not in PERFORMANCE_TOPOLOGIES or topology not in spec.get('allowed_topologies',[]) or not widths: return False
    if left['sections'] != right['sections'] or right_table['column_count'] != PERFORMANCE_TOPOLOGIES[topology] or right_table['grid_widths'] != [str(width) for width in widths]: return False
    template_table=topology_template(left_table,variant,topology)
    if template_table is None: return False
    hide=hidden_paragraph_indices(base_doc, variant, mapping)
    if hide is None: return False
    slots={x['field_id']:x for x in variant['slots']}
    fields=semantic_fields(mapping); lang=variant['language']; inserts={}
    tail_trim=set()
    if variant.get('tail_blank_trim',{}).get('allowed',False):
        for i in range(len(base_doc.paragraphs)-1,-1,-1):
            if base_doc.paragraphs[i].text.strip(): break
            tail_trim.add(i)
    if 'product.features' not in set(hidden_field_ids(mapping)):
        feature_item=fields.get('product.features',{}) or {}
        feature_values=(feature_item.get('normalized_values',feature_item.get('values',{})) or {}).get(lang,'') or ''
        feature_count=len(variant.get('feature_format_contract',{}).get('paragraph_indices',[21,23]))
        extra=max(0,len([x for x in feature_values.splitlines() if x.strip()])-feature_count)
        anchor=variant.get('feature_extension',{}).get('paragraph_template_index',feature_count-1)
        if extra: inserts[anchor]=inserts.get(anchor,0)+extra
    for fid,slot in slots.items():
        if slot.get('kind')!='paragraph' or fid=='product.title' or fid in set(hidden_field_ids(mapping)): continue
        item=fields.get(fid,{}) or {}; values=item.get('normalized_values',item.get('values',{})) or {}
        count=max(0,len([x for x in (values.get(lang,'') or '').splitlines() if x.strip()])-1)
        if count:
            original=slot['locator']['paragraph_index']; inserts[original]=inserts.get(original,0)+count
    expected_slot_texts={}
    for fid,slot in slots.items():
        if fid in set(hidden_field_ids(mapping)): continue
        item=fields.get(fid,{}) or {}; values=item.get('normalized_values',item.get('values',{})) or {}
        if slot.get('kind')=='paragraph':
            expected_slot_texts[slot['locator']['paragraph_index']]='__value__'
        elif slot.get('kind')=='paragraph_list':
            lines=[x.strip() for x in (values.get(lang,'') or '').splitlines() if x.strip()]
            for index,original in enumerate(slot['locator']['paragraph_indices']):
                expected_slot_texts[original]=lines[index] if index<len(lines) else '__value__' if index==0 else ''
    expected_paragraphs=[]
    for original,paragraph in enumerate(base_snapshot['paragraphs']):
        if original in hide or original in tail_trim: continue
        template_paragraph=json.loads(json.dumps(paragraph))
        if original in expected_slot_texts: template_paragraph['text']=expected_slot_texts[original]
        expected_paragraphs.append(template_paragraph)
        for _ in range(inserts.get(original,0)): expected_paragraphs.append(json.loads(json.dumps(template_paragraph)))
    expected_snapshot=json.loads(json.dumps(base_snapshot)); expected_snapshot['paragraphs']=expected_paragraphs
    expected_snapshot['paragraphs']=trim_body_blank_snapshot(expected_snapshot['paragraphs'],variant)
    output_snapshot['paragraphs']=trim_body_blank_snapshot(output_snapshot['paragraphs'],variant)
    total_items=0
    fields=semantic_fields(mapping)
    feature_item=fields.get('product.features',{}) or {}
    feature_values=(feature_item.get('normalized_values',feature_item.get('values',{})) or {}).get(variant['language'],'') or ''
    total_items+=len([x for x in feature_values.splitlines() if x.strip()])
    for fid,slot in slots.items():
        if slot.get('kind')=='paragraph' and fid!='product.title' and fid not in set(hidden_field_ids(mapping)):
            values=(fields.get(fid,{}) or {}).get('normalized_values',(fields.get(fid,{}) or {}).get('values',{})) or {}
            total_items+=len([x for x in (values.get(variant['language'],'') or '').splitlines() if x.strip()])
    apply_vertical_budget_snapshot(expected_snapshot['paragraphs'],variant,total_items)
    apply_vertical_budget_snapshot(output_snapshot['paragraphs'],variant,total_items)
    expected=shape(expected_snapshot)
    right=shape(output_snapshot)
    if expected['paragraphs'] != right['paragraphs']: return False
    source_rows=semantic_rows(mapping)
    if source_rows is not None:
        start=variant.get('performance_table',{}).get('data_start_row_index',1)
        if right_table['rows'][:start] != template_table['rows'][:start]: return False
        for i,row in enumerate(right_table['rows'][start:]):
            expected=template_table['rows'][start+min(i,len(template_table['rows'])-start-1)]
            if row != expected: return False
    else:
        if template_table['rows'][:6] != right_table['rows'][:6]: return False
        extras=mapping.get('performance_extra_rows',[])
        if len(right_table['rows']) != 6+len(extras): return False
        rows=right_table['rows'][6:]
    for row in (rows if source_rows is None else []):
        if [c['shape'] for c in row['cells']] != [c['shape'] for c in template_table['rows'][5]['cells']]: return False
    return True
def audit_layout(base_doc, output_doc, variant_id, variant, mapping):
    """Reject line-break characters that would turn one template slot into an unregistered layout change."""
    errs=[]
    for p in output_doc.paragraphs:
        if '\n' in p.text or '\r' in p.text: errs.append(f'intra_paragraph_line_break:{variant_id}'); break
    if not any(e.startswith('intra_paragraph_line_break') for e in errs):
        for t in output_doc.tables:
            done=False
            for r in t.rows:
                for c in r.cells:
                    for p in c.paragraphs:
                        if '\n' in p.text or '\r' in p.text: errs.append(f'intra_paragraph_line_break:{variant_id}'); done=True; break
                    if done: break
                if done: break
            if done: break
    return errs

def pdf_page_count(path: Path) -> int:
    try:
        from pypdf import PdfReader
        return len(PdfReader(str(path)).pages)
    except Exception:
        data=path.read_bytes()
        count=len(re.findall(rb'/Type\s*/Page\b',data))
        if count: return count
    try:
        result=subprocess.run(['pdfinfo',str(path)],capture_output=True,text=True,encoding='utf-8',errors='replace',check=False)
        match=re.search(r'^Pages:\s*(\d+)$',result.stdout,re.MULTILINE)
        if result.returncode or not match: raise RuntimeError('no PDF page-count reader available')
        return int(match.group(1))
    except FileNotFoundError as exc:
        raise RuntimeError('no PDF page-count reader available: install pypdf or provide pdfinfo') from exc
def main():
    ap=argparse.ArgumentParser(); ap.add_argument('--output-dir',type=Path,required=True); ap.add_argument('--registry',type=Path,required=True); ap.add_argument('--mapping',type=Path,required=True); ap.add_argument('--model',required=True); ap.add_argument('--report',type=Path,required=True); ap.add_argument('--docx-only',action='store_true'); ap.add_argument('--conversion-evidence-dir',type=Path); args=ap.parse_args()
    docx_dir=args.output_dir/'WORD' if (args.output_dir/'WORD').is_dir() else args.output_dir
    pdf_dir=args.output_dir/'PDF' if (args.output_dir/'PDF').is_dir() else args.output_dir
    reg=load(args.registry); mapping=load(args.mapping); results=[]; errors=[]; warnings=[]
    errors.extend(f'source_fidelity_mapping:{item}' for item in source_fidelity_errors(mapping))
    fields=semantic_fields(mapping)
    model=mapping.get('normalized_model',{})
    decisions=mapping.get('decision_ledger',model.get('decision_ledger',[]))
    if mapping.get('schema_version')=='1.3.0':
        if not model: errors.append('missing_normalized_model')
        if model.get('translation',{}).get('source')!='normalized_model': errors.append('translation_not_from_normalized_model')
        if model.get('status')!='approved': errors.append(f'normalized_model_not_approved:{model.get("status","missing")}')
        decision_by_field={item.get('field_id'):item for item in decisions}
        for row in semantic_rows(mapping) or []:
            decision=decision_by_field.get(row.get('field_id'))
            if not decision: errors.append(f'missing_decision_ledger:{row.get("field_id")}'); continue
            if not decision.get('provenance'): errors.append(f'missing_provenance:{row.get("field_id")}')
        pending=sum(1 for item in decisions if item.get('needs_judgment'))
        if pending: warnings.append(f'pending_agent_judgments:{pending}')
    required_fields=['product.title','product.description','product.supply_form','product.features','product.application','product.storage']
    hidden=set(hidden_field_ids(mapping))
    for language in ('zh-CN','en-US'):
        for field in required_fields:
            if field in hidden: continue
            item=fields.get(field,{})
            values=item.get('normalized_values',item.get('values',{}))
            if language not in values: errors.append(f'missing_language_source_value:{language}:{field}')
        if semantic_rows(mapping) is not None:
            for row in semantic_rows(mapping):
                if language not in row.get('normalized_label_values',row.get('label_values',{})) or language not in row.get('normalized_values',row.get('values',{})): errors.append(f'missing_language_source_row:{language}:{row.get("field_id")}')
        else:
            for extra in mapping.get('performance_extra_rows',[]):
                if language not in extra.get('label_values',{}) or language not in extra.get('values',{}): errors.append(f'missing_language_extra_metric:{language}:{extra.get("field_id")}')
    for vid,v in reg['variants'].items():
        stem={'TDS_CN_冠志模板':'TDS_CN_冠志','TDS_CN_国彩模板':'TDS_CN_国彩','TDS_EN_冠志模板':'TDS_EN_冠志','TDS_EN_国彩模板':'TDS_EN_国彩'}[vid]
        out=docx_dir/f'{args.model}_{stem}.docx'
        if not out.is_file(): errors.append(f'missing_docx:{out.name}'); continue
        base=Document(str(ROOT/v['template'])); product=Document(str(out));
        errors.extend(f'typography_contract:{vid}:{item}' for item in typography_contract_errors(base,v))
        geometry_ok=audit_shape(base,product,v,mapping)
        if not geometry_ok: errors.append(f'geometry_changed:{vid}')
        feature_ok=feature_contract_ok(base,product,v,mapping)
        if not feature_ok: errors.append(f'feature_format_contract_changed:{vid}')
        layout_errs=audit_layout(base,product,vid,v,mapping)
        errors.extend(layout_errs)
        errors.extend(f'source_fidelity_output:{vid}:{item}' for item in source_output_fidelity_errors(product,mapping,reg,vid))
        if v.get('body_blank_trim',{}).get('allowed',False):
            blank_count=body_blank_count(product,v)
            expected_blank_count=int(v['body_blank_trim'].get('max_body_blank_paragraphs',v['body_blank_trim'].get('separator_paragraphs',0)))
            if blank_count!=expected_blank_count:
                errors.append(f'body_blank_count_mismatch:{vid}:{blank_count}!={expected_blank_count}')
        errors.extend(f'inter_section_gap:{vid}:{item}' for item in inter_section_gap_errors(product,v))
        parts0=package_inventory(ROOT/v['template']); parts1=package_inventory(out)
        for part,h in parts0.items():
            if part!='word/document.xml' and parts1.get(part)!=h: errors.append(f'package_part_changed:{vid}:{part}')
        text='\n'.join(p.text for p in product.paragraphs)+'\n'+'\n'.join(c.text for t in product.tables for r in t.rows for c in r.cells)
        leaks=[x for x in load(ROOT/'mapping'/'tds_mutation_whitelist.json')['sample_fact_tokens'] if x.lower() in text.lower() and x not in (mapping.get('allowed_source_tokens') or [])]
        if leaks: errors.append(f'sample_fact_leak:{vid}:{leaks}')
        pdf=pdf_dir/f'{out.stem}.pdf'
        page_count=None
        if not args.docx_only:
            if not pdf.is_file(): errors.append(f'missing_pdf:{pdf.name}')
            elif v.get('page_contract',{}).get('max_pages'):
                try:
                    page_count=pdf_page_count(pdf)
                    if page_count>v['page_contract']['max_pages']: errors.append(f'page_count_exceeded:{vid}:{page_count}>{v["page_contract"]["max_pages"]}')
                except Exception as exc:
                    errors.append(f'page_count_unavailable:{vid}:{type(exc).__name__}')
            if args.conversion_evidence_dir is None: errors.append('missing_conversion_evidence_dir')
            else:
                evidence_path=args.conversion_evidence_dir/f'{out.stem}.conversion.json'
                if not evidence_path.is_file(): errors.append(f'missing_conversion_evidence:{evidence_path.name}')
                else:
                    evidence=load(evidence_path)
                    try: source_matches=Path(evidence.get('source_docx','')).resolve()==out.resolve()
                    except (OSError,ValueError): source_matches=False
                    if not source_matches: errors.append(f'conversion_source_mismatch:{vid}')
                    if evidence.get('source_sha256')!=sha256(out): errors.append(f'conversion_source_hash_mismatch:{vid}')
                    if pdf.is_file() and evidence.get('output_sha256')!=sha256(pdf): errors.append(f'conversion_output_hash_mismatch:{vid}')
                    if evidence.get('source_is_final_docx') is not True: errors.append(f'conversion_not_final_docx:{vid}')
                    if evidence.get('independent_pdf_authoring') is not False: errors.append(f'independent_pdf_authoring:{vid}')
            if pdf.stem!=out.stem: errors.append(f'pdf_pair_name_mismatch:{vid}')
        results.append({'variant_id':vid,'docx':str(out),'docx_sha256':sha256(out),'pdf':None if args.docx_only else str(pdf),'pdf_sha256':None if args.docx_only or not pdf.is_file() else sha256(pdf),'page_count':page_count,'conversion_evidence':None if args.docx_only else str(args.conversion_evidence_dir/f'{out.stem}.conversion.json') if args.conversion_evidence_dir else None,'pdf_derived_name_match':None if args.docx_only else pdf.stem==out.stem,'geometry':'pass' if geometry_ok else 'fail','feature_format':'pass' if feature_ok else 'fail'})
        if semantic_rows(mapping) is not None:
            start=v.get('performance_table',{}).get('data_start_row_index',1)
            actual=[[c.text for c in row.cells] for row in product.tables[0].rows[start:]]
            expected=[]
            for source_row in semantic_rows(mapping):
                row_values=[
                    semantic_labels(source_row,v['language']) or ('无数据' if v['language']=='zh-CN' else 'No data available'),
                    semantic_values(source_row,v['language']) or ('无数据' if v['language']=='zh-CN' else 'No data available')]
                topology=performance_topology(mapping,v)
                if topology in ('3-col','4-col'): row_values.append(semantic_units(source_row,v['language']) or source_row.get('unit','') or '')
                if topology=='4-col': row_values.append(semantic_methods(source_row,v['language']) or source_row.get('test_method','') or '')
                expected.append(row_values)
            if actual!=expected: errors.append(f'performance_source_parity:{vid}')
    pdfs=[] if args.docx_only else [p for p in pdf_dir.glob('*.pdf') if '_TDS_' in p.name]
    if not args.docx_only and len(pdfs)!=4: errors.append(f'pdf_count:{len(pdfs)}')
    docxs=[p for p in docx_dir.glob('*.docx') if '_TDS_' in p.name]
    if len(docxs)!=4: errors.append(f'docx_count:{len(docxs)}')
    report={'schema_version':'1.2.0','status':'DOCX_PREFLIGHT_PASS' if args.docx_only and not errors else 'RELEASE_PASS' if not errors else 'RELEASE_FAIL','release_blocker':bool(errors),'docx_count':len(docxs),'pdf_count':len(pdfs),'docx_only':args.docx_only,'errors':sorted(set(errors)),'warnings':sorted(set(warnings)),'normalization_model_status':model.get('status','legacy-mapping'),'translation_source':model.get('translation',{}).get('source','legacy-mapping'),'variants':results,'customer_ready':False,'ready_for_user_proofreading':not errors and not warnings}
    dump(args.report,report); print(f"status={report['status']} errors={len(errors)}")
    raise SystemExit(1 if errors else 0)
if __name__=='__main__': main()
