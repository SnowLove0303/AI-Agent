from __future__ import annotations
import argparse,re
from pathlib import Path
from docx import Document
from tds_common import dump,sha256
HEADINGS={'产品描述':'product.description','Characterization':'product.description','Product Description':'product.description','供应形式':'product.supply_form','Supply Form':'product.supply_form','产品特性':'product.features','Product features':'product.features','Product Features':'product.features','应用':'product.application','Application':'product.application','储存':'product.storage','包装储存':'product.storage','包装与储存':'product.storage','Storage':'product.storage'}
SECTION_BREAKS={'性能指标','Technical Data'}
def main():
    ap=argparse.ArgumentParser(); ap.add_argument('source',type=Path); ap.add_argument('--language',choices=['zh-CN','en-US'],required=True); ap.add_argument('--output',type=Path,required=True); a=ap.parse_args(); d=Document(str(a.source)); sections={}; current=None
    for i,p in enumerate(d.paragraphs):
        h=p.text.strip().strip('【】')
        if h in SECTION_BREAKS: current=None; continue
        if h in HEADINGS: current=HEADINGS[h]; sections.setdefault(current,[]); continue
        if current and p.text.strip(): sections[current].append({'paragraph_index':i,'text':p.text})
    perf=[]
    for ti,t in enumerate(d.tables):
        for ri,row in enumerate(t.rows[1:],1):
            v=[c.text.strip() for c in row.cells]
            if v and any(v): perf.append({'table_index':ti,'row_index':ri,'source_column_count':len(v),'item':v[0],'value':v[1] if len(v)>1 else '','unit':v[2] if len(v)>2 else '','test_method':v[3] if len(v)>3 else '','source_location':f'table[{ti}].row[{ri}]'})
    title=next(({'paragraph_index':i,'text':p.text} for i,p in enumerate(d.paragraphs) if re.search(r'[A-Za-z]{1,8}[- ]?\d{3,6}[A-Za-z0-9]*',p.text) and not p.text.lstrip().startswith(('电话','Tel'))),{'paragraph_index':None,'text':''})
    dump(a.output,{'schema_version':'1.0.0','source_docx':str(a.source.resolve()),'source_docx_sha256':sha256(a.source),'language':a.language,'title':title,'sections':{k:{'text':'\n'.join(x['text'] for x in v),'locations':[x['paragraph_index'] for x in v]} for k,v in sections.items()},'performance_rows':perf}); print(f'extracted={a.output} rows={len(perf)}')
if __name__=='__main__': main()
