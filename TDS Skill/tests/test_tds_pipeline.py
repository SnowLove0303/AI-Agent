from pathlib import Path
from copy import deepcopy
import json
import subprocess
import sys
import pytest

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "scripts"))

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from overwrite_tds import write_variant
from audit_tds_eight import audit_shape, pdf_page_count
from tds_common import body_blank_count, feature_spacing_signature, inter_section_gap_errors, load, numbering_shape, package_inventory, paragraph_shape, source_fidelity_errors, source_output_fidelity_errors, text_sha256, typography_contract_errors


def test_all_four_variants_are_fresh_clones(tmp_path):
    registry = load(ROOT / "mapping" / "template_field_registry.json")
    mapping = load(ROOT / "tests" / "fixtures" / "valid_mapping.json")
    for variant_id, variant in registry["variants"].items():
        output = tmp_path / f"TDS-DEMO_{variant_id}.docx"
        write_variant(mapping, registry, variant_id, output)
        assert output.is_file()
        execution_log = output.with_name(output.name + ".overwrite.log.json")
        assert execution_log.is_file()
        log = json.loads(execution_log.read_text(encoding="utf-8"))
        assert log["status"] == "completed"
        assert [e["event"] for e in log["events"]][-1] == "leak_scan_passed"
        generation = json.loads(output.with_suffix(output.suffix + ".generation.json").read_text(encoding="utf-8"))
        assert generation["execution_log_file"] == execution_log.name
        doc = Document(str(output))
        assert len(doc.tables) == 1
        assert len(doc.tables[0].rows) == 6
        assert "EP-1704" not in "\n".join(p.text for p in doc.paragraphs)
        original = package_inventory(ROOT / variant["template"])
        generated = package_inventory(output)
        assert {k: v for k, v in generated.items() if k != "word/document.xml"} == {k: v for k, v in original.items() if k != "word/document.xml"}


def test_variant_metadata_can_be_separated_from_word_output(tmp_path):
    registry = load(ROOT / "mapping" / "template_field_registry.json")
    mapping = load(ROOT / "tests" / "fixtures" / "valid_mapping.json")
    output = tmp_path / "WORD" / "TDS-DEMO_TDS_CN_冠志.docx"
    log_dir = tmp_path / "audit" / "execution_logs"
    generation_dir = tmp_path / "audit" / "generation"
    write_variant(mapping, registry, "TDS_CN_冠志模板", output, log_dir, generation_dir, tmp_path)
    assert output.is_file()
    assert not list(output.parent.glob("*.json"))
    assert (log_dir / (output.name + ".overwrite.log.json")).is_file()
    generation = json.loads((generation_dir / (output.name + ".generation.json")).read_text(encoding="utf-8"))
    assert generation["execution_log_file"] == "audit/execution_logs/" + output.name + ".overwrite.log.json"


def test_application_multiline_clones_template_body_style(tmp_path):
    registry = load(ROOT / "mapping" / "template_field_registry.json")
    mapping = deepcopy(load(ROOT / "tests" / "fixtures" / "valid_mapping.json"))
    mapping["mapped_fields"]["product.application"]["values"]["zh-CN"] = "1.适用于高光泽涂层。\n2.与各种基材附着力优异。"
    output = tmp_path / "application.docx"
    write_variant(mapping, registry, "TDS_CN_冠志模板", output)
    doc = Document(str(output))
    heading = next(i for i, p in enumerate(doc.paragraphs) if p.text == "【应用】")
    lines = [p.text for p in doc.paragraphs[heading + 1:heading + 4] if p.text.strip()]
    assert lines[:2] == ["1.适用于高光泽涂层。", "2.与各种基材附着力优异。"]
    assert all("\n" not in p.text and "\r" not in p.text for p in doc.paragraphs)
    base = Document(str(ROOT / registry["variants"]["TDS_CN_冠志模板"]["template"]))
    from tds_common import paragraph_shape
    body = [p for p in doc.paragraphs[heading + 1:] if p.text.strip()]
    app_index = next(item["locator"]["paragraph_index"] for item in registry["variants"]["TDS_CN_冠志模板"]["slots"] if item["field_id"] == "product.application")
    assert paragraph_shape(body[0]) == paragraph_shape(base.paragraphs[app_index])
    assert paragraph_shape(body[1]) == paragraph_shape(base.paragraphs[app_index])


def test_language_specific_template_labels_and_grid_widths_are_registered():
    registry = load(ROOT / "mapping" / "template_field_registry.json")
    en_labels = [
        item["label"]
        for item in registry["variants"]["TDS_EN_冠志模板"]["slots"]
        if item["kind"] == "performance_row"
    ]
    assert en_labels == [
        "Emulsion Appearance",
        "Epoxy Equivalent Weight",
        "Solids Content",
        "pH Value (25°C)",
        "Viscosity (25°C)",
    ]
    assert registry["variants"]["TDS_EN_冠志模板"]["template_sha256"] != "4d602128bd397235e76e48d39a9a02b7862b840063734bc89b67661271d0a216"


def test_english_templates_use_clean_western_font_inheritance():
    registry = load(ROOT / "mapping" / "template_field_registry.json")
    for variant_id in ("TDS_EN_冠志模板", "TDS_EN_国彩模板"):
        doc = Document(str(ROOT / registry["variants"][variant_id]["template"]))
        slots = {item["field_id"]: item for item in registry["variants"][variant_id]["slots"]}
        body_indices = [slots[field_id]["locator"]["paragraph_index"] for field_id in ("product.description", "product.supply_form", "product.application", "product.storage")]
        paragraphs = list(doc.paragraphs[9:])
        for index, paragraph in enumerate(paragraphs, 9):
            ppr = paragraph._p.pPr
            ind = ppr.find(qn("w:ind")) if ppr is not None else None
            if index in body_indices:
                assert ind is not None
                assert ind.get(qn("w:firstLineChars")) == "200"
                assert ind.get(qn("w:firstLine")) == "480"
            if index in slots["product.features"]["locator"]["paragraph_indices"]:
                assert ind is not None
                assert ind.get(qn("w:left")) == "840"
                assert ind.get(qn("w:hanging")) == "360"
                assert ind.get(qn("w:firstLineChars")) is None
                assert ind.get(qn("w:firstLine")) is None
            for run in paragraph.runs:
                rpr = run._r.rPr
                rfonts = rpr.find(qn("w:rFonts")) if rpr is not None else None
                assert rfonts is not None
                assert rfonts.get(qn("w:ascii")) == "Times New Roman"
                assert rfonts.get(qn("w:hAnsi")) == "Times New Roman"
                assert rfonts.get(qn("w:cs")) == "Times New Roman"
                assert rfonts.get(qn("w:eastAsia")) is None
                assert rfonts.get(qn("w:hint")) is None
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    for paragraph in cell.paragraphs:
                        for run in paragraph.runs:
                            rpr = run._r.rPr
                            rfonts = rpr.find(qn("w:rFonts")) if rpr is not None else None
                            assert rfonts is not None
                            assert rfonts.get(qn("w:ascii")) == "Times New Roman"
                            assert rfonts.get(qn("w:hAnsi")) == "Times New Roman"
                            assert rfonts.get(qn("w:eastAsia")) is None
                            assert rfonts.get(qn("w:hint")) is None


def test_source_two_column_topology_prunes_unused_template_columns(tmp_path):
    registry = load(ROOT / "mapping" / "template_field_registry.json")
    mapping = deepcopy(load(ROOT / "tests" / "fixtures" / "valid_mapping.json"))
    mapping["performance_rows"] = [
        {"label_values": {"zh-CN": "外观", "en-US": "Appearance"}, "values": {"zh-CN": "乳白色液体", "en-US": "Milky liquid"}, "unit_values": {"zh-CN": "", "en-US": ""}, "test_method_values": {"zh-CN": "", "en-US": ""}},
        {"label_values": {"zh-CN": "粘度", "en-US": "Viscosity"}, "values": {"zh-CN": "<5000", "en-US": "<5000"}, "unit_values": {"zh-CN": "", "en-US": ""}, "test_method_values": {"zh-CN": "", "en-US": ""}},
    ]
    for variant_id, variant in registry["variants"].items():
        output = tmp_path / f"two-col-{variant_id}.docx"
        write_variant(mapping, registry, variant_id, output)
        doc = Document(str(output))
        assert len(doc.tables[0].columns) == 2
        assert [cell.text for cell in doc.tables[0].rows[1].cells] == ["外观" if variant["language"] == "zh-CN" else "Appearance", "乳白色液体" if variant["language"] == "zh-CN" else "Milky liquid"]
        assert audit_shape(Document(str(ROOT / variant["template"])), doc, variant, mapping)


def test_feature_slots_preserve_numbering_and_equal_spacing():
    registry = load(ROOT / "mapping" / "template_field_registry.json")
    for variant in registry["variants"].values():
        doc = Document(str(ROOT / variant["template"]))
        indices = variant["feature_format_contract"]["paragraph_indices"]
        feature_paragraphs = [doc.paragraphs[i] for i in indices]
        numbers = [numbering_shape(p) for p in feature_paragraphs]
        assert all(n is not None and n.get("ilvl") == "0" for n in numbers)
        assert len({n.get("numId") for n in numbers}) == 1
        assert feature_spacing_signature(feature_paragraphs[0]) == feature_spacing_signature(feature_paragraphs[1])
        assert variant["feature_format_contract"]["numbering"] == "enabled"
        assert variant["feature_format_contract"]["number_start_twips"] == 480
        assert variant["feature_format_contract"]["text_start_twips"] == 840
        assert variant["feature_format_contract"]["hanging_twips"] == 360


def test_active_templates_match_standard_typography_contract():
    registry = load(ROOT / "mapping" / "template_field_registry.json")
    for variant in registry["variants"].values():
        assert typography_contract_errors(Document(str(ROOT / variant["template"])), variant) == []


def test_cn_title_is_natively_centered_without_left_indent():
    registry = load(ROOT / "mapping" / "template_field_registry.json")
    for variant_id in ("TDS_CN_冠志模板", "TDS_CN_国彩模板"):
        variant = registry["variants"][variant_id]
        doc = Document(str(ROOT / variant["template"]))
        title_index = next(item["locator"]["paragraph_index"] for item in variant["slots"] if item["field_id"] == "product.title")
        title = doc.paragraphs[title_index]
        assert title.alignment == WD_ALIGN_PARAGRAPH.CENTER
        assert title.paragraph_format.left_indent in (None, 0)
        assert typography_contract_errors(doc, variant) == []


def test_inter_section_gaps_have_no_local_blank_separator(tmp_path):
    registry = load(ROOT / "mapping" / "template_field_registry.json")
    mapping = deepcopy(load(ROOT / "tests" / "fixtures" / "valid_mapping.json"))
    for variant_id, variant in registry["variants"].items():
        output = tmp_path / f"section-gaps-{variant_id}.docx"
        write_variant(mapping, registry, variant_id, output)
        doc = Document(str(output))
        assert body_blank_count(doc, variant) == 0
        assert inter_section_gap_errors(doc, variant) == []


def test_performance_extensions_clone_template_row_style(tmp_path):
    registry = load(ROOT / "mapping" / "template_field_registry.json")
    mapping = deepcopy(load(ROOT / "tests" / "fixtures" / "valid_mapping.json"))
    mapping["performance_extra_rows"] = [{
        "field_id": "performance.extra.001",
        "label_values": {"zh-CN": "粒径", "en-US": "Particle size"},
        "values": {"zh-CN": "≤ 100 nm", "en-US": "≤ 100 nm"},
        "unit": "nm",
        "test_method": "Laser diffraction"
    }]
    for variant_id in registry["variants"]:
        output = tmp_path / f"extended_{variant_id}.docx"
        write_variant(mapping, registry, variant_id, output)
        doc = Document(str(output))
        assert len(doc.tables[0].rows) == 7
        text = "\n".join(c.text for r in doc.tables[0].rows for c in r.cells)
        assert ("粒径" if registry["variants"][variant_id]["language"] == "zh-CN" else "Particle size") in text
        base = Document(str(ROOT / registry["variants"][variant_id]["template"]))
        assert audit_shape(base, doc, registry["variants"][variant_id], mapping)


def test_feature_overflow_clones_template_item_style(tmp_path):
    registry = load(ROOT / "mapping" / "template_field_registry.json")
    mapping = deepcopy(load(ROOT / "tests" / "fixtures" / "valid_mapping.json"))
    mapping["mapped_fields"]["product.features"]["values"]["zh-CN"] += "\n低气味。"
    output = tmp_path / "feature-overflow.docx"
    write_variant(mapping, registry, "TDS_CN_冠志模板", output)
    doc = Document(str(output))
    head = next(i for i, p in enumerate(doc.paragraphs) if p.text == "【产品特性】")
    items = [p for p in doc.paragraphs[head + 1:] if p.text.strip()][:3]
    assert [p.text for p in items] == ["储存稳定性良好；", "粘度稳定。", "低气味。"]
    base = Document(str(ROOT / registry["variants"]["TDS_CN_冠志模板"]["template"]))
    from tds_common import feature_layout_signature, numbering_shape
    anchor = registry["variants"]["TDS_CN_冠志模板"]["feature_extension"]["paragraph_template_index"]
    assert feature_layout_signature(items[2]) == feature_layout_signature(base.paragraphs[anchor])
    assert numbering_shape(items[2]) == numbering_shape(base.paragraphs[anchor])
    assert audit_shape(base, doc, registry["variants"]["TDS_CN_冠志模板"], mapping)


def test_text_replacement_preserves_template_run_shapes(tmp_path):
    registry = load(ROOT / "mapping" / "template_field_registry.json")
    mapping = deepcopy(load(ROOT / "tests" / "fixtures" / "valid_mapping.json"))
    mapping["mapped_fields"]["product.features"]["values"]["zh-CN"] += "\n低气味。"
    for variant_id in ("TDS_CN_冠志模板", "TDS_EN_冠志模板"):
        output = tmp_path / f"run-shape-{variant_id}.docx"
        write_variant(mapping, registry, variant_id, output)
        doc = Document(str(output))
        base = Document(str(ROOT / registry["variants"][variant_id]["template"]))
        indices = registry["variants"][variant_id]["feature_format_contract"]["paragraph_indices"]
        head = next(i for i, p in enumerate(doc.paragraphs) if p.text.strip() == ("【产品特性】" if registry["variants"][variant_id]["language"] == "zh-CN" else "Product Features"))
        next_heading = "【应用】" if registry["variants"][variant_id]["language"] == "zh-CN" else "Application"
        end = next(i for i in range(head + 1, len(doc.paragraphs)) if doc.paragraphs[i].text.strip() == next_heading)
        items = [p for p in doc.paragraphs[head + 1:end] if p.text.strip()]
        expected_indices = indices + ([registry["variants"][variant_id]["feature_extension"]["paragraph_template_index"]] if registry["variants"][variant_id]["language"] == "zh-CN" else [])
        assert [paragraph_shape(p) for p in items] == [paragraph_shape(base.paragraphs[i]) for i in expected_indices]
        assert audit_shape(base, doc, registry["variants"][variant_id], mapping)


def test_unknown_bilingual_metric_is_paired_by_source_order(tmp_path):
    sources = {
        "zh-CN": {"language": "zh-CN", "title": {"text": "TDS-DEMO", "paragraph_index": 0}, "sections": {}, "performance_rows": [{"item": "粒径", "value": "≤ 100 nm", "unit": "nm", "test_method": "方法A", "source_column_count": 4, "source_location": "table[0].row[6]"}]},
        "en-US": {"language": "en-US", "title": {"text": "TDS-DEMO", "paragraph_index": 0}, "sections": {}, "performance_rows": [{"item": "Particle size", "value": "≤ 100 nm", "unit": "nm", "test_method": "Method A", "source_column_count": 4, "source_location": "table[0].row[6]"}]}
    }
    facts = []
    for language, payload in sources.items():
        path = tmp_path / f"{language}.json"
        path.write_text(json.dumps(payload, ensure_ascii=False), encoding="utf-8")
        facts.append(path)
    output = tmp_path / "mapping.json"
    subprocess.run([sys.executable, str(ROOT / "scripts" / "map_tds_fields.py"), "--cn", str(facts[0]), "--en", str(facts[1]), "--registry", str(ROOT / "mapping" / "template_field_registry.json"), "--output", str(output)], check=True)
    result = json.loads(output.read_text(encoding="utf-8"))
    assert result["status"] == "ready"
    assert len(result["performance_extra_rows"]) == 1
    assert set(result["performance_extra_rows"][0]["label_values"]) == {"zh-CN", "en-US"}


def test_common_cn_performance_labels_map_to_fixed_slots(tmp_path):
    facts = tmp_path / "cn.json"
    facts.write_text(json.dumps({
        "language": "zh-CN",
        "title": {"text": "PU-1001"},
        "sections": {},
        "performance_rows": [
            {"item": "外观", "value": "蓝光透明液体", "unit": "", "test_method": "目测", "source_column_count": 4, "source_location": "table[0].row[1]"},
            {"item": "固体份含量", "value": "22±1", "unit": "%", "test_method": "方法", "source_column_count": 4, "source_location": "table[0].row[2]"},
            {"item": "粘度(25℃)", "value": "＜500", "unit": "mPa·s", "test_method": "方法", "source_column_count": 4, "source_location": "table[0].row[3]"},
            {"item": "PH值（1:10稀释在水中）", "value": "7.0-9.0", "unit": "", "test_method": "方法", "source_column_count": 4, "source_location": "table[0].row[4]"},
            {"item": "密度(25℃)", "value": "约1.06", "unit": "g/cm3", "test_method": "方法", "source_column_count": 4, "source_location": "table[0].row[5]"}
        ]
    }, ensure_ascii=False), encoding="utf-8")
    output = tmp_path / "mapping.json"
    subprocess.run([sys.executable, str(ROOT / "scripts" / "map_tds_fields.py"), "--cn", str(facts), "--registry", str(ROOT / "mapping" / "template_field_registry.json"), "--output", str(output)], check=True)
    result = json.loads(output.read_text(encoding="utf-8"))
    assert result["mapped_fields"]["performance.appearance"]["values"]["zh-CN"] == "蓝光透明液体"
    assert result["mapped_fields"]["performance.solid_content"]["values"]["zh-CN"] == "22±1"
    assert result["mapped_fields"]["performance.viscosity_25c"]["values"]["zh-CN"] == "＜500"
    assert "zh-CN" not in result["mapped_fields"]["performance.ph_25c"]["values"]
    assert [x["label_values"]["zh-CN"] for x in result["performance_extra_rows"]] == ["PH值（1:10稀释在水中）", "密度(25℃)"]


def test_mapping_keeps_evidence_separate_from_normalized_model(tmp_path):
    facts = tmp_path / "cn.json"
    facts.write_text(json.dumps({
        "language": "zh-CN",
        "title": {"text": "TDS-DEMO"},
        "sections": {"product.description": {"text": "用于水性涂层。", "locations": [3]}},
        "performance_rows": [{"item": "PH值（1:10稀释在水中）", "value": "7.0-9.0", "unit": "", "test_method": "方法A", "source_column_count": 4, "source_location": "table[0].row[4]"}]
    }, ensure_ascii=False), encoding="utf-8")
    output = tmp_path / "mapping.json"
    subprocess.run([sys.executable, str(ROOT / "scripts" / "map_tds_fields.py"), "--cn", str(facts), "--registry", str(ROOT / "mapping" / "template_field_registry.json"), "--output", str(output)], check=True)
    result = json.loads(output.read_text(encoding="utf-8"))
    row = result["normalized_model"]["performance_rows"][0]
    assert row["source_values"]["zh-CN"] == "7.0-9.0"
    assert row["normalized_values"]["zh-CN"] == "7.0-9.0"
    assert result["normalized_model"]["translation"]["source"] == "normalized_model"
    decision = next(item for item in result["decision_ledger"] if item["field_id"] == "performance.row.001")
    assert decision["decision"] == "preserve_as_source_row"
    assert decision["provenance"]["zh-CN"] == ["table[0].row[4]"]


def test_mapping_records_verbatim_source_hashes(tmp_path):
    facts = tmp_path / "cn.json"
    facts.write_text(json.dumps({
        "language": "zh-CN",
        "title": {"text": "PU-1001"},
        "sections": {"product.description": {"text": "成膜非常柔软。", "locations": [3]}},
        "performance_rows": [{"item": "外观", "value": "蓝光透明液体", "unit": "", "test_method": "目测", "source_column_count": 4, "source_location": "table[0].row[1]"}]
    }, ensure_ascii=False), encoding="utf-8")
    output = tmp_path / "mapping.json"
    subprocess.run([sys.executable, str(ROOT / "scripts" / "map_tds_fields.py"), "--cn", str(facts), "--registry", str(ROOT / "mapping" / "template_field_registry.json"), "--output", str(output)], check=True)
    result = json.loads(output.read_text(encoding="utf-8"))
    assert result["source_fidelity"]["field_sha256"]["product.description"]["zh-CN"] == text_sha256("成膜非常柔软。")
    assert result["source_fidelity"]["performance_row_sha256"]["performance.row.001"]["zh-CN"]["value"] == text_sha256("蓝光透明液体")
    assert source_fidelity_errors(result) == []


def test_generated_cn_output_gate_detects_text_mutation(tmp_path):
    registry = load(ROOT / "mapping" / "template_field_registry.json")
    mapping = load(ROOT / "tests" / "fixtures" / "valid_mapping.json")
    output = tmp_path / "fidelity.docx"
    write_variant(mapping, registry, "TDS_CN_冠志模板", output)
    doc = Document(str(output))
    paragraph = next(p for p in doc.paragraphs if p.text == "用于水性涂层的聚氨酯分散体。")
    paragraph.runs[0].text = "用于水性涂层的聚氨酯树脂。"
    doc.save(str(output))
    errors = source_output_fidelity_errors(Document(str(output)), mapping, registry, "TDS_CN_冠志模板")
    assert "field_output_mismatch:zh-CN:product.description" in errors


def test_overwrite_blocks_cn_normalized_values_that_change_source(tmp_path):
    registry = load(ROOT / "mapping" / "template_field_registry.json")
    mapping = deepcopy(load(ROOT / "tests" / "fixtures" / "valid_mapping.json"))
    mapping["schema_version"] = "1.3.0"
    mapping["normalized_model"] = {
        "status": "approved",
        "translation": {"source": "normalized_model"},
        "fields": {
            field_id: {"field_id": field_id, "normalized_values": dict(item.get("values", {}))}
            for field_id, item in mapping["mapped_fields"].items()
        },
        "performance_rows": [
            {
                **row,
                "normalized_label_values": {"zh-CN": "标准化外观"},
                "normalized_values": {"zh-CN": "标准化液体"},
                "normalized_unit_values": {"zh-CN": ""},
                "normalized_test_method_values": {"zh-CN": "标准化方法"},
            }
            for row in [{"field_id": "performance.row.001", "label_values": {"zh-CN": "乳液外观"}, "values": {"zh-CN": "乳白色液体"}}]
        ],
    }
    output = tmp_path / "normalized.docx"
    with pytest.raises(RuntimeError, match="source fidelity contract"):
        write_variant(mapping, registry, "TDS_CN_冠志模板", output)
    assert not output.exists()


def test_numbered_features_do_not_duplicate_template_numbering(tmp_path):
    registry = load(ROOT / "mapping" / "template_field_registry.json")
    mapping = deepcopy(load(ROOT / "tests" / "fixtures" / "valid_mapping.json"))
    mapping["mapped_fields"]["product.features"]["values"]["zh-CN"] = "1. 成膜柔软；\n2. 耐黄变；"
    output = tmp_path / "numbered.docx"
    write_variant(mapping, registry, "TDS_CN_冠志模板", output)
    text = "\n".join(p.text for p in Document(str(output)).paragraphs)
    assert "1. 1." not in text
    assert "2. 2." not in text


def test_source_led_performance_rows_keep_order_and_conditions(tmp_path):
    registry = load(ROOT / "mapping" / "template_field_registry.json")
    mapping = deepcopy(load(ROOT / "tests" / "fixtures" / "valid_mapping.json"))
    mapping["performance_rows"] = [
        {"field_id": "performance.row.001", "label_values": {"zh-CN": "外观"}, "values": {"zh-CN": "蓝光透明液体"}, "unit_values": {"zh-CN": ""}, "test_method_values": {"zh-CN": "目测"}},
        {"field_id": "performance.row.002", "label_values": {"zh-CN": "固体份含量"}, "values": {"zh-CN": "22±1"}, "unit_values": {"zh-CN": "%"}, "test_method_values": {"zh-CN": "150℃ 30min，鼓风烘箱"}},
        {"field_id": "performance.row.003", "label_values": {"zh-CN": "粘度(25℃)"}, "values": {"zh-CN": "＜500"}, "unit_values": {"zh-CN": "mPa·S"}, "test_method_values": {"zh-CN": "GB/T 2794-2022"}},
        {"field_id": "performance.row.004", "label_values": {"zh-CN": "PH值（1:10稀释在水中）"}, "values": {"zh-CN": "7.0-9.0"}, "unit_values": {"zh-CN": ""}, "test_method_values": {"zh-CN": "GB 6920-86"}},
        {"field_id": "performance.row.005", "label_values": {"zh-CN": "密度(25℃)"}, "values": {"zh-CN": "约1.06"}, "unit_values": {"zh-CN": "g/cm3"}, "test_method_values": {"zh-CN": "GB/T 4472-2011"}}
    ]
    output = tmp_path / "source-led.docx"
    write_variant(mapping, registry, "TDS_CN_冠志模板", output)
    rows = [[c.text for c in row.cells] for row in Document(str(output)).tables[0].rows]
    assert len(rows) == 6
    assert rows[1:] == [
        ["外观", "蓝光透明液体", "", "目测"],
        ["固体份含量", "22±1", "%", "150℃ 30min，鼓风烘箱"],
        ["粘度(25℃)", "＜500", "mPa·S", "GB/T 2794-2022"],
        ["PH值（1:10稀释在水中）", "7.0-9.0", "", "GB 6920-86"],
        ["密度(25℃)", "约1.06", "g/cm3", "GB/T 4472-2011"]
    ]


def test_source_led_table_trims_unused_template_rows(tmp_path):
    registry = load(ROOT / "mapping" / "template_field_registry.json")
    mapping = deepcopy(load(ROOT / "tests" / "fixtures" / "valid_mapping.json"))
    mapping["performance_rows"] = [{"field_id": "performance.row.001", "label_values": {"zh-CN": "外观"}, "values": {"zh-CN": "液体"}, "unit_values": {"zh-CN": ""}, "test_method_values": {"zh-CN": "目测"}}]
    output = tmp_path / "trimmed.docx"
    write_variant(mapping, registry, "TDS_CN_冠志模板", output)
    assert len(Document(str(output)).tables[0].rows) == 2


def test_hidden_no_source_section_is_removed(tmp_path):
    registry = load(ROOT / "mapping" / "template_field_registry.json")
    mapping = deepcopy(load(ROOT / "tests" / "fixtures" / "valid_mapping.json"))
    mapping["schema_version"] = "1.3.0"
    mapping["status"] = "ready"
    mapping["normalized_model"] = {
        "status": "approved",
        "translation": {"source": "normalized_model"},
        "fields": {
            field_id: {"field_id": field_id, "source_values": dict(item.get("values", {})), "normalized_values": dict(item.get("values", {})), "provenance": {}}
            for field_id, item in mapping["mapped_fields"].items()
        },
        "performance_rows": None,
        "decision_ledger": [{"field_id": "product.supply_form", "decision": "hide_no_source", "needs_judgment": False, "provenance": {}}],
    }
    mapping["normalized_model"]["fields"]["product.supply_form"] = {"field_id": "product.supply_form", "source_values": {}, "normalized_values": {}, "provenance": {}}
    for variant_id in ("TDS_CN_冠志模板", "TDS_EN_冠志模板"):
        output = tmp_path / f"hidden_{variant_id}.docx"
        write_variant(mapping, registry, variant_id, output)
        paragraphs = "\n".join(p.text for p in Document(str(output)).paragraphs)
        heading = "【供应形式】" if registry["variants"][variant_id]["language"] == "zh-CN" else "【Supply Form】"
        assert heading not in paragraphs
        assert "无数据" not in paragraphs and "No data available" not in paragraphs
        base = Document(str(ROOT / registry["variants"][variant_id]["template"]))
        assert audit_shape(base, Document(str(output)), registry["variants"][variant_id], mapping)


def test_empty_text_section_becomes_hide_candidate(tmp_path):
    facts = tmp_path / "cn.json"
    facts.write_text(json.dumps({
        "language": "zh-CN",
        "title": {"text": "TDS-DEMO"},
        "sections": {"product.description": {"text": "用于水性涂层。", "locations": [3]}},
        "performance_rows": [{"item": "外观", "value": "液体", "unit": "", "test_method": "目测", "source_column_count": 4, "source_location": "table[0].row[1]"}]
    }, ensure_ascii=False), encoding="utf-8")
    output = tmp_path / "mapping.json"
    subprocess.run([sys.executable, str(ROOT / "scripts" / "map_tds_fields.py"), "--cn", str(facts), "--registry", str(ROOT / "mapping" / "template_field_registry.json"), "--output", str(output)], check=True)
    result = json.loads(output.read_text(encoding="utf-8"))
    decision = next(item for item in result["decision_ledger"] if item["field_id"] == "product.supply_form")
    assert decision["decision"] == "hide_no_source_candidate"
    assert decision["needs_judgment"] is True


def test_packaging_storage_heading_starts_storage_section(tmp_path):
    source = tmp_path / "source.docx"
    doc = Document()
    doc.add_paragraph("【产品特性】")
    doc.add_paragraph("耐水。")
    doc.add_paragraph("【包装储存】")
    doc.add_paragraph("密封保存。")
    doc.save(str(source))
    output = tmp_path / "facts.json"
    subprocess.run([sys.executable, str(ROOT / "scripts" / "extract_tds_source.py"), str(source), "--language", "zh-CN", "--output", str(output)], check=True)
    result = json.loads(output.read_text(encoding="utf-8"))
    assert result["sections"]["product.features"]["text"] == "耐水。"
    assert result["sections"]["product.storage"]["text"] == "密封保存。"


def test_title_preserves_template_positioning_and_style(tmp_path):
    registry = load(ROOT / "mapping" / "template_field_registry.json")
    mapping = deepcopy(load(ROOT / "tests" / "fixtures" / "valid_mapping.json"))
    title = mapping["mapped_fields"]["product.title"]["values"]["zh-CN"]
    output = tmp_path / "title.docx"
    write_variant(mapping, registry, "TDS_CN_冠志模板", output)
    doc = Document(str(output))
    paragraph = next(p for p in doc.paragraphs if p.text == title)
    base = Document(str(ROOT / registry["variants"]["TDS_CN_冠志模板"]["template"]))
    title_index = next(item["locator"]["paragraph_index"] for item in registry["variants"]["TDS_CN_冠志模板"]["slots"] if item["field_id"] == "product.title")
    base_title = base.paragraphs[title_index]
    from tds_common import paragraph_shape
    assert paragraph_shape(paragraph) == paragraph_shape(base_title)


def _mapping_with_hidden_sections():
    mapping = deepcopy(load(ROOT / "tests" / "fixtures" / "valid_mapping.json"))
    for fid in ("product.supply_form", "product.application"):
        mapping["mapped_fields"][fid]["values"] = {}
        mapping["mapped_fields"][fid]["source_values"] = {}
        mapping["mapped_fields"][fid]["sources"] = {}
    mapping["normalized_model"] = {
        "status": "approved",
        "translation": {"source": "normalized_model"},
        "fields": {
            field_id: {"field_id": field_id, "normalized_values": dict(item.get("values", {}))}
            for field_id, item in mapping["mapped_fields"].items()
        },
        "performance_rows": None,
        "decision_ledger": [
            {"field_id": "product.supply_form", "decision": "hide_no_source", "needs_judgment": False},
            {"field_id": "product.application", "decision": "hide_no_source", "needs_judgment": False},
        ],
    }
    return mapping


def test_hidden_sections_remove_adjacent_blank_separators(tmp_path):
    registry = load(ROOT / "mapping" / "template_field_registry.json")
    mapping = _mapping_with_hidden_sections()
    output = tmp_path / "hidden.docx"
    write_variant(mapping, registry, "TDS_CN_冠志模板", output)
    doc = Document(str(output))
    texts = [p.text for p in doc.paragraphs]
    assert "【应用】" not in texts and "【供应形式】" not in texts
    assert "Application" not in texts
    head = next(i for i, p in enumerate(doc.paragraphs) if p.text.strip() == "【产品特性】")
    storage = next(i for i, p in enumerate(doc.paragraphs) if p.text.strip() == "【储存】")
    region = doc.paragraphs[head + 1:storage]
    assert not any(not region[i].text.strip() and not region[i + 1].text.strip() for i in range(len(region) - 1))


def test_multiline_storage_clones_template_body_style(tmp_path):
    registry = load(ROOT / "mapping" / "template_field_registry.json")
    mapping = deepcopy(load(ROOT / "tests" / "fixtures" / "valid_mapping.json"))
    mapping["mapped_fields"]["product.storage"]["values"]["zh-CN"] = "密封储存。\n有效期6个月。"
    output = tmp_path / "storage.docx"
    write_variant(mapping, registry, "TDS_CN_冠志模板", output)
    doc = Document(str(output))
    heading = next(i for i, p in enumerate(doc.paragraphs) if p.text.strip() == "【储存】")
    lines = [p.text for p in doc.paragraphs[heading + 1:heading + 4] if p.text.strip()]
    assert lines[:2] == ["密封储存。", "有效期6个月。"]
    base = Document(str(ROOT / registry["variants"]["TDS_CN_冠志模板"]["template"]))
    from tds_common import paragraph_shape
    body = [p for p in doc.paragraphs[heading + 1:] if p.text.strip()]
    storage_index = next(item["locator"]["paragraph_index"] for item in registry["variants"]["TDS_CN_冠志模板"]["slots"] if item["field_id"] == "product.storage")
    assert paragraph_shape(body[0]) == paragraph_shape(base.paragraphs[storage_index])
    assert paragraph_shape(body[1]) == paragraph_shape(base.paragraphs[storage_index])


def test_no_intra_paragraph_line_breaks_in_any_variant(tmp_path):
    registry = load(ROOT / "mapping" / "template_field_registry.json")
    mapping = deepcopy(load(ROOT / "tests" / "fixtures" / "valid_mapping.json"))
    for variant_id in ("TDS_CN_冠志模板", "TDS_EN_冠志模板"):
        output = tmp_path / f"breaks_{variant_id}.docx"
        write_variant(mapping, registry, variant_id, output)
        doc = Document(str(output))
        assert all("\n" not in p.text and "\r" not in p.text for p in doc.paragraphs)
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    assert all("\n" not in p.text and "\r" not in p.text for p in cell.paragraphs)


def test_body_blank_trim_removes_placeholder_rows_and_section_gaps(tmp_path):
    registry = load(ROOT / "mapping" / "template_field_registry.json")
    mapping = deepcopy(load(ROOT / "tests" / "fixtures" / "valid_mapping.json"))
    for variant_id, variant in registry["variants"].items():
        output = tmp_path / f"body-blank-{variant_id}.docx"
        write_variant(mapping, registry, variant_id, output)
        doc = Document(str(output))
        assert body_blank_count(doc, variant) == 0
        start = next(i for i, p in enumerate(doc.paragraphs) if p.text.strip() in {"【产品描述】", "Product Description"})
        assert all(paragraph.text.strip() for paragraph in doc.paragraphs[start:])
        generation = json.loads(output.with_suffix(output.suffix + ".generation.json").read_text(encoding="utf-8"))
        assert generation["body_blank_paragraph_count"] == 0
        assert audit_shape(Document(str(ROOT / variant["template"])), doc, variant, mapping)


def test_english_vertical_budget_preserves_template_indent(tmp_path):
    registry = load(ROOT / "mapping" / "template_field_registry.json")
    mapping = deepcopy(load(ROOT / "tests" / "fixtures" / "valid_mapping.json"))
    values = mapping["mapped_fields"]
    values["product.description"]["values"]["en-US"] = "Description line one.\nDescription line two."
    values["product.supply_form"]["values"]["en-US"] = "Supply line one.\nSupply line two."
    values["product.features"]["values"]["en-US"] = "Feature one.\nFeature two.\nFeature three.\nFeature four.\nFeature five."
    values["product.application"]["values"]["en-US"] = "Application line one.\nApplication line two.\nApplication line three.\nApplication line four.\nApplication line five."
    values["product.storage"]["values"]["en-US"] = "Storage line one.\nStorage line two.\nStorage line three."
    variant = registry["variants"]["TDS_EN_冠志模板"]
    output = tmp_path / "english-budget.docx"
    write_variant(mapping, registry, "TDS_EN_冠志模板", output)
    doc = Document(str(output))
    headings = {"Product Description", "Supply Form", "Product Features", "Application", "Storage"}
    body_start = next(i for i, p in enumerate(doc.paragraphs) if p.text.strip() == "Product Description")
    body = [p for p in doc.paragraphs[body_start:] if p.text.strip() and p.text.strip() not in headings]
    assert body
    generation = json.loads(output.with_suffix(output.suffix + ".generation.json").read_text(encoding="utf-8"))
    expected = {1: ("280", "30"), 2: ("260", "20")} [generation["english_vertical_budget_level"]]
    for paragraph in body:
        spacing = paragraph._p.pPr.find(qn("w:spacing"))
        assert spacing is not None and spacing.get(qn("w:line")) == expected[0] and spacing.get(qn("w:after")) == expected[1]
    ind = body[0]._p.pPr.find(qn("w:ind"))
    assert ind is not None and ind.get(qn("w:firstLineChars")) == "200" and ind.get(qn("w:firstLine")) == "480"
    feature = next(p for p in body if p.text == "Feature one.")
    feature_ind = feature._p.pPr.find(qn("w:ind"))
    assert feature_ind is not None and feature_ind.get(qn("w:left")) == "840" and feature_ind.get(qn("w:hanging")) == "360"
    assert generation["english_vertical_budget_applied"] is True
    assert audit_shape(Document(str(ROOT / variant["template"])), doc, variant, mapping)


def test_pdf_page_count_uses_pure_python_fallback_when_pypdf_is_unavailable(tmp_path, monkeypatch):
    pdf = tmp_path / "fallback.pdf"
    pdf.write_bytes(b"%PDF-1.4\n1 0 obj << /Type /Page >>\n2 0 obj << /Type /Page >>\n")
    monkeypatch.setitem(sys.modules, "pypdf", None)
    assert pdf_page_count(pdf) == 2
