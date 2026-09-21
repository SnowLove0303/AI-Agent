import base64
from pathlib import Path

import pytest
from docx import Document
from docx.shared import Inches

import sys
sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "scripts"))

from ghs_pictogram_policy import insert_source_pictogram
from section2_ghs_policy import (
    apply_precautionary_layout,
    format_label_elements,
    is_missing_section2_value,
    is_explicit_other_hazards_row,
    validate_s2_semantics,
    project_source_cn_facts,
    project_source_cn_headings,
    suppress_missing_section2_rows_and_renumber,
    normalize_non_hazard_category,
    normalize_pictogram_value,
)
from template_mutation_whitelist import (
    TemplateSlotRegistry,
    clear_value_cells,
    write_row_values,
)


ROOT = Path(__file__).resolve().parents[1]


def test_label_elements_are_explicit_and_line_separated():
    assert format_label_elements("zh", ["亲水脂肪族聚异氰酸酯"]) == "必须列在标签上的有害成分：\n亲水脂肪族聚异氰酸酯"
    assert format_label_elements("en", ["Hydrophilic aliphatic polyisocyanate"]) == "Hazardous ingredients required to be listed on the label:\nHydrophilic aliphatic polyisocyanate"


def test_non_hazard_and_pictogram_values_use_approved_customer_text():
    assert normalize_non_hazard_category("无") == "根据 GHS 不属于危险物"
    assert normalize_non_hazard_category("") == "根据 GHS 不属于危险物"
    assert normalize_non_hazard_category("类别 3") == "类别 3"
    assert normalize_pictogram_value("无") == "无象形图"
    assert normalize_pictogram_value("") == "无象形图"


def test_precautionary_layout_splits_headings_and_indents_details():
    document = Document(ROOT / "examples" / "template_reference.docx")
    row = document.tables[1].rows[7]
    row.cells[-1].text = "预防措施：\n第一条。\n第二条。\n事故响应：\n第三条。"
    result = apply_precautionary_layout(document)
    paragraphs = document.tables[1].rows[7].cells[-1].paragraphs
    assert result["changed"] is True
    assert [p.text for p in paragraphs] == ["预防措施：", "第一条。", "第二条。", "事故响应：", "第三条。"]
    assert paragraphs[0].paragraph_format.left_indent == 0
    assert paragraphs[1].paragraph_format.left_indent.pt == 18
    assert paragraphs[3].paragraph_format.left_indent == 0


def test_label_elements_without_ingredients_leave_empty_value_for_suppression():
    assert format_label_elements("zh", []) == ""
    assert format_label_elements("en", []) == ""
    document = Document(ROOT / "examples" / "template_reference.docx")
    table = document.tables[1]
    table.rows[3].cells[-1].text = format_label_elements("zh", [])
    result = suppress_missing_section2_rows_and_renumber(document, lambda p, text: setattr(p, "text", text))
    labels = [row.cells[0].text.strip() for row in document.tables[1].rows[1:]]
    assert not any("标签要素" in label for label in labels)
    assert result["removed_count"] >= 1


def test_section2_missing_rows_are_removed_and_unique_items_renumbered():
    document = Document(ROOT / "examples" / "template_reference.docx")
    table = document.tables[1]
    values = {
        1: "无数据",
        2: "易燃液体，类别3（H226）",
        3: format_label_elements("zh", ["亲水脂肪族聚异氰酸酯"]),
        4: "",
        5: "警告",
        6: "H226 易燃液体和蒸气。",
        7: "P210 远离热源。",
        8: "无数据",
        9: "吸入：吸入有害。",
        10: "食入：无数据",
        11: "皮肤：可能致敏。",
        12: "眼睛：无数据",
        13: "症状和体征：无数据",
        14: "无数据",
        15: "存在风险。",
    }
    for row_index, value in values.items():
        table.rows[row_index].cells[-1].text = value

    result = suppress_missing_section2_rows_and_renumber(document, lambda p, text: setattr(p, "text", text))
    labels = [row.cells[0].text.strip() for row in document.tables[1].rows[1:]]
    assert result["removed_count"] == 7
    assert labels == [
        "2.1  GHS危险性类别：",
        "2.2  GHS标签要素：",
        "2.3  信号词：",
        "2.4  危险性说明：",
        "2.5  防范说明：",
        "2.6  健康危害",
        "2.6  健康危害",
        "2.7  其他危害",
    ]
    assert is_missing_section2_value("眼睛：无数据")
    assert not is_missing_section2_value("眼睛：无刺激")


def test_s28_prefix_only_rows_are_removed_after_whitelist_clear():
    document = Document(ROOT / "examples" / "template_reference.docx")
    registry = TemplateSlotRegistry.from_document(document)
    clear_value_cells(document, registry)
    for row_index in range(9, 14):
        row = document.tables[1].rows[row_index]
        write_row_values(
            row, [row.cells[0].text, ""],
            table_index=1, row_index=row_index, registry=registry,
        )
    result = suppress_missing_section2_rows_and_renumber(
        document, lambda paragraph, text: setattr(paragraph, "text", text)
    )
    assert result["removed_count"] >= 5
    assert not any(
        row.cells[0].text.strip().startswith("2.8")
        for row in document.tables[1].rows[1:]
    )


def test_section2_explicit_number_map_keeps_source_requested_numbers():
    document = Document(ROOT / "examples" / "template_reference.docx")
    table = document.tables[1]
    for row in table.rows[1:]:
        row.cells[-1].text = ""
    table.rows[2].cells[-1].text = ""
    table.rows[3].cells[-1].text = "根据GHS不属于危害化学品"
    table.rows[15].cells[-1].text = "无适用资料。"
    result = suppress_missing_section2_rows_and_renumber(
        document, lambda p, text: setattr(p, "text", text), number_map={3: 2, 10: 3}
    )
    labels = [row.cells[0].text.strip() for row in document.tables[1].rows[1:]]
    assert labels == ["2.2  GHS标签要素：", "2.3  其他危害"]
    assert result["number_map"] == {"2.3": "2.2", "2.10": "2.3"}


def test_blank_other_hazards_exception_is_suppressed_but_explicit_missing_is_kept():
    document = Document(ROOT / "examples" / "template_reference.docx")
    table = document.tables[1]
    other_row = table.rows[15]
    other_row.cells[-1].text = ""
    assert not is_explicit_other_hazards_row(other_row)
    result = suppress_missing_section2_rows_and_renumber(
        document, lambda p, text: setattr(p, "text", text), number_map={3: 2, 10: 3}
    )
    assert "2.10  其他危害" in result["removed_labels"] or "2.10 其他危害" in result["removed_labels"]

    explicit = Document(ROOT / "examples" / "template_reference.docx")
    explicit_row = explicit.tables[1].rows[15]
    explicit_row.cells[-1].text = "无适用资料。"
    assert is_explicit_other_hazards_row(explicit_row)


def test_source_cn_s2_projects_by_semantic_slot_not_list_position():
    rows, number_map = project_source_cn_facts({
        "ghs_classes": ["根据 GHS 不属于危害化学品"],
        "label_elements": ["根据 GHS 不属于危害化学品"],
        "other_hazards": "无适用资料。",
    })
    assert rows[1] == ["2.2 GHS危险性类别：", "根据 GHS 不属于危害化学品"]
    assert rows[2] == ["2.3 GHS标签要素：", "根据 GHS 不属于危害化学品"]
    assert rows[14] == ["2.10 其他危害：", "无适用资料。"]
    assert number_map == {2: 1, 3: 2, 10: 3}


def test_source_cn_s2_projects_grouped_precautionary_value_and_suppresses_empty_group():
    rows, _ = project_source_cn_facts({
        "precautionary_groups": [
            {
                "group_key": "prevention",
                "source_heading": "\u9884\u9632\u63aa\u65bd\uff1a",
                "statements": [{"code": "P210", "text": "P210 \u8fdc\u79bb\u70ed\u6e90\u3002"}],
            },
            {
                "group_key": "storage",
                "source_heading": "\u5b89\u5168\u50a8\u5b58\uff1a",
                "statements": [],
            },
            {
                "group_key": "disposal",
                "source_heading": "\u5e9f\u5f03\u5904\u7f6e\uff1a",
                "statements": [{"code": "P501", "text": "P501 \u6309\u7167\u89c4\u5b9a\u5904\u7f6e\u3002"}],
            },
        ],
    })
    precautionary = rows[6][1]
    assert precautionary.splitlines() == [
        "\u9884\u9632\u63aa\u65bd\uff1a",
        "P210 \u8fdc\u79bb\u70ed\u6e90\u3002",
        "\u5e9f\u5f03\u5904\u7f6e\uff1a",
        "P501 \u6309\u7167\u89c4\u5b9a\u5904\u7f6e\u3002",
    ]


def test_release_audit_blocks_missing_or_out_of_order_precautionary_groups():
    from audit_section2_release import run

    document = Document(ROOT / "examples" / "template_reference.docx")
    table = document.tables[1]
    table.rows[7].cells[-1].text = (
        "\u9884\u9632\u63aa\u65bd\uff1a\nP210 \u8fdc\u79bb\u70ed\u6e90\u3002\n"
        "\u5e9f\u5f03\u5904\u7f6e\uff1a\nP501 \u6309\u89c4\u5b9a\u5904\u7f6e\u3002"
    )
    errors, info = run(
        None,
        document=document,
        expected_precautionary_groups=["prevention", "response"],
    )
    assert not info["pass"]
    assert any("precautionary group headings are missing or out of order" in error for error in errors)


def test_release_audit_rejects_template_precautionary_group_when_source_has_none():
    from audit_section2_release import run

    document = Document(ROOT / "examples" / "template_reference.docx")
    errors, _info = run(
        None,
        document=document,
        expected_precautionary_groups=[],
    )
    assert any("precautionary group headings are missing or out of order" in error for error in errors)


def test_en_drafter_protects_controlled_precautionary_heading_from_generic_glossary():
    from draft_en_facts import draft_section

    review = []
    rows = [["2.6 \u9632\u8303\u8bf4\u660e\uff1a", "\u9884\u9632\u63aa\u65bd\uff1a\nP210 \u8fdc\u79bb\u70ed\u6e90\u3002"]]
    drafted = draft_section(rows, [("\u9884\u9632\u63aa\u65bd", "Wrong glossary heading")], "s2", review)
    assert drafted[0][1].splitlines()[0] == "Prevention:"


def test_s2_semantic_contract_separates_label_ingredients_and_signal_word():
    rows = [
        ["2.1 紧急情况概述", ""],
        ["2.2 GHS危险性类别：", "类别 3"],
        ["2.3 GHS标签要素：", "必须列在标签上的有害成分：\n基于HDI的亲水脂肪族聚异氰酸酯"],
        [" GHS象形图", ""],
        ["2.4 信号词：", "警告"],
    ]
    assert validate_s2_semantics(rows, "zh") == []

    misplaced = [list(row) for row in rows]
    misplaced[2][1] = "警告"
    misplaced[4][1] = "必须列在标签上的有害成分：\n基于HDI的亲水脂肪族聚异氰酸酯"
    errors = validate_s2_semantics(misplaced, "zh")
    assert any("label-elements slot contains only a signal word" in error for error in errors)
    assert any("signal-word slot contains label-ingredient prose" in error for error in errors)


@pytest.mark.parametrize(
    "signal", ["无信号词", "无", "None", "No signal word", "Not applicable"]
)
def test_s2_semantic_contract_accepts_controlled_non_hazard_signal_values(signal):
    rows = [
        ["2.1 紧急情况概述", ""],
        ["2.2 GHS危险性类别：", "根据 GHS 不属于危害化学品"],
        ["2.3 GHS标签要素：", ""],
        [" GHS象形图", ""],
        ["2.4 信号词：", signal],
    ]
    assert validate_s2_semantics(rows, "zh") == []


def test_source_pictogram_is_inserted_as_picture(tmp_path):
    image_path = tmp_path / "source.png"
    # A tiny valid PNG keeps this regression test independent of Pillow.
    image_path.write_bytes(base64.b64decode(
        "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mNk+A8AAQUBAScY42YAAAAASUVORK5CYII="
    ))
    source = tmp_path / "source.docx"
    source_doc = Document()
    source_doc.add_paragraph().add_run().add_picture(str(image_path))
    source_doc.save(source)

    output = Document(ROOT / "examples" / "template_reference.docx")
    audit = insert_source_pictogram(output, source)
    assert output.tables[1].rows[4].cells[-1]._tc.xpath(".//w:drawing")
    assert audit["target_width_inches"] <= 1.0


def test_source_pictogram_preserves_declared_physical_width(tmp_path):
    image_path = tmp_path / "source.png"
    image_path.write_bytes(base64.b64decode(
        "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mNk+A8AAQUBAScY42YAAAAASUVORK5CYII="
    ))
    source = tmp_path / "source-sized.docx"
    source_doc = Document()
    source_doc.add_paragraph().add_run().add_picture(str(image_path), width=Inches(0.89))
    source_doc.save(source)

    output = Document(ROOT / "examples" / "template_reference.docx")
    audit = insert_source_pictogram(output, source)
    assert audit["source_width_inches"] == pytest.approx(0.89, abs=0.002)
    assert audit["target_width_inches"] == pytest.approx(0.89, abs=0.002)


def test_cn_source_headings_are_projected_without_layout_drift():
    template = Document(ROOT / "examples" / "template_reference.docx")
    output = Document(ROOT / "examples" / "template_reference.docx")
    try:
        project_source_cn_headings(output)
    except Exception as exc:
        assert "locked Section 2 labels are immutable" in str(exc)
    else:
        raise AssertionError("legacy Section 2 label-rewrite path was not blocked")
    assert [row.cells[0].text for row in output.tables[1].rows] == [
        row.cells[0].text for row in template.tables[1].rows
    ]


def test_non_hazard_section2_suppression_and_release_audit():
    """Verify that non-hazardous products have absent H/P/route rows suppressed,
    surviving items renumbered continuously (2.1, 2.2, 2.3), and illegal CMR P-statements blocked.
    """
    from audit_section2_release import run as audit_s2
    document = Document(ROOT / "examples" / "template_reference.docx")
    table = document.tables[1]

    # Populate non-hazard inputs: Category 2.2 populated as non-hazard, signal word, other hazards;
    # all H statements, P statements, health routes left empty.
    table.rows[1].cells[-1].text = ""  # 2.1 emergency overview empty -> suppressed
    table.rows[2].cells[-1].text = "未被分类"  # 2.2 GHS category -> becomes 2.1
    table.rows[3].cells[-1].text = ""  # 2.3 label elements empty -> suppressed
    table.rows[4].cells[-1].text = "无危险的象形图"  # pictogram
    table.rows[5].cells[-1].text = "无信号词"  # 2.4 signal word -> becomes 2.2
    table.rows[6].cells[-1].text = ""  # 2.5 H statements empty -> suppressed
    table.rows[7].cells[-1].text = ""  # 2.6 P statements empty -> suppressed
    table.rows[8].cells[-1].text = ""  # 2.7 physical/chemical empty -> suppressed
    for r in range(9, 14):
        table.rows[r].cells[-1].text = ""  # 2.8 health hazard routes empty -> suppressed
    table.rows[14].cells[-1].text = ""  # 2.9 environmental empty -> suppressed
    table.rows[15].cells[-1].text = "无适用资料。"  # 2.10 other hazards -> becomes 2.3

    result = suppress_missing_section2_rows_and_renumber(
        document, lambda p, text: setattr(p, "text", text)
    )
    labels = [row.cells[0].text.strip() for row in document.tables[1].rows[1:]]
    assert labels == [
        "2.1  GHS危险性类别：",
        "GHS象形图：",
        "2.2  信号词：",
        "2.3  其他危害",
    ]
    errors, info = audit_s2("dummy.docx", document=document)
    assert not errors, f"Expected audit pass on clean non-hazard S2, got: {errors}"

    # Invalidate by adding prohibited P405 to non-hazardous substance
    table.rows[-1].cells[-1].text = "P405 储存处须加锁。"
    errors, _ = audit_s2("dummy.docx", document=document)
    assert any("P405" in err for err in errors)
