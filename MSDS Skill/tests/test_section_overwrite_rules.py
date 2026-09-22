from pathlib import Path
import sys

import pytest
from docx import Document

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "scripts"))

from msds_pipeline import (  # noqa: E402
    ReleaseBlocked,
    align_note_section_rows,
    align_s9_rows,
    validate_template_baselines,
)
from section2_ghs_policy import (  # noqa: E402
    normalize_non_hazard_category,
    normalize_ghs_classification_text,
    normalize_s2_projected_rows,
    sanitize_label_elements_text,
)
from section_overwrite_rules import (  # noqa: E402
    LOCAL_SECTION_POLICIES,
    SECTION_RULES,
    local_policy_for,
    SectionRuleViolation,
    validate_section_payload,
    validate_section_template,
)
from missing_data_policy import _payload_cells  # noqa: E402


def test_active_templates_match_all_sixteen_section_rules():
    assert set(SECTION_RULES) == set(range(1, 17))
    for language, filename in (("zh", "template_reference.docx"),
                               ("en", "template_reference_en.docx")):
        validate_section_template(Document(str(ROOT / "examples" / filename)), language)


def test_component_payload_cannot_use_positional_two_column_rows():
    template = Document(str(ROOT / "examples" / "template_reference.docx"))
    with pytest.raises(SectionRuleViolation, match="exactly name/CAS/content"):
        validate_section_payload(3, [["name", "CAS"]], template.tables[2])


def test_modified_active_template_cannot_become_the_audit_baseline(tmp_path):
    cn = ROOT / "examples" / "template_reference.docx"
    en = ROOT / "examples" / "template_reference_en.docx"
    en_source = ROOT / "examples" / "template_reference_en_source.docx"
    tampered = tmp_path / "template_reference.docx"
    tampered.write_bytes(cn.read_bytes() + b"tampered")
    with pytest.raises(ReleaseBlocked, match="template baseline"):
        validate_template_baselines(tampered, en, en_source)


def test_local_rules_are_semantic_and_cover_known_high_risk_sections():
    assert {2, 8, 9, 10, 11, 12, 13, 14} <= set(LOCAL_SECTION_POLICIES)
    assert local_policy_for(8)["match_basis"] == "ppe_and_control_meaning"
    assert local_policy_for(9)["match_basis"] == "property_alias"
    assert local_policy_for(11)["match_basis"] == "endpoint_study_field"


def test_section_11_2_three_column_middle_label_is_not_a_value():
    class Cell:
        def __init__(self, text):
            self.text = text

    cells = [Cell("11.2 毒性："), Cell("经口："), Cell("LD50 > 5000 mg/kg")]
    assert [cell.text for cell in _payload_cells(cells, 11)] == ["LD50 > 5000 mg/kg"]


def test_section2_keeps_explicit_category_and_line_breaks_special_note():
    assert normalize_non_hazard_category("GHS危险性类别: 无") == "无"
    assert sanitize_label_elements_text("请注意以下物质：，N,N-二甲基乙醇胺\n特定阈值浓度≥5%") == (
        "请注意以下物质：\nN,N-二甲基乙醇胺，特定阈值浓度≥5%"
    )


def test_section2_classification_keeps_h_code_with_class_and_corrects_reviewed_typo():
    assert normalize_ghs_classification_text(
        "依然液体3 H226；急性毒性4 吸入性 H332；皮肤腐蚀1B H314"
    ) == (
        "易燃液体3 H226\n急性毒性4 吸入性 H332\n皮肤腐蚀1B H314"
    )


def test_section2_h_and_p_values_split_before_the_guarded_writer():
    rows = normalize_s2_projected_rows([
        ["2.5 危险性说明：", "H226 易燃液体。 H332 吸入有害。"],
        ["2.6 防范说明：", "预防措施： P280 戴防护手套。 P270 操作时不得进食。"],
    ])
    assert rows == [
        ["2.5 危险性说明：", "H226 易燃液体。\nH332 吸入有害。"],
        ["2.6 防范说明：", "预防措施：\nP280 戴防护手套。\nP270 操作时不得进食。"],
    ]


def test_section2_english_classification_uses_the_same_line_policy():
    rows = normalize_s2_projected_rows([
        ["2.2 GHS classification", "Flammable liquid 3 H226; Skin corrosion 1B H314"],
        ["2.3 GHS label elements", "Please note the following substance:\nN,N-dimethylethanolamine\nspecific concentration limit >=5%"],
    ])
    assert rows == [
        ["2.2 GHS classification", "Flammable liquid 3 H226\nSkin corrosion 1B H314"],
        ["2.3 GHS label elements", "Please note the following substance:\nN,N-dimethylethanolamine, specific concentration limit >=5%"],
    ]


def test_section9_matches_properties_by_label_not_source_position():
    template = Document(str(ROOT / "examples" / "template_reference.docx"))
    rows = align_s9_rows([
        {"label": "9.3 pH值：", "value": "7-9"},
        {"label": "外 观：", "value": "透明装液体"},
    ], template.tables[8])
    assert rows[0][1] == "透明装液体"
    assert next(value for label, value in rows if "ph" in label.casefold()) == "7-9"


def test_section13_combines_two_notes_into_one_note_slot():
    template = Document(str(ROOT / "examples" / "template_reference.docx"))
    rows = align_note_section_rows([
        ["第一条说明", ""],
        ["第二条说明", ""],
        ["处理方法：", "处置值"],
    ], template.tables[12])
    assert rows == [["第一条说明\n第二条说明"], ["处理方法：", "处置值"]]
