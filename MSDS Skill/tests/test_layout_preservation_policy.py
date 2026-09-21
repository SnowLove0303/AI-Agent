from pathlib import Path
import sys

from docx import Document
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "scripts"))

from layout_preservation_policy import (  # noqa: E402
    audit_s114_vertical_alignment,
    audit_s82,
)
from template_mutation_whitelist import unique_cells, write_s82_top_rows  # noqa: E402


TEMPLATE = ROOT / "examples" / "template_reference.docx"


def test_s82_hidden_block_is_complete_and_allowed():
    template = Document(str(TEMPLATE))
    output = Document(str(TEMPLATE))
    write_s82_top_rows(output.tables[7], [], "zh")
    report = audit_s82(template, output, language="cn", expected_present=False)
    assert report["status"] == "passed", report["errors"]
    assert report["state"] == "hidden"


def test_s82_data_block_inherits_five_column_grid_and_four_logical_cells():
    template = Document(str(TEMPLATE))
    output = Document(str(TEMPLATE))
    write_s82_top_rows(
        output.tables[7],
        [["物质A", "GB/T 1234", "TWA", "1 mg/m3"]],
        "zh",
    )
    report = audit_s82(template, output, language="cn", expected_present=True)
    assert report["status"] == "passed", report["errors"]
    assert report["grid_column_count"] == 5
    assert report["data_row_count"] == 1


def test_s82_formal_template_is_self_consistent():
    template = Document(str(TEMPLATE))
    report = audit_s82(template, template, language="cn")
    assert report["status"] == "passed", report["errors"]


def test_s82_grid_drift_is_blocked():
    template = Document(str(TEMPLATE))
    output = Document(str(TEMPLATE))
    write_s82_top_rows(
        output.tables[7],
        [["物质A", "GB/T 1234", "TWA", "1 mg/m3"]],
        "zh",
    )
    output.tables[7]._tbl.tblGrid.gridCol_lst[1].set(qn("w:w"), "999")
    report = audit_s82(template, output, language="cn", expected_present=True)
    assert report["status"] == "failed"
    assert any("table grid changed" in error for error in report["errors"])


def test_s82_basis_gridspan_drift_is_blocked():
    template = Document(str(TEMPLATE))
    output = Document(str(TEMPLATE))
    write_s82_top_rows(
        output.tables[7],
        [["物质A", "GB/T 1234", "TWA", "1 mg/m3"]],
        "zh",
    )
    basis = unique_cells(output.tables[7].rows[14])[1]
    basis._tc.tcPr.find(qn("w:gridSpan")).set(qn("w:val"), "1")
    report = audit_s82(template, output, language="cn", expected_present=True)
    assert report["status"] == "failed"
    assert any("gridSpan=2" in error for error in report["errors"])


def test_s114_exact_vertical_alignment_inherits_template():
    template = Document(str(TEMPLATE))
    output = Document(str(TEMPLATE))
    report = audit_s114_vertical_alignment(template, output)
    assert report["status"] == "passed", report["errors"]
    assert report["state"] == "present"


def test_s114_vertical_alignment_drift_is_blocked():
    template = Document(str(TEMPLATE))
    output = Document(str(TEMPLATE))
    value_cell = unique_cells(output.tables[10].rows[9])[1]
    tc_pr = value_cell._tc.get_or_add_tcPr()
    align = tc_pr.find(qn("w:vAlign"))
    if align is None:
        align = tc_pr.makeelement(qn("w:vAlign"), {})
        tc_pr.append(align)
    align.set(qn("w:val"), "bottom")
    report = audit_s114_vertical_alignment(template, output)
    assert report["status"] == "failed"
    assert any("vertical alignment changed" in error for error in report["errors"])
