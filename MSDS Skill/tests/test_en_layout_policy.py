from pathlib import Path
import sys

from docx import Document

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "scripts"))
from normalize_en_layout import normalize_en_document


def test_en_layout_policy_removes_duplicate_first_heading_and_wrap_risk(tmp_path):
    source = ROOT / "examples" / "template_reference_en.docx"
    output = tmp_path / "en.docx"
    doc = Document(str(source))
    first_heading = doc.tables[0].cell(0, 0).paragraphs[0]
    first_heading.text = "Identification"
    section11 = doc.tables[10]
    for row_index, text in {
        3: "Oral",
        4: "Inhalation",
        5: "Dermal",
        12: "Fertility",
        13: "Developmental toxicity",
        14: "In-vitro genetic toxicity",
    }.items():
        section11.rows[row_index].cells[1].text = text
    normalize_en_document(doc)
    doc.save(output)

    checked = Document(str(output))
    assert checked.tables[0].cell(0, 0).text == "Identification"
    assert checked.tables[10].cell(4, 1).text == "Inhalation"
    assert checked.tables[10].cell(5, 1).text == "Dermal"
    assert checked.tables[10].cell(13, 1).text == "Developmental toxicity"
    assert checked.tables[10].cell(14, 1).text == "In-vitro genetic toxicity"
    assert checked.tables[12].cell(2, 1).paragraphs[0].alignment == 0
