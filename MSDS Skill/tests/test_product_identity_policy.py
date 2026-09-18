import importlib.util
import sys
from pathlib import Path

mod_path = Path(__file__).parents[1] / 'scripts' / 'product_identity_policy.py'
spec = importlib.util.spec_from_file_location('identity', mod_path)
identity = importlib.util.module_from_spec(spec); sys.modules['identity']=identity; spec.loader.exec_module(identity)

def test_ep1704_expected_identity():
    x = identity.expected_identity('水性环氧乳液', 'EP-1704')
    assert x.product_name_value == ''
    assert x.chinese_name_value == '水性环氧乳液 EP-1704'

def test_english_product_name_requires_reviewed_name_and_adds_model_suffix():
    assert identity.expected_english_product_name(
        'Waterborne polyurethane brightener resin', 'PU-1004'
    ) == 'Waterborne polyurethane brightener resin PU-1004'

def test_english_product_name_does_not_duplicate_model_suffix():
    assert identity.expected_english_product_name(
        'Waterborne polyurethane brightener resin PU-1004', 'PU-1004'
    ) == 'Waterborne polyurethane brightener resin PU-1004'

def test_english_product_name_rejects_blank_model_only_and_chinese_values():
    assert identity.english_product_name_errors('', 'PU-1004')
    assert identity.english_product_name_errors('PU-1004', 'PU-1004')
    assert identity.english_product_name_errors('水性聚氨酯树脂', 'PU-1004')

def test_english_identity_audit_requires_canonical_value():
    assert identity.audit_english_identity(
        product_name_value='Waterborne polyurethane brightener resin PU-1004',
        model='PU-1004',
    ) == []

def test_rendered_english_identity_gate_rejects_blank_and_accepts_reviewed_value():
    sys.path.insert(0, str(Path(__file__).parents[1] / 'scripts'))
    from msds_pipeline import gate_product_identity
    from docx import Document
    template = Path(__file__).parents[1] / 'examples' / 'template_reference_en.docx'
    blank = Document(str(template))
    assert gate_product_identity(
        template, 'en', 'PU-1004', document=blank
    )
    reviewed = Document(str(template))
    reviewed.tables[0].rows[1].cells[1].text = (
        'Waterborne polyurethane brightener resin PU-1004'
    )
    assert gate_product_identity(
        template, 'en', 'PU-1004', document=reviewed
    ) == []

def test_reject_model_in_product_name():
    errs = identity.audit_identity(chinese_name='水性环氧乳液', model='EP-1704',
        product_name_value='EP-1704', chinese_name_value='水性环氧乳液 EP-1704',
        header_text='EP-1704', footer_text='EP-1704-MSDS')
    assert any('产品名称' in e for e in errs)

def test_no_duplicate_model():
    x = identity.expected_identity('水性环氧乳液 EP-1704', 'EP-1704')
    assert x.chinese_name_value == '水性环氧乳液 EP-1704'

def test_no_duplicate_model_when_source_name_has_compact_suffix():
    x = identity.expected_identity('脂肪族水性聚氨酯接着树脂PU-1001', 'PU-1001')
    assert x.chinese_name_value == '脂肪族水性聚氨酯接着树脂PU-1001'

def test_compact_model_suffix_match_is_case_insensitive_and_whitespace_tolerant():
    x = identity.expected_identity('水性环氧乳液 p u - 1 7 0 4', 'PU-1704')
    assert x.chinese_name_value == '水性环氧乳液 p u - 1 7 0 4'
