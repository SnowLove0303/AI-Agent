from pathlib import Path

from docx import Document

ROOT = Path(__file__).resolve().parents[1]
import sys
sys.path.insert(0, str(ROOT / "scripts"))

from agent_execution_contract import (  # noqa: E402
    load_spec,
    validate_agent_execution_contract,
    validate_spec_installation,
)
from audit_openspec_overwrite import (  # noqa: E402
    audit_empty_value_rows,
    audit_value_whitespace,
)
from missing_data_policy import apply_source_absence_policy  # noqa: E402
from template_mutation_whitelist import unique_cells  # noqa: E402


TEMPLATE = ROOT / "examples" / "template_reference.docx"


def reviewed_execution():
    spec = load_spec()
    return {
        "spec_id": spec["spec_id"],
        "spec_version": spec["version"],
        "status": "reviewed",
        "read_mode": "full",
        "read_sources": spec["read_before_action"],
        "acknowledged": {key: True for key in spec["required_acknowledgements"]},
        "operation_order": spec["operation_order"],
        "execution_mode": spec["execution_mode"],
        "agent_mutation_boundary": spec["agent_mutation_boundary"],
    }


def reviewed_sop():
    sop = load_spec()["sop"]
    return {
        "version": sop["version"],
        "status": "reviewed",
        "stages": [{"stage": stage, "status": "completed"} for stage in sop["required_stages"]],
        "loaded_local_sections": sop["required_local_sections"],
        "cross_section_routes": [],
        "audit_plan": ["source_fidelity", "template_lock", "section_local_rules", "render_qa"],
    }


def test_openspec_normative_sources_are_installed():
    assert validate_spec_installation() == []


def test_agent_execution_requires_full_review_record():
    errors = validate_agent_execution_contract({})
    assert any("agent_execution contract is missing" in error for error in errors)


def test_reviewed_agent_execution_record_is_accepted():
    assert validate_agent_execution_contract({
        "agent_execution": reviewed_execution(), "overwrite_sop": reviewed_sop()
    }) == []


def test_agent_execution_requires_reviewed_overwrite_sop():
    errors = validate_agent_execution_contract({"agent_execution": reviewed_execution()})
    assert any("overwrite_sop record is missing" in error for error in errors)


def test_sop_keeps_cross_section_routing_explicit():
    record = reviewed_sop()
    record["cross_section_routes"] = [{
        "source_section": "s3",
        "target_section": "s2",
        "source_fact_ids": ["FACT-S3-001"],
        "reason": "GHS label explanation is semantically a Section 2 label element",
    }]
    assert record["cross_section_routes"][0]["source_section"] != record["cross_section_routes"][0]["target_section"]


def test_agent_mutation_boundary_cannot_be_expanded():
    record = reviewed_execution()
    record["agent_mutation_boundary"] = dict(record["agent_mutation_boundary"])
    record["agent_mutation_boundary"]["allowed_operations"] = [
        *record["agent_mutation_boundary"]["allowed_operations"],
        "edit_locked_label_text",
    ]
    errors = validate_agent_execution_contract({"agent_execution": record})
    assert any("agent_mutation_boundary" in error for error in errors)


def test_pipeline_blocks_before_template_clone_without_execution_record():
    from hashlib import sha256
    from msds_pipeline import ReleaseBlocked, validate_approved_facts

    source = ROOT / "examples" / "regression_HPU-7660_source.docx"
    facts = {
        "model": "HPU-7660",
        "source_sha256": sha256(source.read_bytes()).hexdigest(),
        "source_mapping": {},
    }
    try:
        validate_approved_facts(facts, source, "HPU-7660")
    except ReleaseBlocked as exc:
        assert "agent_execution" in str(exc)
    else:
        raise AssertionError("missing execution contract was not blocked")


def test_empty_value_row_is_a_release_blocker():
    document = Document(str(TEMPLATE))
    unique_cells(document.tables[9].rows[1])[1].text = ""
    problems = audit_empty_value_rows(document)
    assert any(item["type"] == "empty_value_row_remains" for item in problems)


def test_missing_phone_value_row_is_hidden_and_following_row_shifts_up():
    document = Document(str(TEMPLATE))
    phone_row = document.tables[0].rows[8]
    fax_row = document.tables[0].rows[9]
    phone_label = unique_cells(phone_row)[0].text
    fax_label = unique_cells(fax_row)[0].text
    unique_cells(phone_row)[1].text = ""
    facts = {"s1": [["source", "value"] for _ in range(9)]}

    result = apply_source_absence_policy(document, facts, unique_cells)

    assert any(phone_label in item for item in result["source_absent_removed"])
    labels = [unique_cells(row)[0].text for row in document.tables[0].rows]
    assert phone_label not in labels
    assert labels[-1] == fax_label


def test_intentional_blank_slots_are_not_classified_as_empty_value_rows():
    document = Document(str(TEMPLATE))
    problems = audit_empty_value_rows(document)
    assert not any(
        item["type"] == "empty_value_row_remains"
        and item["table"] == 0
        and item["row"] in {1, 5}
        for item in problems
    )


def test_extra_empty_component_rows_are_suppressed():
    document = Document(str(TEMPLATE))
    unique_cells(document.tables[2].rows[5])[0].text = ""
    unique_cells(document.tables[2].rows[5])[1].text = ""
    unique_cells(document.tables[2].rows[5])[2].text = ""
    facts = {"s3": [
        ["产品类型：", "混合物", ""],
        ["成分", "", ""],
        ["化学品名称", "CAS编号", "含量%（w/w）"],
        ["组分A", "123-45-6", "10"],
    ]}
    result = apply_source_absence_policy(document, facts, unique_cells)
    assert result["source_absent_removed"]
    assert len(document.tables[2].rows) == 5


def test_blank_lines_and_fake_spacing_are_release_blockers():
    document = Document(str(TEMPLATE))
    cell = unique_cells(document.tables[9].rows[1])[1]
    cell.paragraphs[0].text = "first\n\nsecond   "
    problems = audit_value_whitespace(document)
    assert any(item["type"] == "blank_line_in_value" for item in problems)
    assert any(item["type"] == "artificial_spacing_in_value" for item in problems)
