from pathlib import Path
import sys

import pytest
from docx import Document

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "scripts"))

from extract_source_facts import Extraction, extract_s2, extract_s8  # noqa: E402
from msds_pipeline import write_body  # noqa: E402
from section2_fact_router import ROUTER_VERSION, audit, semantic_target  # noqa: E402
from section2_ghs_policy import project_source_cn_facts  # noqa: E402
from section_overwrite_rules import sanitize_section_payload  # noqa: E402
from s8_ppe_policy import (  # noqa: E402
    align_s8_rows,
    audit_s8_ppe_label_boundaries,
)
from template_mutation_whitelist import (  # noqa: E402
    MutationViolation,
    TemplateSlotRegistry,
    clear_value_cells,
    compare_locked_skeleton,
    english_body_cells,
    is_s15_locked_heading_row,
    unique_cells,
    write_row_values,
)


def _add_two_cell_row(table, label, value=""):
    row = table.add_row()
    row.cells[0].text = label
    row.cells[1].text = value
    return row


def test_s2_extracts_health_routes_without_reusing_overall_classification():
    document = Document()
    table = document.add_table(rows=1, cols=1)
    overall = "\u672a\u88ab\u5206\u7c7b"
    ingestion = "\u6b63\u5e38\u4f7f\u7528\u65f6\u53ea\u6709\u8f7b\u5fae\u7684\u6444\u5165\u5371\u5bb3"
    skin = "\u53ef\u80fd\u5f15\u8d77\u8f7b\u5fae\u76ae\u80a4\u523a\u6fc0"
    eyes = "\u53ef\u80fd\u5f15\u8d77\u773c\u775b\u523a\u6fc0\u3001\u53d1\u7ea2\u3001\u6d41\u6cea"
    for line in (
        f"2.1 GHS\u5371\u9669\u6027\u5206\u7c7b\uff1a{overall}",
        f"\u5371\u9669\u6027\uff1a{overall}",
        f"\u98df\u5165\u5371\u5bb3\uff1a{ingestion}",
        f"\u76ae\u80a4\u523a\u6fc0\uff1a{skin}",
        f"\u773c\u775b\u523a\u6fc0\uff1a{eyes}",
    ):
        table.add_row().cells[0].text = line

    data = extract_s2(table, Extraction())

    assert overall in data["ghs_classes"]
    assert data["health_hazards"]["inhalation"] == []
    assert data["health_hazards"]["ingestion"] == [ingestion]
    assert data["health_hazards"]["skin"] == [skin]
    assert data["health_hazards"]["eyes"] == [eyes]

    rows, _ = project_source_cn_facts(data)
    projected = {
        semantic_target(label): value
        for label, value in rows
        if "route=" in label
    }
    assert projected["health_hazards.inhalation"] == ""
    assert projected["health_hazards.ingestion"] == ingestion
    assert projected["health_hazards.skin"] == skin
    assert projected["health_hazards.eyes"] == eyes


def test_section2_route_alias_wins_over_visible_2_8_number():
    assert semantic_target("2.8 Health hazards [route=skin]") == "health_hazards.skin"
    assert semantic_target("s2.health_hazards.ingestion[1]") == "health_hazards.ingestion"
    assert semantic_target("2.8 \u5065\u5eb7\u5371\u5bb3 [route=eyes]") == "health_hazards.eyes"


def test_s2_pipeline_writes_sparse_routes_before_omission_and_renumbering():
    document = Document(str(ROOT / "examples" / "template_reference.docx"))
    facts = {f"s{section}": [] for section in range(1, 17)}
    facts["s2"] = project_source_cn_facts({
        "ghs_classes": ["\u672a\u88ab\u5206\u7c7b"],
        "label_elements": [],
        "label_ingredients": [],
        "signal": "\u65e0\u4fe1\u53f7\u8bcd",
        "h_statements": [],
        "p_statements": [],
        "health_hazards": {
            "inhalation": [],
            "ingestion": ["\u98df\u5165\u5371\u5bb3"],
            "skin": ["\u76ae\u80a4\u5371\u5bb3"],
            "eyes": [],
            "symptoms_signs": [],
        },
        "other_hazards": "",
    })[0]

    write_body(document, facts, "zh")
    rows = [unique_cells(row) for row in document.tables[1].rows[1:]]
    assert any("GHS\u5371\u9669\u6027\u7c7b\u522b" in cells[0].text
               and cells[1].text == "\u672a\u88ab\u5206\u7c7b" for cells in rows)
    assert any("\u4fe1\u53f7\u8bcd" in cells[0].text
               and cells[1].text == "\u65e0\u4fe1\u53f7\u8bcd" for cells in rows)
    route_rows = [cells for cells in rows if "\u5065\u5eb7\u5371\u5bb3" in cells[0].text]
    assert len(route_rows) == 2
    assert {cells[1].text for cells in route_rows} == {
        "\u98df\u5165\uff1a\n\u98df\u5165\u5371\u5bb3",
        "\u76ae\u80a4\uff1a\n\u76ae\u80a4\u5371\u5bb3",
    }
    assert all("\u672a\u88ab\u5206\u7c7b" not in cells[1].text for cells in route_rows)
    assert all("\u773c\u775b\uff1a" not in cells[1].text
               and "\u75c7\u72b6\u548c\u4f53\u5f81\uff1a" not in cells[1].text
               for cells in rows)
    assert compare_locked_skeleton(
        Document(str(ROOT / "examples" / "template_reference.docx")),
        document,
    ) == []


def test_section2_rejects_overall_not_classified_fact_as_route_value():
    value = "\u672a\u88ab\u5206\u7c7b"
    facts = {
        "fact_ledger": [{
            "fact_id": "FACT-CLASS",
            "source_section": "s2",
            "source_locator": "s2.ghs_classes[1]",
            "source_text": value,
            "normalized_value": value,
            "evidence_type": "explicit",
        }],
        "source_mapping": {"items": [{
            "fact_id": "FACT-CLASS",
            "source_section": "s2",
            "source_locator": "s2.ghs_classes[1]",
            "source_text": value,
            "decision": "mapped",
            "target_section": "s2",
            "target_slot": "s2.health_hazards.inhalation[1]",
        }]},
        "section2_routing": {
            "version": ROUTER_VERSION,
            "status": "reviewed",
            "items": [{
                "fact_id": "FACT-CLASS",
                "semantic_target": "health_hazards.inhalation",
                "usage": "exclusive",
            }],
        },
        "output_traceability": {"items": [{
            "target_section": "s2",
            "target_slot": "s2.health_hazards.inhalation",
            "decision": "written",
            "source_fact_ids": ["FACT-CLASS"],
            "output_values": {"zh": value, "en": "Not classified"},
        }]},
        "zh": {"s2": [["2.8 Health hazards [route=inhalation]", value]]},
        "en": {"s2": [["2.8 Health hazards [route=inhalation]", "Not classified"]]},
    }
    report = audit(facts)
    assert report["status"] == "failed"
    assert any("overall classification fact FACT-CLASS" in error for error in report["errors"])


def test_irregular_section8_rows_are_recovered_and_aligned_by_label():
    document = Document()
    table = document.add_table(rows=1, cols=2)
    _add_two_cell_row(table, "8.1 \u66b4\u9732\u63a7\u5236\uff1a")
    _add_two_cell_row(table, "\u547c\u5438\u7cfb\u7edf\u9632\u62a4\uff1a", "\u55b7\u6d82\u8fc7\u7a0b\u4e2d\u8981\u6c42\u6709\u547c\u5438\u9632\u62a4\u8bbe\u5907\u3002")
    _add_two_cell_row(table, "\u624b\u90e8\u9632\u62a4\uff1a\t\u55b7\u6d82\u8fc7\u7a0b\u4e2d\u8981\u6c42\u6709\u547c\u5438\u9632\u62a4\u8bbe\u5907\u3002", "\u5efa\u8bae\u6234\u4e0a\u9632\u62a4\u624b\u5957\u3002")
    _add_two_cell_row(table, "\u9632\u62a4\u624b\u5957\u7684\u5408\u9002\u6750\u6599\uff1a", "EN 374-3")
    _add_two_cell_row(table, "\u6c1f\u5316\u6a61\u80f6 \u2013 FKM:", "\u539a\u5ea6\u22650.4mm\uff1b\u7a7f\u900f\u65f6\u95f4\u2265480min.")
    _add_two_cell_row(table, "\u4e01\u57fa\u6a61\u80f6 \u2013 IIR:", "\u539a\u5ea6\u22650.5mm\uff1b\u7a7f\u900f\u65f6\u95f4\u2265480min.")
    _add_two_cell_row(table, "\u4e01\u8148\u6a61\u80f6 \u2013 NBR:", "\u539a\u5ea6\u22650.35mm\uff1b\u7a7f\u900f\u65f6\u95f4\u2265480min.")
    _add_two_cell_row(table, "\u5efa\u8bae\uff1a", "\u6c61\u67d3\u7684\u624b\u5957\u5e94\u5e9f\u5f03\u3002")
    _add_two_cell_row(table, "\u773c\u775b\u9632\u62a4\uff1a", "\u6234\u62a4\u76ee\u955c/\u9762\u7f69\u3002")
    _add_two_cell_row(table, "\u8eab\u4f53\u9632\u62a4\uff1a", "\u7a7f\u7740\u9002\u5f53\u7684\u9632\u62a4\u670d\u3002")

    ext = Extraction()
    rows, controls = extract_s8(table, ext)
    labels = [row[0] for row in rows[1:-1]]
    aligned = {row[0]: row[1] for row in align_s8_rows(rows)[1:-1]}

    assert labels == [
        "\u547c\u5438\u7cfb\u7edf\u9632\u62a4\uff1a",
        "\u624b\u90e8\u9632\u62a4\uff1a",
        "\u9632\u62a4\u624b\u5957\u7684\u5408\u9002\u6750\u6599\uff1a",
        "\u6c1f\u5316\u6a61\u80f6 \u2013 FKM:",
        "\u4e01\u57fa\u6a61\u80f6 \u2013 IIR:",
        "\u4e01\u8148\u6a61\u80f6 \u2013 NBR:",
        "\u5efa\u8bae\uff1a",
        "\u773c\u775b\u9632\u62a4\uff1a",
        "\u8eab\u4f53\u9632\u62a4\uff1a",
    ]
    assert aligned["\u547c\u5438\u7cfb\u7edf\u9632\u62a4\uff1a"].startswith("\u55b7\u6d82")
    assert aligned["\u624b\u90e8\u9632\u62a4\uff1a"].startswith("\u5efa\u8bae\u6234")
    assert "\u55b7\u6d82\u8fc7\u7a0b\u4e2d\u8981\u6c42\u6709\u547c\u5438\u9632\u62a4" not in aligned["\u624b\u90e8\u9632\u62a4\uff1a"]
    assert aligned["\u6c1f\u5316\u6a61\u80f6 \u2013 FKM:"].startswith("\u539a\u5ea6")
    assert aligned["\u773c\u775b\u9632\u62a4\uff1a"].startswith("\u6234\u62a4")
    assert aligned["\u8eab\u4f53\u9632\u62a4\uff1a"].startswith("\u7a7f\u7740")
    assert controls == []
    assert any(item["issue"] == "inline-label-value-contamination" for item in ext.review)


def test_section8_sanitizer_keeps_sparse_slots_until_semantic_alignment():
    rows = [["8.1 \u66b4\u9732\u63a7\u5236\uff1a", ""], ["\u624b\u90e8\u9632\u62a4\uff1a", ""]]
    assert sanitize_section_payload(8, rows) == rows


def test_section8_source_boundary_is_reviewed_and_output_label_uses_baseline_diff():
    source_errors = audit_s8_ppe_label_boundaries([
        ["手部防护：\t源文件中的旧尾部", "建议戴上防护手套。"],
    ])
    assert any("S8 source PPE row" in error for error in source_errors)
    assert audit_s8_ppe_label_boundaries([["手部防护：", "建议戴上防护手套。"]]) == []

    template = Document(str(ROOT / "examples" / "template_reference.docx"))
    output = Document(str(ROOT / "examples" / "template_reference.docx"))
    assert compare_locked_skeleton(template, output) == []
    label_cell = unique_cells(output.tables[7].rows[3])[0]
    label_cell.paragraphs[0].add_run("\t不应写入标签单元格的值")
    errors = compare_locked_skeleton(template, output)
    assert any("locked" in error and "text changed" in error for error in errors)


def test_section2_sanitizer_keeps_sparse_slots_until_semantic_omission():
    rows = [["2.1 \u7d27\u6025\u60c5\u51b5\u6982\u8ff0", ""],
            ["2.2 GHS\u5371\u9669\u6027\u7c7b\u522b\uff1a", "\u672a\u88ab\u5206\u7c7b"]]
    assert sanitize_section_payload(2, rows) == rows


@pytest.mark.parametrize(
    ("filename", "heading1", "heading2"),
    [
        ("template_reference.docx", "\u5176\u5b83\u7684\u89c4\u5b9a\uff1a", "\u7b26\u5408\u4e0b\u5217\u6cd5\u89c4\u8981\u6c42\uff1a"),
        ("template_reference_en.docx", "Other provisions: ", "Complies with the following regulations: "),
    ],
)
def test_section15_bold_headings_are_locked_not_one_cell_values(filename, heading1, heading2):
    template = Document(str(ROOT / "examples" / filename))
    output = Document(str(ROOT / "examples" / filename))
    table = output.tables[14]
    assert table.rows[2].cells[0].text == heading1
    assert table.rows[3].cells[0].text == heading2
    assert is_s15_locked_heading_row(table.rows[2])
    assert is_s15_locked_heading_row(table.rows[3])
    assert all(run.bold is True for run in table.rows[2].cells[0].paragraphs[0].runs if run.text.strip())
    assert all(run.bold is True for run in table.rows[3].cells[0].paragraphs[0].runs if run.text.strip())

    registry = TemplateSlotRegistry.from_document(output)
    assert (14, 2) not in registry.slots
    assert (14, 3) not in registry.slots
    assert english_body_cells(14, 2, table.rows[2]) == []
    assert english_body_cells(14, 3, table.rows[3]) == []
    clear_value_cells(output, registry)
    with pytest.raises(MutationViolation):
        write_row_values(
            table.rows[2], [heading1, "must not overwrite heading"],
            table_index=14, row_index=2, registry=registry,
        )
    with pytest.raises(MutationViolation):
        write_row_values(
            table.rows[3], [heading2, "must not overwrite heading"],
            table_index=14, row_index=3, registry=registry,
        )
    assert compare_locked_skeleton(template, output) == []
