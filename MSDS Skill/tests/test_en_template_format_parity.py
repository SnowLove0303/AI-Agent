from copy import deepcopy
from pathlib import Path
import sys

from docx import Document
from lxml import etree

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "scripts"))
from normalize_en_layout import EN_BODY_SIZE, _approved_body_rpr, sync_en_template_text_format
from template_mutation_whitelist import (
    TemplateSlotRegistry,
    clear_value_cells,
    compare_format_anchors,
    unique_cells,
    write_row_values,
)


def _xml(element):
    if element is None:
        return ()

    def node(item):
        return (
            etree.QName(item).localname,
            tuple(sorted((etree.QName(key).localname, value)
                         for key, value in item.attrib.items())),
            tuple(node(child) for child in item),
        )

    return node(element)


def test_en_text_format_is_restored_from_the_active_en_template():
    template = ROOT / "examples" / "template_reference_en.docx"
    document = Document(str(template))
    target = document.tables[0].cell(1, 1).paragraphs[0]
    target.text = "OS-9015"
    target.paragraph_format.alignment = 3
    target.runs[0].font.name = "Courier New"
    target.runs[0].font.size = None

    reference = Document(str(template))
    ref = reference.tables[0].cell(1, 1).paragraphs[0]
    sync_en_template_text_format(document, template)

    assert _xml(target._p.pPr) == _xml(ref._p.pPr)
    body_rpr = _approved_body_rpr(reference)
    assert _xml(target.runs[0]._r.rPr) == _xml(body_rpr)


def test_en_blank_value_slots_use_the_same_body_format_as_populated_slots():
    template = ROOT / "examples" / "template_reference_en.docx"
    document = Document(str(template))
    populated = document.tables[10].cell(10, 1).paragraphs[0]
    populated.text = "Polyurethane dispersion"
    blank_slot = document.tables[10].cell(12, 3).paragraphs[0]
    blank_slot.text = "Polyurethane polymer\nNo data available."

    reference = Document(str(template))
    sync_en_template_text_format(document, template, reference_document=reference)

    expected = _xml(_approved_body_rpr(reference))
    assert _xml(populated.runs[0]._r.rPr) == expected
    assert _xml(blank_slot.runs[0]._r.rPr) == expected


def test_en_s28_normalization_preserves_prefix_and_normalizes_value_tail():
    template = ROOT / "examples" / "template_reference_en.docx"
    document = Document(str(template))
    registry = TemplateSlotRegistry.from_document(document)
    row = document.tables[1].rows[9]
    prefix_run = unique_cells(row)[1].paragraphs[0].runs[0]
    prefix_rpr = _xml(prefix_run._r.rPr)
    write_row_values(
        row, [unique_cells(row)[0].text, "source-backed health hazard"],
        table_index=1, row_index=9, registry=registry,
    )
    sync_en_template_text_format(document, template)
    cell = unique_cells(row)[1]
    assert cell.text == f"{prefix_run.text}\nsource-backed health hazard"
    assert _xml(cell.paragraphs[0].runs[0]._r.rPr) == prefix_rpr
    assert _xml(cell.paragraphs[0].runs[-1]._r.rPr) == _xml(_approved_body_rpr(Document(str(template))))
    assert compare_format_anchors(Document(str(template)), document, language="en") == []


def test_en_compatibility_layout_constant_is_not_the_legacy_10_5_pt_style():
    assert EN_BODY_SIZE.pt == 12
