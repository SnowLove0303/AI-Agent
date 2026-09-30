from __future__ import annotations

import io
import sys
import struct
import zlib
import binascii
from pathlib import Path

from docx import Document
from docx.shared import Inches

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "scripts"))

from audit_openspec_overwrite import audit_empty_value_rows, audit_value_whitespace
from ghs_pictogram_policy import extract_all_embedded_images, insert_source_pictogram


TEMPLATE = ROOT / "examples" / "template_reference.docx"


def _png_bytes(rgb: tuple[int, int, int]) -> bytes:
    """Create a valid 1x1 RGB PNG without an optional imaging dependency."""
    raw = b"\x00" + bytes(rgb)

    def chunk(kind: bytes, payload: bytes) -> bytes:
        checksum = binascii.crc32(kind + payload) & 0xFFFFFFFF
        return (struct.pack(">I", len(payload)) + kind + payload
                + struct.pack(">I", checksum))

    header = struct.pack(">IIBBBBB", 1, 1, 8, 2, 0, 0, 0)
    return (b"\x89PNG\r\n\x1a\n" + chunk(b"IHDR", header)
            + chunk(b"IDAT", zlib.compress(raw)) + chunk(b"IEND", b""))


def _create_multi_image_docx(path: Path) -> Path:
    doc = Document()
    for color in ((255, 0, 0), (0, 0, 0)):
        doc.add_picture(io.BytesIO(_png_bytes(color)), width=Inches(0.5))
    doc.save(str(path))
    return path


def test_extract_all_embedded_images_reads_media_parts(tmp_path):
    multi_doc_path = _create_multi_image_docx(tmp_path / "multi_image.docx")
    images = extract_all_embedded_images(multi_doc_path)
    assert isinstance(images, list)
    assert len(images) == 2
    for name, payload in images:
        assert name.endswith((".png", ".jpeg", ".jpg", ".emf"))
        assert len(payload) > 0


def test_insert_source_pictogram_embeds_in_single_run_without_spacers(tmp_path):
    multi_doc_path = _create_multi_image_docx(tmp_path / "multi_image.docx")
    doc = Document(str(TEMPLATE))
    report = insert_source_pictogram(doc, multi_doc_path)
    assert report["image_count"] == 2
    assert report["target_table"] == 1
    assert report["target_row"] == 4

    cell = doc.tables[1].rows[4].cells[-1]
    # Verify exactly one paragraph and exactly one run with no blank spacers
    assert len(cell.paragraphs) == 1
    p = cell.paragraphs[0]
    assert len(p.runs) == 1
    run = p.runs[0]
    # Check that drawings are inside this single run
    drawings = run._r.xpath(".//w:drawing")
    assert len(drawings) == 2

    # Audit empty value rows should pass on this cell
    empty_issues = audit_empty_value_rows(doc)
    assert not any(issue.get("table") == 1 and issue.get("row") == 4 for issue in empty_issues)

    # Whitespace audit should have no fake spacing in this cell
    ws_issues = audit_value_whitespace(doc)
    assert not any(issue.get("table") == 1 and issue.get("row") == 4 for issue in ws_issues)
