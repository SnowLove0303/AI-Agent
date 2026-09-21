from pathlib import Path

import pytest

import sys
sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "scripts"))

from source_ingest import (  # noqa: E402
    SourceSelectionError,
    discover_source,
    prepare_source,
    require_section_extraction,
)


def test_pdf_is_publication_only_and_cannot_be_a_source(tmp_path):
    source = tmp_path / "PU-2411.pdf"
    source.write_bytes(b"source")
    with pytest.raises(SourceSelectionError, match="publication-only"):
        discover_source(source)


def test_directory_discovery_requires_one_unambiguous_candidate(tmp_path):
    (tmp_path / "~$PU-2411.docx").write_bytes(b"lock")
    source = tmp_path / "PU-2411.docx"
    source.write_bytes(b"docx")
    selected = discover_source(tmp_path, model="PU-2411")
    assert selected.original_path == source.resolve()

    (tmp_path / "PU-2411.xlsx").write_bytes(b"spreadsheet")
    with pytest.raises(SourceSelectionError, match="ambiguous"):
        discover_source(tmp_path, model="PU-2411")


def test_direct_docx_source_is_the_only_extraction_path_without_conversion(tmp_path):
    source = tmp_path / "source.docx"
    source.write_bytes(b"docx")
    selected = discover_source(source)
    with prepare_source(selected) as prepared:
        assert prepared.adapter == "direct-docx"
        assert prepared.extraction_path == source.resolve()
        assert require_section_extraction(prepared) == source.resolve()


def test_formal_output_cannot_be_reused_as_source(tmp_path):
    output = tmp_path / "PU-2411_MSDS_CN_冠志.docx"
    output.write_bytes(b"formal output")
    with pytest.raises(SourceSelectionError, match="formal MSDS output"):
        discover_source(output)


def test_source_under_overwrite_output_directory_is_allowed(tmp_path):
    output_dir = tmp_path / "覆写产出"
    output_dir.mkdir()
    source = output_dir / "reviewed_source.docx"
    source.write_bytes(b"original source")
    selected = discover_source(source)
    assert selected.original_path == source.resolve()


def test_docx_preview_checkpoint_cannot_be_reused_as_source(tmp_path):
    preview_dir = tmp_path / "_docx_preview"
    preview_dir.mkdir()
    preview = preview_dir / "reviewed_source.docx"
    preview.write_bytes(b"audited DOCX checkpoint")
    with pytest.raises(SourceSelectionError, match="generated/output directory"):
        discover_source(preview)


def test_legacy_word_conversion_is_reused_by_source_hash(tmp_path, monkeypatch):
    import source_ingest
    from docx import Document

    source = tmp_path / "legacy.doc"
    source.write_bytes(b"legacy-source-v1")
    cache = tmp_path / ".msds_cache"
    calls = []

    def fake_run(command, **kwargs):
        calls.append(command)
        out_dir = Path(command[5])
        converted = out_dir / "legacy.docx"
        Document().save(converted)
        return type("Result", (), {"returncode": 0, "stdout": "", "stderr": ""})()

    monkeypatch.setattr(source_ingest, "_find_soffice", lambda: "soffice.com")
    monkeypatch.setattr(source_ingest.subprocess, "run", fake_run)
    selected = discover_source(source)
    with prepare_source(selected, cache_dir=cache) as prepared:
        assert prepared.adapter == "libreoffice-docx-cache"
        assert prepared.extraction_path.is_file()
        assert prepared.cache_reused is False
    with prepare_source(selected, cache_dir=cache) as prepared:
        assert prepared.adapter == "libreoffice-docx-cache"
        assert prepared.extraction_path.is_file()
        assert prepared.cache_reused is True
    assert len(calls) == 1

    source.write_bytes(b"legacy-source-v2")
    changed = discover_source(source)
    with prepare_source(changed, cache_dir=cache) as prepared:
        assert prepared.extraction_path.is_file()
    assert len(calls) == 2


def test_legacy_grounding_uses_prepared_source_but_keeps_original_identity(tmp_path):
    from docx import Document
    from source_grounding import audit

    source = tmp_path / "legacy.doc"
    source.write_bytes(b"legacy-source")
    prepared = tmp_path / "prepared.docx"
    document = Document()
    document.add_paragraph("Known prepared source value")
    document.save(prepared)
    facts = {
        "fact_ledger": [{
            "fact_id": "fact-1",
            "source_text": "Known prepared source value",
            "normalized_value": "Known prepared source value",
        }],
        "zh": {"s1": [["Product name", "Known prepared source value"]]},
    }

    report = audit(facts, source, "EFF-TEST", prepared_source=prepared)

    assert report["status"] == "passed"
    assert report["source"] == str(source)
    assert report["prepared_source"] == str(prepared)
    assert report["source_search_path"] == str(prepared)
