import sys
from pathlib import Path
ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "scripts"))

from copy import deepcopy
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn

from template_mutation_whitelist import (
    audit_cross_page_contract,
    audit_value_typography_contract,
    clear_value_cells,
    composite_value_text,
    MutationViolation,
    compare_format_anchors,
    compare_locked_skeleton,
    set_sequence_prefix,
    TemplateSlotRegistry,
    unique_cells,
    write_s82_top_rows,
    write_row_values,
)
from template_runtime import ensure_source_data_rows, sanitize_template_artifacts


import sys
ROOT = Path(__file__).resolve().parents[1]
TEMPLATE = ROOT / "examples" / "template_reference.docx"
TEMPLATE_EN = ROOT / "examples" / "template_reference_en.docx"
TEMPLATE_EN_SOURCE = ROOT / "examples" / "template_reference_en_source.docx"


def _unique_texts(row):
    seen = set()
    texts = []
    for cell in row.cells:
        key = id(cell._tc)
        if key not in seen:
            seen.add(key)
            texts.append(cell.text)
    return texts


def test_value_write_preserves_locked_label_and_skeleton():
    template = Document(str(TEMPLATE))
    output = Document(str(TEMPLATE))
    row = output.tables[0].rows[1]
    original_label = row.cells[0].text
    write_row_values(row, ["unused source label", "new value"], table_index=0, row_index=1)
    assert row.cells[0].text == original_label
    assert not compare_locked_skeleton(template, output)


def test_global_value_writer_removes_duplicate_row_heading_only():
    output = Document(str(TEMPLATE))
    row = output.tables[3].rows[1]
    write_row_values(
        row,
        ["4.1 一般措施：", "一般措施：立即脱掉所有被污染的衣物。"],
        table_index=3, row_index=1,
    )
    assert unique_cells(row)[1].text == "立即脱掉所有被污染的衣物。"


def test_label_text_change_is_blocked():
    template = Document(str(TEMPLATE))
    output = Document(str(TEMPLATE))
    row = output.tables[0].rows[1]
    paragraph = row.cells[0].paragraphs[0]
    for run in paragraph.runs:
        if run.text.strip():
            run.text = "tampered label"
            break
    errors = compare_locked_skeleton(template, output)
    assert any("locked label text changed" in error for error in errors)


def test_label_format_change_is_blocked():
    template = Document(str(TEMPLATE))
    output = Document(str(TEMPLATE))
    paragraph = output.tables[0].rows[1].cells[0].paragraphs[0]
    paragraph.paragraph_format.left_indent = 999000
    errors = compare_locked_skeleton(template, output)
    assert any("paragraph properties changed" in error for error in errors)


def test_value_format_change_is_blocked():
    template = Document(str(TEMPLATE))
    output = Document(str(TEMPLATE))
    value = output.tables[0].rows[1].cells[1].paragraphs[0]
    value.paragraph_format.space_after = 123
    errors = compare_format_anchors(template, output)
    assert any("format anchor changed" in error for error in errors)


def test_cross_page_contract_preserves_template_table_and_row_settings():
    template = Document(str(TEMPLATE))
    output = Document(str(TEMPLATE))
    report = audit_cross_page_contract(template, output)
    assert report["errors"] == []
    assert len(report["tables"]) == 16
    assert all(item["template_allows_cross_page"] for item in report["tables"])
    assert all(item["output_allows_cross_page"] for item in report["tables"])
    assert all(item["template_row_cant_split"] == item["output_row_cant_split"] for item in report["tables"])
    assert all(item["template_all_rows_breakable"] == item["output_all_rows_breakable"] for item in report["tables"])


def test_cross_page_contract_blocks_row_split_setting_drift():
    template = Document(str(TEMPLATE))
    output = Document(str(TEMPLATE))
    tr_pr = output.tables[0].rows[1]._tr.get_or_add_trPr()
    tr_pr.append(tr_pr.makeelement(qn("w:cantSplit"), {}))
    report = audit_cross_page_contract(template, output)
    assert any("cross-page row settings changed" in error for error in report["errors"])


def test_blank_s117_value_inherits_template_body_run_format():
    template = Document(str(TEMPLATE))
    output = Document(str(TEMPLATE))
    row = output.tables[10].rows[12]
    write_row_values(
        row,
        ["11.7 生殖毒性：", "生育力", "无数据"],
        table_index=10, row_index=12,
    )
    value_run = row.cells[-1].paragraphs[0].runs[0]
    assert value_run._r.rPr is not None
    assert value_run._r.rPr.find(qn("w:sz")).get(qn("w:val")) == "24"
    assert not compare_format_anchors(template, output)


def test_template_label_artifacts_are_preserved_exactly():
    template = Document(str(TEMPLATE))
    output = Document(str(TEMPLATE))
    original_label = template.tables[7].rows[3].cells[0].text
    sanitize_template_artifacts(output)
    assert output.tables[7].rows[3].cells[0].text == original_label
    assert len(output.tables[7].rows) == len(template.tables[7].rows)
    assert not compare_locked_skeleton(template, output)
    assert not compare_format_anchors(template, output)


def test_one_cell_write_does_not_append_an_empty_line():
    output = Document(str(TEMPLATE))
    row = output.tables[14].rows[1]
    write_row_values(row, ["法规文本", ""], table_index=14, row_index=1)
    assert row.cells[0].text == "法规文本"


def test_sequence_renumber_is_content_only():
    template = Document(str(TEMPLATE))
    output = Document(str(TEMPLATE))
    cell = output.tables[1].rows[1].cells[0]
    set_sequence_prefix(cell, 2, 9)
    assert not compare_locked_skeleton(template, output)


def test_s3_component_row_writes_all_three_data_cells():
    output = Document(str(TEMPLATE))
    row = output.tables[2].rows[4]
    write_row_values(row, ["Component", "123-45-6", "10"], table_index=2, row_index=4)
    assert [cell.text for cell in row.cells] == ["Component", "123-45-6", "10"]
    assert all(paragraph.alignment == 1 for cell in row.cells for paragraph in cell.paragraphs)


def test_global_value_typography_contract_covers_s3_and_section11_values():
    output = Document(str(TEMPLATE))
    registry = TemplateSlotRegistry.from_document(output)
    clear_value_cells(output, registry)
    write_row_values(
        output.tables[2].rows[4], ["Component", "123-45-6", "10"],
        table_index=2, row_index=4, registry=registry,
    )
    write_row_values(
        output.tables[10].rows[1], ["该产品无可用的毒理学研究。"],
        table_index=10, row_index=1, registry=registry,
    )
    assert audit_value_typography_contract(output, "cn") == []

    run = output.tables[10].rows[1].cells[-1].paragraphs[0].runs[0]
    run._r.get_or_add_rPr().remove(run._r.rPr.find(qn("w:sz")))
    errors = audit_value_typography_contract(output, "cn")
    assert any("explicit size mismatch" in error for error in errors)


def test_registry_records_locked_composite_and_multi_column_topology_deterministically():
    for template_path in (TEMPLATE, TEMPLATE_EN):
        first = TemplateSlotRegistry.from_document(Document(str(template_path)))
        second = TemplateSlotRegistry.from_document(Document(str(template_path)))
        assert first.slots == second.slots
        if first.slots[(1, 9)].role == "s2_composite_value":
            assert first.slots[(1, 9)].cell_indices == (1,)
            assert first.slots[(1, 9)].locked_cell_indices == (0,)
            assert first.slots[(1, 9)].special_policy == "s2_route_prefix"
        else:
            assert first.slots[(1, 9)].role == "field"
        assert first.slots[(2, 4)].cell_indices == (0, 1, 2)
        if (7, 14) in first.slots:
            assert first.slots[(7, 14)].cell_indices == (0, 1, 2, 3)
            assert first.slots[(7, 14)].special_policy == "four_column_data"


def test_s28_route_prefixes_survive_clear_and_overwrite_in_cn_and_en():
    for template_path in (TEMPLATE, TEMPLATE_EN):
        template = Document(str(template_path))
        output = Document(str(template_path))
        registry = TemplateSlotRegistry.from_document(output)
        if registry.slots[(1, 9)].role != "s2_composite_value":
            clear_value_cells(output, registry)
            for index in range(9, 14):
                row = output.tables[1].rows[index]
                write_row_values(row, [unique_cells(row)[0].text, "源文件危害说明"],
                                 table_index=1, row_index=index, registry=registry)
                assert unique_cells(row)[1].text == "源文件危害说明"
            assert compare_locked_skeleton(template, output) == []
            continue
        expected_prefixes = [
            unique_cells(template.tables[1].rows[index])[1].text
            for index in range(9, 14)
        ]
        clear_value_cells(output, registry)
        assert [unique_cells(output.tables[1].rows[index])[1].text for index in range(9, 14)] == expected_prefixes
        for index in range(9, 14):
            row = output.tables[1].rows[index]
            write_row_values(
                row, [unique_cells(row)[0].text, "源文件危害说明"],
                table_index=1, row_index=index, registry=registry,
            )
        for index, prefix in zip(range(9, 14), expected_prefixes):
            cell = unique_cells(output.tables[1].rows[index])[1]
            assert cell.text == prefix + "\n源文件危害说明"
            assert composite_value_text(output.tables[1].rows[index]) == "源文件危害说明"
        language = "en" if template_path == TEMPLATE_EN else "cn"
        assert compare_locked_skeleton(template, output) == []
        assert compare_format_anchors(template, output, language=language) == []


def test_s28_health_rows_keep_composite_handling_after_authorized_renumbering():
    for template_path in (TEMPLATE, TEMPLATE_EN):
        template = Document(str(template_path))
        output = Document(str(template_path))
        row = output.tables[1].rows[9]
        set_sequence_prefix(unique_cells(row)[0], 2, 7)
        registry = TemplateSlotRegistry.from_document(output)
        if registry.slots[(1, 9)].role != "s2_composite_value":
            clear_value_cells(output, registry)
            write_row_values(
                row, [unique_cells(row)[0].text, "重排后的健康危害"],
                table_index=1, row_index=9, registry=registry,
            )
            assert unique_cells(row)[1].text == "重排后的健康危害"
            assert compare_locked_skeleton(template, output) == []
            continue
        clear_value_cells(output, registry)
        write_row_values(
            row, [unique_cells(row)[0].text, "重排后的健康危害"],
            table_index=1, row_index=9, registry=registry,
        )
        assert composite_value_text(row) == "重排后的健康危害"
        assert unique_cells(row)[1].text.endswith("\n重排后的健康危害")
        assert compare_locked_skeleton(template, output) == []


def test_s28_route_prefix_mutation_is_blocked():
    template = Document(str(TEMPLATE))
    output = Document(str(TEMPLATE))
    cell = unique_cells(output.tables[1].rows[9])[1]
    cell.paragraphs[0].runs[0].text = "错误前缀："
    errors = compare_locked_skeleton(template, output)
    assert any("route prefix" in error for error in errors)


def test_s11_middle_sublabel_text_is_locked_but_final_value_is_writable():
    template = Document(str(TEMPLATE))
    output = Document(str(TEMPLATE))
    row = output.tables[10].rows[4]
    cells = unique_cells(row)
    original_sublabel = cells[1].text
    write_row_values(
        row, [cells[0].text, original_sublabel, "source endpoint value"],
        table_index=10, row_index=4,
        registry=TemplateSlotRegistry.from_document(output),
    )
    assert cells[1].text == original_sublabel
    assert cells[-1].text == "source endpoint value"
    assert compare_locked_skeleton(template, output) == []
    cells[1].paragraphs[0].runs[0].text = "tampered sublabel"
    errors = compare_locked_skeleton(template, output)
    assert any("sub-label/header text changed" in error for error in errors)


def test_s11_2_three_column_variant_locks_first_two_cells():
    output = Document(str(TEMPLATE))
    row = output.tables[10].rows[12]
    cells = unique_cells(row)
    cells[0].paragraphs[0].runs[0].text = cells[0].text.replace("11.7", "11.2", 1)
    original_sublabel = cells[1].text
    registry = TemplateSlotRegistry.from_document(output)
    write_row_values(
        row, [cells[0].text, original_sublabel, "source endpoint value"],
        table_index=10, row_index=12, registry=registry,
    )
    assert cells[1].text == original_sublabel
    assert cells[2].text == "source endpoint value"


def test_three_column_structure_drift_is_blocked():
    template = Document(str(TEMPLATE))
    output = Document(str(TEMPLATE))
    cell = unique_cells(output.tables[2].rows[4])[1]
    tc_pr = cell._tc.get_or_add_tcPr()
    grid_span = tc_pr.find(qn("w:gridSpan"))
    if grid_span is None:
        grid_span = tc_pr.makeelement(qn("w:gridSpan"), {})
        tc_pr.append(grid_span)
    grid_span.set(qn("w:val"), "2")
    errors = compare_format_anchors(template, output)
    assert any("format anchor changed" in error for error in errors)


def test_extra_s9_source_row_uses_cloned_style_and_source_label_only_on_insert():
    template = Document(str(TEMPLATE))
    output = Document(str(TEMPLATE))
    ensure_source_data_rows(output, 9, 24)
    row = output.tables[8].rows[24]
    write_row_values(
        row, ["9.24  New physical property:", "source value"],
        table_index=8, row_index=24, inserted_data_row=True,
    )
    assert row.cells[0].text == "9.24  New physical property:"
    assert row.cells[1].text == "source value"
    assert not compare_locked_skeleton(template, output)


def test_section9_sequence_renumber_can_rebalance_prefix_spacing():
    template = Document(str(TEMPLATE))
    output = Document(str(TEMPLATE))
    cell = output.tables[8].rows[12].cells[0]  # template 9.12 has one separator space
    updated = set_sequence_prefix(cell, 9, 5, prefix_width=5)
    assert updated.startswith("9.5  ")
    assert not compare_locked_skeleton(template, output)
    assert not compare_format_anchors(template, output)


def test_extra_s15_regulation_row_uses_cloned_note_style():
    template = Document(str(TEMPLATE))
    output = Document(str(TEMPLATE))
    ensure_source_data_rows(output, 15, 9)
    row = output.tables[14].rows[9]
    write_row_values(
        row, ["Source-backed additional regulation"],
        table_index=14, row_index=9, inserted_data_row=True,
    )
    assert row.cells[0].text == "Source-backed additional regulation"
    assert not compare_locked_skeleton(template, output)
    assert not compare_format_anchors(template, output)


def test_s81_parent_node_is_not_a_writable_note_slot():
    template = Document(str(TEMPLATE))
    output = Document(str(TEMPLATE))
    original = output.tables[7].rows[1].cells[0].text
    write_row_values(
        output.tables[7].rows[1],
        ["8.1 source label", "source value must not enter the parent node"],
        table_index=7,
        row_index=1,
    )
    assert output.tables[7].rows[1].cells[0].text == original
    assert not compare_locked_skeleton(template, output)


def test_s82_top_data_write_preserves_locked_header_and_parent_label():
    template = Document(str(TEMPLATE_EN_SOURCE))
    output = Document(str(TEMPLATE_EN_SOURCE))
    table = output.tables[7]
    original_parent = table.rows[12].cells[0].text
    audit = write_s82_top_rows(
        table,
        [["Substance A", "CN OEL", "TWA", "0.03 mg/m3"]],
        "en",
    )
    assert audit == {"record_count": 1, "row_count": 1, "placeholder": False}
    assert _unique_texts(table.rows[13]) == ["Substance", "Basis", "Type", "Value"]
    assert _unique_texts(table.rows[14]) == ["Substance A", "CN OEL", "TWA", "0.03 mg/m3"]
    assert len(table.rows) == 15
    assert table.rows[12].cells[0].text == original_parent
    assert not compare_locked_skeleton(template, output)


def test_s82_empty_records_hide_workplace_component_block():
    template = Document(str(TEMPLATE))
    output = Document(str(TEMPLATE))
    table = output.tables[7]
    audit = write_s82_top_rows(table, [], "zh")
    assert audit == {"record_count": 0, "row_count": 0, "placeholder": False, "hidden": True}
    assert len(table.rows) == 12
    assert not compare_locked_skeleton(template, output)


def test_s82_extra_records_clone_styled_data_row():
    template = Document(str(TEMPLATE_EN_SOURCE))
    output = Document(str(TEMPLATE_EN_SOURCE))
    table = output.tables[7]
    records = [
        ["Substance A", "CN OEL", "TWA", "0.03 mg/m3"],
        ["Substance B", "CN OEL", "STEL", "0.06 mg/m3"],
        ["Substance C", "CN OEL", "TWA", "100 mg/m3"],
    ]
    write_s82_top_rows(table, records, "en")
    assert len(table.rows) == 17
    assert _unique_texts(table.rows[16]) == ["Substance C", "CN OEL", "TWA", "100 mg/m3"]
    assert not compare_locked_skeleton(template, output)


def test_s82_top_header_mutation_is_blocked():
    template = Document(str(TEMPLATE))
    output = Document(str(TEMPLATE))
    output.tables[7].rows[13].cells[1].paragraphs[0].runs[0].text = "Regulation"
    errors = compare_locked_skeleton(template, output)
    assert any("locked S8.2 header text changed" in error for error in errors)


def test_s82_generic_row_write_is_blocked():
    output = Document(str(TEMPLATE))
    try:
        write_row_values(
            output.tables[7].rows[13],
            ["header", "Substance", "Basis", "Type", "Value"],
            table_index=7,
            row_index=13,
        )
    except MutationViolation:
        pass
    else:
        raise AssertionError("expected MutationViolation for S8.2 header write")
