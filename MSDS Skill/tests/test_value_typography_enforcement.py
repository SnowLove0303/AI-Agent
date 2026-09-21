# -*- coding: utf-8 -*-
"""Unit tests for unified value cell typography and audit."""
from pathlib import Path
import sys
import docx
from docx.oxml.ns import qn

SCRIPTS_DIR = Path(__file__).resolve().parents[1] / "scripts"
if str(SCRIPTS_DIR) not in sys.path:
    sys.path.insert(0, str(SCRIPTS_DIR))

from section2_ghs_policy import set_cell_value_unified
from audit_section2_release import audit_value_cell_typography


def test_set_cell_value_unified_zh():
    doc = docx.Document()
    table = doc.add_table(rows=1, cols=2)
    cell = table.cell(0, 1)

    set_cell_value_unified(cell, "第一行测试\n第二行测试", lang="zh", bold=False, size_pt=12.0)
    assert len(cell.paragraphs) == 2
    for p in cell.paragraphs:
        assert len(p.runs) == 1
        r = p.runs[0]
        assert r.font.size.pt == 12.0
        assert r.font.name == "宋体"
        assert r.font.bold is False
        assert p.alignment == 0
    assert cell._tc.tcPr.find(qn("w:vAlign")).get(qn("w:val")) == "center"


def test_set_cell_value_unified_en():
    doc = docx.Document()
    table = doc.add_table(rows=1, cols=2)
    cell = table.cell(0, 1)

    set_cell_value_unified(cell, "Line 1\nLine 2", lang="en", bold=False, size_pt=12.0)
    assert len(cell.paragraphs) == 2
    for p in cell.paragraphs:
        assert len(p.runs) == 1
        r = p.runs[0]
        assert r.font.size.pt == 12.0
        assert r.font.name == "Times New Roman"
        assert r.font.bold is False
        assert p.alignment == 0
    assert cell._tc.tcPr.find(qn("w:vAlign")).get(qn("w:val")) == "center"


def test_audit_value_cell_typography():
    doc = docx.Document()
    table0 = doc.add_table(rows=1, cols=1)  # Table 0
    table1 = doc.add_table(rows=2, cols=2)  # Table 1 (Section 2)
    
    # Row 0: Header
    table1.rows[0].cells[0].text = "2. 危险性概述"
    table1.rows[0].cells[1].text = "2. 危险性概述"

    # Row 1: Valid value cell
    table1.rows[1].cells[0].text = "2.1 GHS危险性类别："
    set_cell_value_unified(table1.rows[1].cells[1], "根据GHS不属于危险物", lang="zh")

    errors = audit_value_cell_typography(doc)
    assert not errors, f"Expected no errors, got: {errors}"
