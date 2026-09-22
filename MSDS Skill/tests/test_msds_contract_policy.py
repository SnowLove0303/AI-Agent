from pathlib import Path
import sys

import pytest
from docx import Document

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "scripts"))

from missing_data_policy import (  # noqa: E402
    SourceState,
    apply_source_absence_policy,
    classify_source_value,
)
from msds_pipeline import (  # noqa: E402
    apply_post_overwrite_fine_tuning,
    format_revision_date,
    plan_body_write,
    suppress_empty_s82_engineering_control,
    write_body,
    write_header_footer,
)
from section11_alignment import align_s11_rows  # noqa: E402
from template_mutation_whitelist import (  # noqa: E402
    TemplateSlotRegistry,
    compare_format_anchors,
    normalize_value_text,
    write_row_values,
    unique_cells,
)


TEMPLATE = ROOT / "examples" / "template_reference.docx"


def test_source_states_are_distinct():
    assert classify_source_value("厚度≧0.4mm；穿透时间≧480min.") == SourceState.SUPPORTED
    assert classify_source_value("无数据。") == SourceState.EXPLICIT_MISSING
    assert classify_source_value("不适用") == SourceState.NOT_APPLICABLE
    assert classify_source_value("") == SourceState.ABSENT


def test_s8_recommendation_value_is_source_gated_but_gloves_keep_mapping():
    document = Document(str(TEMPLATE))
    registry = TemplateSlotRegistry.from_document(document)
    table = document.tables[7]

    write_row_values(
        table.rows[8], ["建议：", "污染的手套应废弃。"],
        table_index=7, row_index=8, registry=registry,
    )
    write_row_values(
        table.rows[5], ["氟化橡胶 –FKM:", "厚度≧0.4mm；穿透时间≧480min."],
        table_index=7, row_index=5, registry=registry,
    )

    assert unique_cells(table.rows[8])[1].text == "污染的手套应废弃。"
    assert any(
        run.text.strip() and run.bold is not True
        for paragraph in unique_cells(table.rows[8])[1].paragraphs
        for run in paragraph.runs
    )
    assert compare_format_anchors(
        Document(str(TEMPLATE)), document, language="cn"
    ) == []
    assert unique_cells(table.rows[5])[1].text == "厚度≧0.4mm；穿透时间≧480min."

    # Source-absent rows may be removed before the recommendation row.  The
    # recommendation's non-bold value format must remain auditable by its
    # semantic label after that approved reordering.
    reordered = Document(str(TEMPLATE))
    reordered_registry = TemplateSlotRegistry.from_document(reordered)
    write_row_values(
        reordered.tables[7].rows[8], ["建议：", "污染的手套应废弃。"],
        table_index=7, row_index=8, registry=reordered_registry,
    )
    reordered.tables[7]._tbl.remove(reordered.tables[7].rows[7]._tr)
    assert compare_format_anchors(
        Document(str(TEMPLATE)), reordered, language="cn"
    ) == []

    absent = Document(str(TEMPLATE))
    absent_registry = TemplateSlotRegistry.from_document(absent)
    write_row_values(
        absent.tables[7].rows[8], ["建议：", ""],
        table_index=7, row_index=8, registry=absent_registry,
    )
    assert unique_cells(absent.tables[7].rows[8])[1].text == ""


def test_synthetic_spaced_slash_becomes_semantic_break_but_compact_slash_survives():
    assert normalize_value_text("第一句 / 第二句") == "第一句\n第二句"
    assert normalize_value_text("通风/排气；有/无") == "通风/排气；有/无"
    assert normalize_value_text("第一句 / / 第二句") == "第一句\n第二句"


def test_empty_s82_engineering_control_row_is_hidden():
    document = Document(str(TEMPLATE))
    table = document.tables[7]
    table.rows[11].cells[-1].text = ""
    result = suppress_empty_s82_engineering_control(document)
    assert result["hidden"] is True
    assert not any(row.cells[0].text.strip().startswith("8.2") for row in table.rows)


def test_pu1001_source_text_is_written_without_summary_loss():
    document = Document(str(TEMPLATE))
    registry = TemplateSlotRegistry.from_document(document)
    ingestion = "正常使用时只有轻微的摄入危害，可能引起胃不适导致呕吐引起胃损伤"
    write_row_values(
        document.tables[1].rows[10], ["2.8  健康危害", ingestion],
        table_index=1, row_index=10, registry=registry,
    )
    assert unique_cells(document.tables[1].rows[10])[1].text == "食入：\n" + ingestion


def test_s10_missing_rows_are_hidden():
    document = Document(str(TEMPLATE))
    table = document.tables[9]  # Section 10
    unique_cells(table.rows[4])[1].text = ""
    unique_cells(table.rows[5])[1].text = "无数据"
    facts = {"s10": [
        ["10.1", "稳定"], ["10.2", "分解"], ["10.3", "无反应"],
        ["10.4", ""], ["10.5", "无数据"],
    ]}

    audit = apply_source_absence_policy(document, facts, unique_cells)
    labels = [unique_cells(row)[0].text for row in table.rows[1:]]
    assert any("10.4" in label for label in audit["source_absent_removed"])
    assert not any("10.4" in label for label in labels)
    assert not any("10.5" in label for label in labels)


def test_s11_7_keeps_explicit_missing_but_hides_unmatched_children():
    document = Document(str(TEMPLATE))
    table = document.tables[10]
    cells = unique_cells(table.rows[12])
    cells[-1].text = "无数据"
    facts = {"s11": [["说明"]] + [["endpoint", "value"] for _ in range(10)] + [
        ["11.7 生殖毒性：", "生育力", "无数据"],
        ["11.7 生殖毒性：", "致畸形", ""],
        ["11.7 生殖毒性：", "体外遗传毒性", ""],
    ]}
    apply_source_absence_policy(document, facts, unique_cells)
    labels = [unique_cells(row)[0].text for row in table.rows[1:]]
    assert any("生殖毒性" in label for label in labels)
    assert len([label for label in labels if "11.7" in label]) == 1


def test_s11_7_explicit_missing_writes_to_blank_template_value_slot():
    document = Document(str(TEMPLATE))
    registry = TemplateSlotRegistry.from_document(document)
    row = document.tables[10].rows[12]
    write_row_values(
        row,
        ["11.7 生殖毒性：", "生育力", "无数据"],
        table_index=10,
        row_index=12,
        registry=registry,
    )
    assert unique_cells(row)[-1].text == "无数据"


def test_s11_non_11_7_source_missing_endpoint_is_preserved():
    document = Document(str(TEMPLATE))
    table = document.tables[10]
    registry = TemplateSlotRegistry.from_document(document)
    row = table.rows[11]
    write_row_values(
        row,
        ["11.6 致癌性：", "", "无数据资料。"],
        table_index=10,
        row_index=11,
        registry=registry,
    )
    facts = {"s11": [["说明"]] + [["endpoint", "value"] for _ in range(10)] + [
        ["11.6 致癌性：", "", "无数据资料。"],
    ]}
    apply_source_absence_policy(document, facts, unique_cells)
    labels = [unique_cells(row)[0].text for row in table.rows[1:]]
    assert any("11.6" in label for label in labels)


def test_s11_compressed_facts_are_aligned_by_endpoint_before_rows_are_hidden():
    document = Document(str(TEMPLATE))
    facts = {f"s{section}": [] for section in range(1, 17)}
    facts["s11"] = [
        ["该产品无可用的毒理学研究。"],
        ["11.1 急性毒性：", "经口：", "半数致死剂量（LD50）/大鼠：>2,000 mg/kg"],
        ["11.2 主要皮肤刺激性：", "物种：兔子\n结果：轻微刺激"],
        ["11.4 致敏性：", "物种：豚鼠\n结果：皮肤接触可能致敏"],
        ["11.7 生殖毒性：", "致畸形", "无数据资料。"],
        ["11.10 附加信息：", "眼睛接触可能造成刺激。"],
    ]
    facts["s8_control_parameters"] = {"zh": [], "en": []}

    write_body(document, facts, "zh")
    table = document.tables[10]
    visible = []
    for row in table.rows[1:]:
        cells = unique_cells(row)
        visible.append((cells[0].text, [cell.text for cell in cells[1:]]))

    skin = next(item for item in visible if "11.2" in item[0])
    sensitization = next(
        item for item in visible
        if any("皮肤接触可能致敏" in value for value in item[1])
    )
    teratogenicity = next(item for item in visible if item[1] and "致畸形" in item[1][0])
    assert "轻微刺激" in skin[1][0]
    assert any("皮肤接触可能致敏" in value for value in sensitization[1])
    assert teratogenicity[1][-1] == "无数据资料。"
    assert not any("吸入：" in label or "经皮：" in label for label, _ in visible)
    assert all("\n\n" not in value for _, values in visible for value in values)


@pytest.mark.parametrize(
    ("template_name", "source_row", "expected"),
    [
        (
            "template_reference.docx",
            ["11.1 急性毒性：", "经口：", "毒性：经口（鼠）LD50：＞5000 mg/kg"],
            "（鼠）LD50：＞5000 mg/kg",
        ),
        (
            "template_reference_en.docx",
            ["11.1 Acute toxicity:", "Oral:", "Toxicity: Oral (rat) LD50: >5000 mg/kg"],
            "(rat) LD50: >5000 mg/kg",
        ),
    ],
)
def test_s11_value_drops_only_locked_endpoint_and_route_headings(
    template_name, source_row, expected
):
    document = Document(str(ROOT / "examples" / template_name))
    aligned = align_s11_rows([source_row], document.tables[10])
    row = next(item for item in aligned if item[0].strip().startswith("11.1"))
    assert row[-1] == expected


def test_write_plan_is_resolved_before_row_mutation():
    document = Document(str(TEMPLATE))
    facts = {f"s{section}": [] for section in range(1, 17)}
    facts["s9"] = [[f"9.{index}", f"value-{index}"] for index in range(1, 25)]
    facts["s11"] = [["该产品无可用的毒理学研究。"]]
    plans, policy_facts = plan_body_write(document, facts, "zh")
    assert len(document.tables[8].rows) == 24
    s9 = next(plan for plan in plans if plan.section == 9)
    assert s9.inserted_data_rows == 1
    assert s9.semantic_mode == "field_rows"
    assert len(policy_facts["s11"]) == len(next(plan for plan in plans if plan.section == 11).rows)
    assert len(policy_facts["s11"]) > len(facts["s11"])


def test_post_overwrite_fine_tuning_is_explicit_and_bounded():
    document = Document(str(TEMPLATE))
    facts = {f"s{section}": [] for section in range(1, 17)}
    facts["s11"] = [["该产品无可用的毒理学研究。"]]
    result = apply_post_overwrite_fine_tuning(document, facts)
    assert result["source_presence_policy"]["note_only"]["s11"] is True
    assert result["section2_policy"]["removed_count"] >= 1


def test_s11_alignment_rejects_unclassified_endpoint_instead_of_guessing():
    document = Document(str(TEMPLATE))
    with pytest.raises(ValueError, match="no matching template row"):
        align_s11_rows([["11.99 未知毒理项目：", "错误归类"]], document.tables[10])


def test_s12_template_only_explanation_is_removed():
    document = Document(str(TEMPLATE))
    table = document.tables[11]
    unique_cells(table.rows[1])[0].text = "该产品无可用的生态毒理学研究。"
    facts = {"s12": [
        ["该产品无可用的生态毒理学研究。"],
        ["12.1 生态毒性：", "禁止排入环境。"],
    ]}
    apply_source_absence_policy(document, facts, unique_cells)
    visible = [unique_cells(row)[0].text for row in table.rows[1:]]
    assert "以下为类似产品的生态毒理学参考数据：" not in visible


def test_s12_missing_endpoint_is_hidden_but_source_note_remains():
    document = Document(str(TEMPLATE))
    table = document.tables[11]
    note_rows = [unique_cells(table.rows[index])[0].text for index in (1, 2)]
    for row in table.rows[1:]:
        cells = unique_cells(row)
        if len(cells) > 1:
            cells[-1].text = ""
    unique_cells(table.rows[4])[-1].text = "有效数据"
    facts = {"s12": [
        [note_rows[0]], [note_rows[1]],
        ["12.1 生态毒性：", "无数据资料。"],
        ["12.2 持久性和降解性：", "有效数据"],
        ["12.3 其他不利的影响：", "无数据资料。"],
    ]}
    apply_source_absence_policy(document, facts, unique_cells)
    visible = [unique_cells(row)[0].text for row in document.tables[11].rows[1:]]
    assert len(visible) == 3
    assert any("12.2" in label for label in visible)
    assert not any("12.1" in label or "12.3" in label for label in visible)


def test_s15_trailing_blank_row_is_removed():
    document = Document(str(TEMPLATE))
    facts = {"s15": [[f"rule-{i}"] for i in range(7)]}
    apply_source_absence_policy(document, facts, unique_cells)
    assert len(document.tables[14].rows) == 8


def test_note_only_sections_keep_only_the_explanation_row():
    document = Document(str(TEMPLATE))
    facts = {
        "s11": [["该产品无可用的毒理学研究。"]],
        "s12": [["该产品无可用的生态毒理学研究。"]],
    }
    audit = apply_source_absence_policy(document, facts, unique_cells)
    assert audit["note_only"] == {"s11": True, "s12": True}
    assert len(document.tables[10].rows) == 2
    assert len(document.tables[11].rows) == 2


def test_header_version_is_inherited_from_template():
    document = Document(str(TEMPLATE))
    version_before = next(
        paragraph.text for paragraph in document.sections[0].header.paragraphs
        if "Version" in paragraph.text
    )
    write_header_footer(document, "zh", "guanzhi", "PU-1001", "2025/2/22")
    version_after = next(
        paragraph.text for paragraph in document.sections[0].header.paragraphs
        if "Version" in paragraph.text
    )
    assert version_after == version_before == "Version：V1.0"


def test_revision_date_is_currently_formatted_and_template_p_guard_survives():
    document = Document(str(TEMPLATE))
    write_header_footer(document, "zh", "guanzhi", "PU-1001", "2026/9/11")
    footer = document.sections[0].footer.tables[0].rows[0]
    assert unique_cells(footer)[1].text.startswith("P修订日期：2026年9月11日")
    assert format_revision_date("en", "2026年9月11日") == "September 11, 2026"

    english = Document(str(ROOT / "examples" / "template_reference_en.docx"))
    write_header_footer(english, "en", "guanzhi", "PU-1001", "2026年9月11日")
    en_footer = english.sections[0].footer.tables[0].rows[0]
    assert unique_cells(en_footer)[1].text.startswith("PRevision date: September 11, 2026")
