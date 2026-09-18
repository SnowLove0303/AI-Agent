import json
from pathlib import Path
import sys

import pytest
from docx import Document

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "scripts"))

from family_profile import FamilyProfileError, load_profile, review_profile, validate_profile  # noqa: E402
from preflight_facts import run as run_preflight  # noqa: E402


def _write_profile(path: Path, candidate: dict, model: str = "EFF-TEST") -> Path:
    path.write_text(json.dumps({
        "profile_id": "test-profile",
        "family": "test-family",
        "models": [model],
        "candidates": [candidate],
    }, ensure_ascii=False), encoding="utf-8")
    return path


def test_unapproved_candidate_stays_a_suggestion(tmp_path):
    source = tmp_path / "source.docx"
    Document().save(source)
    profile = _write_profile(tmp_path / "profile.json", {
        "candidate_id": "candidate-1",
        "target": "s9.1",
        "value": "not in source",
        "evidence_required": True,
        "disposition": "candidate",
    })

    report = review_profile(profile, source, "EFF-TEST")

    assert report["status"] == "passed"
    assert len(report["suggestions"]) == 1
    assert report["confirmed"] == []
    assert report["writes_facts"] is False
    assert report["writes_docx"] is False


def test_confirmed_candidate_requires_current_source_anchor(tmp_path):
    source = tmp_path / "source.docx"
    document = Document()
    document.add_paragraph("confirmed source property")
    document.save(source)
    profile = _write_profile(tmp_path / "profile.json", {
        "candidate_id": "candidate-1",
        "target": "s9.1",
        "value": "confirmed source property",
        "evidence_required": True,
        "disposition": "confirmed",
        "evidence": {
            "type": "current_source",
            "source_locator": "paragraph[1]",
            "source_text": "confirmed source property",
        },
    })

    report = review_profile(profile, source, "EFF-TEST")

    assert report["status"] == "passed"
    assert len(report["confirmed"]) == 1


def test_conflict_or_missing_source_anchor_is_blocked(tmp_path):
    source = tmp_path / "source.docx"
    Document().save(source)
    conflict = _write_profile(tmp_path / "conflict.json", {
        "candidate_id": "conflict-1",
        "target": "s11.1",
        "value": "conflict",
        "evidence_required": True,
        "disposition": "conflict",
    })
    missing_anchor = _write_profile(tmp_path / "missing.json", {
        "candidate_id": "confirmed-1",
        "target": "s11.1",
        "value": "not in source",
        "evidence_required": True,
        "disposition": "confirmed",
        "evidence": {
            "type": "current_source",
            "source_locator": "paragraph[1]",
            "source_text": "not in source",
        },
    })

    conflict_report = review_profile(conflict, source, "EFF-TEST")
    missing_report = review_profile(missing_anchor, source, "EFF-TEST")

    assert conflict_report["status"] == "failed"
    assert missing_report["status"] == "failed"
    assert conflict_report["blocked"][0]["candidate_id"] == "conflict-1"


def test_profile_shape_and_scope_are_blocked(tmp_path):
    malformed = {"profile_id": "x", "family": "x", "models": ["OTHER"], "candidates": []}
    assert any("requested model" in error for error in validate_profile(malformed, "EFF-TEST"))
    path = tmp_path / "profile.txt"
    path.write_text("{}", encoding="utf-8")
    with pytest.raises(FamilyProfileError, match="JSON"):
        load_profile(path)


def test_preflight_reports_family_profile_blocker_before_clone(tmp_path):
    source = tmp_path / "source.docx"
    Document().save(source)
    profile = _write_profile(tmp_path / "profile.json", {
        "candidate_id": "conflict-1",
        "target": "s11.1",
        "value": "conflict",
        "evidence_required": True,
        "disposition": "conflict",
    })
    facts = tmp_path / "facts.json"
    facts.write_text(json.dumps({"model": "EFF-TEST"}), encoding="utf-8")

    result = run_preflight(source, facts, "EFF-TEST", family_profile=profile)

    assert result["status"] == "blocked"
    assert any("family profile" in error for error in result["errors"])
    assert result["template_clone_started"] is False
