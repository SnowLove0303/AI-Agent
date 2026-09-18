"""Focused v3.25 regressions from the defect/evolution report."""

from pathlib import Path
import subprocess
import sys

from docx import Document
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "scripts"))

from section_overwrite_rules import sanitize_section_payload  # noqa: E402
import convert_docx_to_pdf as converter  # noqa: E402
from source_grounding import audit as audit_source_grounding  # noqa: E402
from s11_layout_policy import audit as audit_s11_layout  # noqa: E402
from s11_layout_policy import normalize as normalize_s11_layout  # noqa: E402
from template_mutation_whitelist import (  # noqa: E402
    TemplateSlotRegistry,
    audit_values_nonbold,
    clear_value_cells,
    compare_format_anchors,
    set_s28_composite_value_cell,
    set_value_cell_text,
    unique_cells,
    write_s82_top_rows,
    write_row_values,
)


TEMPLATE = ROOT / "examples" / "template_reference.docx"


def _row_height(row):
    tr_pr = row._tr.trPr
    node = tr_pr.find(qn("w:trHeight")) if tr_pr is not None else None
    return int(node.get(qn("w:val"))) if node is not None else None


def test_s15_and_s16_one_cell_payloads_do_not_duplicate_source_label():
    document = Document(str(TEMPLATE))
    registry = TemplateSlotRegistry.from_document(document)
    clear_value_cells(document, registry)

    write_row_values(
        document.tables[14].rows[1],
        ["适用法律法规", "适用法律法规：GB/T 16483"],
        table_index=14, row_index=1, registry=registry,
    )
    write_row_values(
        document.tables[15].rows[1],
        ["免责声明", "免责声明：仅供安全使用参考"],
        table_index=15, row_index=1, registry=registry,
    )

    assert unique_cells(document.tables[14].rows[1])[0].text == "GB/T 16483"
    assert unique_cells(document.tables[15].rows[1])[0].text == "仅供安全使用参考"


def test_empty_intermediate_rows_are_removed_before_capacity_check():
    rows = [["", ""], ["法规A", "法规A：适用"], ["", ""], ["法规B"]]
    assert sanitize_section_payload(15, rows) == [
        ["法规A", "法规A：适用"], ["法规B"],
    ]


def test_s5_and_s13_empty_capacity_rows_are_removed_without_reordering_values():
    for section in (5, 13):
        rows = [["模板标签A", ""], ["源标签", "源值"], ["模板标签B", ""]]
        assert sanitize_section_payload(section, rows) == [["源标签", "源值"]]


def test_source_grounding_rejects_template_only_value(tmp_path):
    source = tmp_path / "source.docx"
    source_doc = Document()
    source_doc.add_paragraph("产品A")
    source_doc.save(source)
    facts = {
        "model": "TEST-1",
        "source_sha256": "unused-by-unit-test",
        "fact_ledger": [{
            "fact_id": "FACT-0001",
            "source_text": "产品A",
            "normalized_value": "产品A",
        }],
        "output_traceability": {
            "items": [{"output_values": {"zh": "产品A", "en": "Product A"}}]
        },
        "zh": {"s1": [["产品名称", "模板示例水和沸点"]]},
        "en": {"s1": [["Product name", "Product A"]]},
    }
    report = audit_source_grounding(facts, source, "TEST-1")
    assert report["status"] == "failed"
    assert any("zh.s1[1]" in error for error in report["errors"])


def test_source_grounding_accepts_reviewed_translation_value(tmp_path):
    source = tmp_path / "source.docx"
    source_doc = Document()
    source_doc.add_paragraph("产品A")
    source_doc.save(source)
    facts = {
        "model": "TEST-1",
        "fact_ledger": [{
            "fact_id": "FACT-0001",
            "source_text": "产品A",
            "normalized_value": "产品A",
        }],
        "output_traceability": {
            "items": [{"output_values": {"zh": "产品A", "en": "Product A"}}]
        },
        "zh": {"s1": [["产品名称", "产品A"]]},
        "en": {"s1": [["Product name", "Product A"]]},
    }
    report = audit_source_grounding(facts, source, "TEST-1")
    assert report["status"] == "passed"


def test_source_grounding_reads_header_and_footer_source_anchors(tmp_path):
    source = tmp_path / "source.docx"
    source_doc = Document()
    source_doc.add_paragraph("正文")
    source_doc.sections[0].header.paragraphs[0].text = "HeaderFact-1234"
    source_doc.sections[0].footer.paragraphs[0].text = "FooterFact-5678"
    source_doc.save(source)
    facts = {
        "model": "TEST-1",
        "fact_ledger": [
            {"fact_id": "FACT-H", "source_text": "HeaderFact-1234", "normalized_value": "HeaderFact-1234"},
            {"fact_id": "FACT-F", "source_text": "FooterFact-5678", "normalized_value": "FooterFact-5678"},
        ],
        "output_traceability": {"items": []},
        "zh": {"s1": [["标签", "HeaderFact-1234"], ["标签", "FooterFact-5678"]]},
        "en": {"s1": []},
    }
    report = audit_source_grounding(facts, source, "TEST-1")
    assert report["status"] == "passed"


def test_s114_short_row_height_is_output_only_and_auditable():
    template = Document(str(TEMPLATE))
    output = Document(str(TEMPLATE))
    row = output.tables[10].rows[9]
    unique_cells(row)[1].text = "短文本"
    assert _row_height(template.tables[10].rows[9]) == 1169
    exception = normalize_s11_layout(
        output, [["11.4 致敏性：", "短文本"]]
    )
    assert exception["applied"] is True
    assert _row_height(output.tables[10].rows[9]) == 285
    assert audit_s11_layout(template, output, exception)["errors"] == []
    assert audit_s11_layout(template, output, None)["errors"]
    assert _row_height(template.tables[10].rows[9]) == 1169


def test_s28_multi_run_prefix_is_rebuilt_from_original_run_fragments():
    document = Document(str(TEMPLATE))
    cell = unique_cells(document.tables[1].rows[9])[1]
    paragraph = cell.paragraphs[0]
    first = paragraph.runs[0]
    original = first.text
    first.text = original[:2]
    second = paragraph.add_run(original[2:])
    second.bold = first.bold

    set_s28_composite_value_cell(cell, "源文件健康危害说明")

    assert paragraph.text == original + "\n源文件健康危害说明"
    assert len(paragraph.runs) >= 4
    assert paragraph.runs[0].text + paragraph.runs[1].text == original


def test_writable_values_strip_template_bold_and_keep_other_inherited_style():
    for template_path in (TEMPLATE, ROOT / "examples" / "template_reference_en.docx"):
        output = Document(str(template_path))
        registry = TemplateSlotRegistry.from_document(output)
        clear_value_cells(output, registry)
        cell = unique_cells(output.tables[0].rows[1])[1]
        paragraph = cell.paragraphs[0]
        paragraph.runs[0].bold = True
        set_value_cell_text(cell, "source-backed value")
        assert audit_values_nonbold(output, language="en" if "_en" in template_path.name else "cn") == []
        assert paragraph.runs[0].bold is not True
        assert paragraph.runs[0]._r.rPr.find(qn("w:sz")) is not None


def test_value_nonbold_override_wins_over_bold_paragraph_inheritance():
    output = Document(str(TEMPLATE))
    registry = TemplateSlotRegistry.from_document(output)
    clear_value_cells(output, registry)
    cell = unique_cells(output.tables[0].rows[1])[1]
    p_pr = cell.paragraphs[0]._p.get_or_add_pPr()
    r_pr = p_pr.find(qn("w:rPr"))
    if r_pr is None:
        r_pr = p_pr.makeelement(qn("w:rPr"), {})
        p_pr.append(r_pr)
    r_pr.append(r_pr.makeelement(qn("w:b"), {}))
    set_value_cell_text(cell, "source-backed value")
    assert audit_values_nonbold(output) == []


def test_nonbold_value_audit_blocks_manual_bold_value_but_ignores_locked_labels():
    output = Document(str(TEMPLATE))
    registry = TemplateSlotRegistry.from_document(output)
    clear_value_cells(output, registry)
    cell = unique_cells(output.tables[0].rows[1])[1]
    set_value_cell_text(cell, "source-backed value")
    cell.paragraphs[0].runs[0].bold = True
    problems = audit_values_nonbold(output)
    assert any(item["type"] == "bold_writable_value" for item in problems)
    assert not any(item["table"] == 0 and item["cell"] == 0 for item in problems)


def test_format_audit_allows_only_bold_removal_on_writable_values():
    template = Document(str(TEMPLATE))
    output = Document(str(TEMPLATE))
    registry = TemplateSlotRegistry.from_document(output)
    clear_value_cells(output, registry)
    write_row_values(
        output.tables[14].rows[1],
        ["适用法律法规", "适用法律法规：GB/T 16483"],
        table_index=14, row_index=1, registry=registry,
    )
    assert compare_format_anchors(template, output, language="cn") == []


def test_nonbold_gate_covers_all_supported_value_topologies():
    for template_path, language, s82_language in (
        (TEMPLATE, "cn", "zh"),
        (ROOT / "examples" / "template_reference_en.docx", "en", "en"),
    ):
        output = Document(str(template_path))
        registry = TemplateSlotRegistry.from_document(output)
        clear_value_cells(output, registry)
        write_row_values(
            output.tables[1].rows[9],
            [unique_cells(output.tables[1].rows[9])[0].text, "health route value"],
            table_index=1, row_index=9, registry=registry,
        )
        write_row_values(
            output.tables[2].rows[4],
            ["component", "123-45-6", "10"],
            table_index=2, row_index=4, registry=registry,
        )
        write_s82_top_rows(
            output.tables[7], [["substance", "basis", "type", "value"]], s82_language
        )
        write_row_values(
            output.tables[10].rows[4],
            [unique_cells(output.tables[10].rows[4])[0].text, "sublabel", "endpoint"],
            table_index=10, row_index=4, registry=registry,
        )
        write_row_values(
            output.tables[14].rows[1], ["regulation", "source regulation"],
            table_index=14, row_index=1, registry=registry,
        )
        assert audit_values_nonbold(output, language=language) == []


def test_wps_timeout_kills_and_reaps_owned_process(monkeypatch):
    class HungProcess:
        def __init__(self):
            self.returncode = None
            self.killed = False
            self.communicate_calls = 0

        def communicate(self, timeout=None):
            self.communicate_calls += 1
            if timeout is not None and not self.killed:
                raise subprocess.TimeoutExpired(["wpscli.exe"], timeout)
            self.returncode = -9
            return "", ""

        def kill(self):
            self.killed = True

        def poll(self):
            return self.returncode

    process = HungProcess()
    monkeypatch.setattr(converter.subprocess, "Popen", lambda *args, **kwargs: process)

    try:
        converter._run_command(["wpscli.exe", "word2pdf"], timeout=1)
    except subprocess.TimeoutExpired:
        pass
    else:  # pragma: no cover - assertion branch
        raise AssertionError("a hung converter must surface TimeoutExpired")

    assert process.killed is True
    assert process.communicate_calls == 2
    assert process.returncode == -9
