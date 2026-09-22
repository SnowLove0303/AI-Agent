import json
from pathlib import Path

import pytest

import sys
sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "scripts"))

from extract_source_facts import (
    Extraction,
    _extract_docx,
    extract,
    extract_s11,
    extract_s12,
    extract_s2,
    extract_s8,
    extract_s9,
    main,
    split_inline_protective_material,
)
from source_ingest import SourceSelectionError


ROOT = Path(__file__).resolve().parents[1]
FIXTURE = ROOT / "examples" / "regression_HPU-7660_source.docx"


def test_extractor_detects_model_and_covers_every_cell():
    data = extract(FIXTURE)
    assert data["model"] == "HPU-7660"
    assert len(data["source_sha256"]) == 64
    assert set(data["sections"]) == {f"s{i}" for i in range(1, 17)}
    assert data["coverage"]["unmapped"] == []
    assert data["coverage"]["heading_skipped"]
    mapping = data["source_mapping"]
    assert mapping["status"] == "needs-review"
    assert {item["source_section"] for item in mapping["items"]} == {
        f"s{i}" for i in range(1, 17)
    }
    assert mapping["source_sha256"] == data["source_sha256"]
    assert mapping["unresolved"]
    assert all("['']" not in item["source_text"] for item in mapping["items"])
    assert data["source_coverage"]["status"] == "ready"
    assert data["source_coverage"]["unmapped"] == []
    assert data["source_coverage"]["unreadable"] == []
    assert data["source_coverage"]["counts"]["source_unit_count"] == len(
        data["source_coverage"]["source_units"]
    )
    assert data["fact_ledger"]
    assert len({item["fact_id"] for item in data["fact_ledger"]}) == len(data["fact_ledger"])
    assert all(item["source_unit_ids"] for item in data["fact_ledger"])
    assert data["output_traceability"]["status"] == "needs-review"


def test_extractor_splits_components_and_flags_judgment_points():
    data = extract(FIXTURE)
    issues = {(r["section"], r["issue"]) for r in data["review"]}
    assert ("s2", "per-route-synthesis") in issues
    assert ("s11", "alias-mucosa-to-eye") in issues
    assert ("s11", "repro-split") in issues
    components = [row for row in data["sections"]["s3"] if len(row) == 3]
    assert components


def test_extractor_en_skeleton_copies_only_language_independent_facts():
    data = extract(FIXTURE)
    skeleton = data["sections_en_skeleton"]
    assert skeleton["s3"]
    for row in skeleton["s3"]:
        assert row["name_en"] == ""
        assert row["cas"]
    assert skeleton["note"]


def test_extractor_marks_missing_sentinels_without_deciding():
    data = extract(FIXTURE)
    s9 = data["sections"]["s9"]
    assert any(entry.get("omit") for entry in s9)


def test_cli_rejects_requested_model_mismatch(monkeypatch, tmp_path):
    monkeypatch.setattr(sys, "argv", [
        "extract_source_facts.py", str(FIXTURE), "--model", "WRONG-999",
        "--out", str(tmp_path / "facts.json"),
    ])
    with pytest.raises(SourceSelectionError, match="does not match requested"):
        main()


def test_extractor_splits_inline_glove_material_values():
    assert split_inline_protective_material(
        "氟化橡胶 –FKM:厚度≧0.4mm；穿透时间≧480min."
    ) == ["氟化橡胶 –FKM:", "厚度≧0.4mm；穿透时间≧480min."]


def test_s2_keeps_category_value_and_source_other_hazards():
    from docx import Document

    document = Document()
    table = document.add_table(rows=1, cols=1)
    for text in ("2.1 物质或混合物的分类", "GHS危险性类别:",
                 "根据GHS不属于危害化学品", "2.2 标签要素", "GHS-象形图",
                 "根据GHS不属于危害化学品", "2.3 其他危险 无适用资料。"):
        row = table.add_row()
        row.cells[0].text = text
    data = extract_s2(table, Extraction())
    assert data["ghs_classes"] == ["根据GHS不属于危害化学品"]
    assert data["label_elements"] == ["根据GHS不属于危害化学品"]
    assert data["other_hazards"] == "无适用资料。"


def test_s2_recognizes_unnumbered_other_hazards_line_without_inventing_value():
    from docx import Document

    document = Document()
    table = document.add_table(rows=1, cols=1)
    for text in ("2.1 物质或混合物的分类", "其他危险：无适用资料。"):
        row = table.add_row()
        row.cells[0].text = text
    data = extract_s2(table, Extraction())
    assert data["other_hazards"] == "无适用资料。"

    empty_document = Document()
    empty_table = empty_document.add_table(rows=1, cols=1)
    row = empty_table.add_row()
    row.cells[0].text = "其他危害："
    empty_data = extract_s2(empty_table, Extraction())
    assert empty_data["other_hazards"] == ""


def test_s2_keeps_label_ingredient_explanation_out_of_signal_word():
    from docx import Document

    document = Document()
    table = document.add_table(rows=1, cols=1)
    for text in (
        "2.2 标签要素",
        "必须列在标签上的有害成分：",
        "基于HDI的亲水脂肪族聚异氰酸酯",
        "信号词：警告",
    ):
        row = table.add_row()
        row.cells[0].text = text
    data = extract_s2(table, Extraction())
    assert data["label_ingredients"] == ["基于HDI的亲水脂肪族聚异氰酸酯"]
    assert data["label_elements"] == [
        "必须列在标签上的有害成分：\n基于HDI的亲水脂肪族聚异氰酸酯"
    ]
    assert data["signal"] == "警告"
    assert "基于HDI的亲水脂肪族聚异氰酸酯" not in data["other"]


def test_s2_does_not_treat_english_label_elements_heading_as_ingredient():
    from docx import Document

    document = Document()
    table = document.add_table(rows=1, cols=1)
    for text in (
        "2.2 Label Elements",
        "Hazardous ingredients required to be listed on the label:",
        "HDI-based hydrophilic aliphatic polyisocyanate",
        "Signal Word: Warning",
    ):
        row = table.add_row()
        row.cells[0].text = text
    data = extract_s2(table, Extraction())
    assert data["label_ingredients"] == ["HDI-based hydrophilic aliphatic polyisocyanate"]
    assert data["signal"] == "Warning"


def test_s11_splits_combined_preamble_before_endpoint_mapping():
    from docx import Document

    document = Document()
    table = document.add_table(rows=1, cols=1)
    row = table.add_row()
    row.cells[0].text = "该产品无可用的毒理学研究。下面是这些成分的毒理学数据。 11.1 毒理学效应"
    row = table.add_row()
    row.cells[0].text = "急性毒性，经口"
    row.cells[0].add_paragraph("半数致死剂量（LD50）/大鼠：> 2,000 mg/kg")
    rows = extract_s11(table, Extraction())
    assert rows[0] == ["该产品无可用的毒理学研究。 下面是这些成分的毒理学数据。"] or rows[0][0] == "该产品无可用的毒理学研究。"


def test_s8_maps_control_parameters_to_engineering_and_preserves_recommendation():
    from docx import Document

    document = Document()
    table = document.add_table(rows=1, cols=2)
    for label, value in (
        ("8.1 控制参数", "8.1 控制参数"),
        ("根据EC指令2006/121/EG,无可用的接触限值信息", "根据EC指令2006/121/EG,无可用的接触限值信息"),
        ("8.2 暴露控制", "8.2 暴露控制"),
        ("呼吸系统防护：", "喷涂过程中要求有呼吸防护设备。"),
        ("建议：", "污染的手套应废弃。"),
    ):
        row = table.add_row()
        row.cells[0].text = label
        row.cells[1].text = value
    rows, control_parameters = extract_s8(table, Extraction())
    assert rows[0] == ["8.1 暴露控制：", ""]
    assert rows[1] == ["呼吸系统防护：", "喷涂过程中要求有呼吸防护设备。"]
    assert rows[2] == ["建议：", "污染的手套应废弃。"]
    assert rows[-1] == ["8.2 工程控制：", "根据EC指令2006/121/EG,无可用的接触限值信息"]
    assert control_parameters == []


@pytest.mark.parametrize(
    "signal", ["无信号词", "无", "None", "No signal word", "Not applicable"]
)
def test_s2_extracts_controlled_non_hazard_signal_values(signal):
    from docx import Document

    document = Document()
    table = document.add_table(rows=1, cols=1)
    for text in ("2.1 GHS classification", f"Signal Word: {signal}"):
        row = table.add_row()
        row.cells[0].text = text
    data = extract_s2(table, Extraction())
    assert data["signal"] == signal


def test_s8_extracts_nested_control_parameter_records_and_coverage(tmp_path):
    from docx import Document

    document = Document()
    for index in range(16):
        table = document.add_table(rows=1, cols=2)
        table.rows[0].cells[0].text = f"Section {index + 1}"
        if index == 0:
            row = table.add_row()
            row.cells[0].text = "Product name"
            row.cells[1].text = "TEST-1234"
        if index == 7:
            row = table.add_row()
            row.cells[0].text = "工作场所组分控制参数"
            nested = row.cells[0].add_table(rows=2, cols=4)
            for cell, value in zip(nested.rows[0].cells, ("物质", "依据", "类型", "数值")):
                cell.text = value
            for cell, value in zip(
                nested.rows[1].cells,
                ("六亚甲基-1,6-二异氰酸酯", "CN OEL", "TWA", "0.03 mg/m3"),
            ):
                cell.text = value
    source = tmp_path / "TEST-1234.docx"
    document.save(source)

    data = _extract_docx(source)
    assert data["s8_control_parameters"]["zh"] == [[
        "六亚甲基-1,6-二异氰酸酯", "CN OEL", "TWA", "0.03 mg/m3"
    ]]
    assert data["source_coverage"]["unmapped"] == []
    assert data["source_coverage"]["counts"]["nested_table_count"] == 1
    assert data["source_coverage"]["counts"]["nested_source_unit_count"] == 8
    assert any(
        unit["text"] == "0.03 mg/m3" and "nested_table" in unit["source_locator"]
        for unit in data["source_coverage"]["source_units"]
    )


def test_s9_splits_inline_nco_content_into_a_dedicated_property_row():
    from docx import Document

    document = Document()
    table = document.add_table(rows=1, cols=2)
    row = table.add_row()
    row.cells[0].text = "其他信息："
    row.cells[1].text = "NCO含量：16.6±0.5%\n上述数据非产品指标，产品指标请参见产品技术信息表。"

    rows = extract_s9(table, Extraction())
    assert rows[0] == {"label": "NCO含量：", "value": "16.6±0.5%"}
    assert rows[1]["label"] == "其他信息："
    assert rows[1]["value"].startswith("上述数据非产品指标")


def test_s12_keeps_source_121_and_suppresses_only_template_notes():
    from docx import Document

    document = Document()
    table = document.add_table(rows=1, cols=2)
    for label, value in (
        ("生态毒性：", "该产品无可用的生态毒理学研究。禁止排入下水道，废水或土壤中。"),
        ("持久性和降解性：", "生物降解性：＜60%，28d。"),
        ("其他：", "无数据资料"),
    ):
        row = table.add_row()
        row.cells[0].text = label
        row.cells[1].text = value
    rows = extract_s12(table, Extraction())
    assert rows[:2] == [[""], [""]]
    assert rows[2][0].startswith("12.1")
    assert rows[3][0].startswith("12.2")
    assert rows[4][0].startswith("12.3")


def test_s2_extracts_precautionary_groups_and_does_not_leak_headings_to_other():
    from docx import Document

    document = Document()
    table = document.add_table(rows=1, cols=2)
    row = table.add_row()
    row.cells[0].text = "S2"
    row.cells[1].text = (
        "2.4 \u9632\u8303\u8bf4\u660e\uff1a\n"
        "\u9884\u9632\u63aa\u65bd\uff1a\n"
        "P210 \u8fdc\u79bb\u70ed\u6e90\u3002\n"
        "\u4e8b\u6545\u54cd\u5e94\uff1a\n"
        "P304+P340 \u5982\u679c\u5438\u5165\uff1a\u5c06\u4eba\u5458\u79fb\u5230\u7a7a\u6c14\u65b0\u9c9c\u5904\u3002\n"
        "\u5b89\u5168\u50a8\u5b58\uff1a\n"
        "\u5e9f\u5f03\u5904\u7f6e\uff1a\n"
        "P501 \u6309\u7167\u89c4\u5b9a\u5904\u7f6e\u3002"
    )

    data = extract_s2(table, Extraction())
    assert [group["group_key"] for group in data["precautionary_groups"]] == [
        "prevention", "response", "storage", "disposal",
    ]
    assert [len(group["statements"]) for group in data["precautionary_groups"]] == [1, 1, 0, 1]
    assert all("\u9884\u9632\u63aa\u65bd" not in line for line in data["other"])
    assert all("\u4e8b\u6545\u54cd\u5e94" not in line for line in data["other"])
    assert all("\u5b89\u5168\u50a8\u5b58" not in line for line in data["other"])
    assert all("\u5e9f\u5f03\u5904\u7f6e" not in line for line in data["other"])


def test_s2_does_not_route_precautionary_response_lines_as_health_hazards():
    from docx import Document

    document = Document()
    table = document.add_table(rows=1, cols=1)
    table.add_row().cells[0].text = (
        "2.6 防范说明：\n事故响应：\n吸入：一旦吸入，如有不适，就医。\n"
        "食入：立即漱口。\n2.7 健康危害：\n吸入危害：可能引起呼吸道刺激。\n"
        "皮肤刺激：可能造成皮肤刺激。"
    )
    data = extract_s2(table, Extraction())
    assert data["health_hazards"]["inhalation"] == ["可能引起呼吸道刺激。"]
    assert data["health_hazards"]["skin"] == ["可能造成皮肤刺激。"]
    response_text = [
        statement["text"]
        for statement in data["precautionary_groups"][0]["statements"]
    ]
    assert "吸入：一旦吸入，如有不适，就医。" in response_text
    assert "食入：立即漱口。" in response_text


def test_s2_extracts_inline_group_boundary_without_merging_heading_into_statement():
    from docx import Document

    document = Document()
    table = document.add_table(rows=1, cols=2)
    row = table.add_row()
    row.cells[0].text = "S2"
    row.cells[1].text = (
        "\u9884\u9632\u63aa\u65bd\uff1aP201 \u4f7f\u7528\u524d\u53d6\u5f97\u4e13\u95e8\u8bf4\u660e\u3002"
        "\u4e8b\u6545\u54cd\u5e94\uff1aP304+P340 \u5982\u679c\u5438\u5165\uff0c\u5c06\u4eba\u5458\u79fb\u5230\u7a7a\u6c14\u65b0\u9c9c\u5904\u3002"
    )
    data = extract_s2(table, Extraction())
    assert [group["group_key"] for group in data["precautionary_groups"]] == ["prevention", "response"]
    assert data["precautionary_groups"][0]["statements"][0]["code"] == "P201"
    assert data["precautionary_groups"][1]["statements"][0]["code"] == "P304+P340"


def test_s2_does_not_promote_heading_like_prose_outside_precautionary_context():
    from docx import Document

    document = Document()
    table = document.add_table(rows=1, cols=2)
    row = table.add_row()
    row.cells[0].text = "S2"
    row.cells[1].text = (
        "2.1 \u5e94\u6025\u60c5\u51b5\u6982\u8ff0\uff1a\u4e8b\u6545\u54cd\u5e94\uff1a\u6309\u7167\u5e94\u6025\u7a0b\u5e8f\u5904\u7406\u3002\n"
        "2.2 GHS \u5206\u7c7b\uff1a\u672a\u5206\u7c7b"
    )
    data = extract_s2(table, Extraction())
    assert data["precautionary_groups"] == []
    assert data["p_statements"] == []


def test_s2_records_orphan_group_for_review_but_projection_suppresses_it():
    from docx import Document

    document = Document()
    table = document.add_table(rows=1, cols=2)
    row = table.add_row()
    row.cells[0].text = "S2"
    row.cells[1].text = (
        "2.4 \u9632\u8303\u8bf4\u660e\uff1a\n"
        "\u5b89\u5168\u50a8\u5b58\uff1a"
    )
    ext = Extraction()
    data = extract_s2(table, ext)
    assert data["precautionary_groups"][0]["group_key"] == "storage"
    assert data["precautionary_groups"][0]["statements"] == []
    assert any(item["issue"] == "precautionary-group-orphan" for item in ext.review)


def test_source_mapping_draft_keeps_precautionary_group_as_one_semantic_fact():
    from extract_source_facts import source_mapping_draft

    heading = "\u9884\u9632\u63aa\u65bd\uff1a"
    statement = "P210 \u8fdc\u79bb\u70ed\u6e90\u3002"
    sections = {f"s{number}": {} for number in range(1, 17)}
    sections["s2"] = {
        "precautionary_groups": [{
            "group_key": "prevention",
            "source_heading": heading,
            "source_locator": "s2.line[2]",
            "statements": [{"code": "P210", "text": statement}],
        }],
    }
    coverage = {
        "source_units": [
            {"unit_id": "SRC-1", "source_section": "s2", "text": heading},
            {"unit_id": "SRC-2", "source_section": "s2", "text": statement},
        ]
    }
    result = source_mapping_draft(sections, [], [], coverage, "PU-1002", "hash")
    group_items = [item for item in result["items"] if item.get("source_kind") == "precautionary_group"]
    assert len(group_items) == 1
    assert group_items[0]["target_slot"] == "s2.precautionary_statements.group[prevention]"
    assert group_items[0]["source_unit_ids"] == ["SRC-1", "SRC-2"]
