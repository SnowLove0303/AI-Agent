import json
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1]


def test_snapshot_pins_current_template_geometry_and_parts():
    snapshot = json.loads((ROOT / "tests/template_snapshot.json").read_text(encoding="utf-8"))
    assert snapshot["source_sha256"] == "a69a447f7f39599b10c30c7f92c4d5101d94fd85bfdcc4d89de429d1c796af1c"
    assert len(snapshot["tables"]) == 16
    assert [t["row_count"] for t in snapshot["tables"]] == [10, 16, 6, 6, 5, 4, 3, 16, 24, 6, 18, 6, 3, 5, 9, 2]
    assert [t["column_count"] for t in snapshot["tables"]] == [2, 2, 3, 2, 2, 2, 2, 5, 2, 2, 4, 2, 2, 2, 1, 1]
    assert snapshot["sections"][0]["header"]["part_xml_hash"]
    assert snapshot["sections"][0]["footer"]["part_xml_hash"]


def test_snapshot_records_revised_sections_and_example_fact_boundary():
    snapshot = json.loads((ROOT / "tests/template_snapshot.json").read_text(encoding="utf-8"))
    section8 = "\n".join(c["text"] for r in snapshot["tables"][7]["rows"] for c in r["cells"])
    section11 = "\n".join(c["text"] for r in snapshot["tables"][10]["rows"] for c in r["cells"])
    assert "8.2" in section8
    assert "11.10" in section11
    assert "PEA-4139" in snapshot["sections"][0]["header"]["tables"][0]["rows"][0]["cells"][0]["text"]


def test_snapshot_pins_section8_2_top_level_control_parameter_rows():
    snapshot = json.loads((ROOT / "tests/template_snapshot.json").read_text(encoding="utf-8"))
    rows = snapshot["tables"][7]["rows"]
    assert len(rows) == 16
    assert rows[12]["cells"][0]["text"].strip() == "工作场所组分控制参数"
    assert [cell["text"] for cell in rows[13]["cells"]] == ["物质", "依据", "类型", "数值"]
    assert [cell["text"] for cell in rows[14]["cells"]] == ["六亚甲基-1,6-二异氰酸酯", "CN OEL", "TWA", "0.03 mg/m3"]


def test_snapshot_pins_current_row_page_settings():
    document = Document(ROOT / "examples/template_reference.docx")
    cant_split_counts = {
        table_index: sum(
            row._tr.trPr is not None
            and row._tr.trPr.find(qn("w:cantSplit")) is not None
            for row in table.rows
        )
        for table_index, table in enumerate(document.tables)
    }
    assert {index: count for index, count in cant_split_counts.items() if count} == {
        1: 1,
        5: 3,
        10: 16,
    }


def test_geometry_audit_script_is_packaged():
    assert (ROOT / "scripts/audit_template_geometry.py").exists()


def _border_value(cell, edge):
    borders = cell._tc.tcPr.find(qn("w:tcBorders"))
    element = borders.find(qn(f"w:{edge}")) if borders is not None else None
    return None if element is None else element.get(qn("w:val"))


def test_current_template_section8_and_section11_border_baseline():
    document = Document(ROOT / "examples/template_reference.docx")
    section8 = document.tables[7]
    assert _border_value(section8.rows[1].cells[0], "bottom") == "dotted"
    assert _border_value(section8.rows[2].cells[0], "top") == "dotted"
    assert _border_value(section8.rows[2].cells[0], "right") == "dotted"
    # Formal baseline: the second cell's left edge is single; an older nested
    # table is not an active template authority.
    assert _border_value(section8.rows[2].cells[1], "left") == "single"

    section11 = document.tables[10]
    assert _border_value(section11.rows[2].cells[0], "bottom") == "dotted"
    for cell in section11.rows[3].cells[:3]:
        assert _border_value(cell, "top") == "dotted"
