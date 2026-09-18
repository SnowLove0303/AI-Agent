from pathlib import Path

from docx import Document

ROOT = Path(__file__).resolve().parents[1]
import sys
sys.path.insert(0, str(ROOT / "scripts"))

from audit_english_terminology import run  # noqa: E402


def test_english_terminology_audit_scans_visible_docx_parts(tmp_path):
    path = tmp_path / "bad.docx"
    document = Document()
    document.add_paragraph("Product name: Resin树脂： 150℃")
    document.save(path)
    issues = run(path)
    assert any("CHINESE_REMAINDER" in issue and "word/document.xml" in issue for issue in issues)
    assert any("FULLWIDTH_COLON" in issue for issue in issues)
    assert any("INVALID_UNIT_TYPOGRAPHY" in issue for issue in issues)


def test_english_terminology_audit_accepts_canonical_visible_text(tmp_path):
    path = tmp_path / "good.docx"
    document = Document()
    document.add_paragraph("Product name: Waterborne polyurethane resin PU-1004")
    document.add_paragraph("Visual inspection; 150°C, 30 min (blast oven)")
    document.save(path)
    assert run(path) == []


def test_plain_text_audit_keeps_part_location(tmp_path):
    path = tmp_path / "bad.txt"
    path.write_text("标签：", encoding="utf-8")
    issues = run(path)
    assert any(issue.startswith("text: ") for issue in issues)
