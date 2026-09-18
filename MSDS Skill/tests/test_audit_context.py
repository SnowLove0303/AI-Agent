from pathlib import Path

from docx import Document

ROOT = Path(__file__).resolve().parents[1]
import sys
sys.path.insert(0, str(ROOT / "scripts"))

from audit_context import AuditContext  # noqa: E402
from audit_openspec_overwrite import audit as audit_openspec  # noqa: E402
from audit_template_mutation_whitelist import audit as audit_template  # noqa: E402
from audit_whitespace import run as audit_whitespace  # noqa: E402


TEMPLATE = ROOT / "examples" / "template_reference.docx"


def test_context_indexes_without_mutating_document():
    document = Document(str(TEMPLATE))
    before = document.element.body.xml
    table_before = tuple(table._tbl.xml for table in document.tables)

    context = AuditContext(document)

    assert document.element.body.xml == before
    assert tuple(table._tbl.xml for table in document.tables) == table_before
    assert context.summary()["table_count"] == 16
    assert context.summary()["row_count"] >= 16
    assert context.writable_boundaries()


def test_context_reuses_the_same_observations_as_direct_audits():
    direct_doc = Document(str(TEMPLATE))
    indexed_doc = Document(str(TEMPLATE))
    context = AuditContext(indexed_doc)

    assert audit_whitespace(str(TEMPLATE), document=direct_doc) == audit_whitespace(
        str(TEMPLATE), document=indexed_doc, context=context
    )

    direct_template = Document(str(TEMPLATE))
    direct_output = Document(str(TEMPLATE))
    indexed_template = Document(str(TEMPLATE))
    indexed_output = Document(str(TEMPLATE))
    indexed_context = AuditContext(indexed_output)
    direct_template_report = audit_template(
        TEMPLATE, TEMPLATE, template=direct_template, output=direct_output
    )
    indexed_template_report = audit_template(
        TEMPLATE, TEMPLATE, template=indexed_template, output=indexed_output,
        context=indexed_context,
    )
    assert indexed_template_report == direct_template_report

    direct_template = Document(str(TEMPLATE))
    direct_output = Document(str(TEMPLATE))
    indexed_template = Document(str(TEMPLATE))
    indexed_output = Document(str(TEMPLATE))
    indexed_context = AuditContext(indexed_output)
    direct_report = audit_openspec(
        TEMPLATE, TEMPLATE, template=direct_template, output=direct_output
    )
    indexed_report = audit_openspec(
        TEMPLATE, TEMPLATE, template=indexed_template, output=indexed_output,
        context=indexed_context,
    )
    assert indexed_report == direct_report


def test_context_is_invalidated_by_construction_only_and_does_not_offer_writes():
    document = Document(str(TEMPLATE))
    context = AuditContext(document)
    assert not any(name.startswith(("write", "set", "mutate")) for name in dir(context))
