from pathlib import Path
import sys
sys.path.insert(0, str(Path(__file__).resolve().parents[1] / 'scripts'))
from docx import Document
from renumber_visible_items import collect, audit, renumber

def make_doc(path):
    d=Document(); t=d.add_table(rows=6, cols=2)
    vals=['2.2  GHS危险性类别：','2.3  GHS标签要素：','11.1  急性毒性：','11.1  急性毒性：','11.2  主要皮肤刺激性：','11.3  主要眼睛刺激性：']
    for row,v in zip(t.rows,vals):
        # deliberately split prefix across runs, matching real template behavior
        p=row.cells[0].paragraphs[0]
        sec,item,rest=v.split('.',1)[0],v.split('.',1)[1].split()[0],v.split(None,1)[1]
        p.add_run(sec); p.add_run('.'); p.add_run(item+'  '); p.add_run(rest)
    d.save(path)

def test_gap_and_repeated_subrows(tmp_path):
    p=tmp_path/'a.docx'; make_doc(p)
    d=Document(p); assert audit(collect(d))
    renumber(d); out=tmp_path/'b.docx'; d.save(out)
    d2=Document(out); probs=audit(collect(d2)); assert not probs, probs
    texts=[x[3] for x in collect(d2)]
    assert texts[:2][0].startswith('2.1') and texts[:2][1].startswith('2.2')
    s11=[x for x in texts if x.startswith('11.')]
    assert s11[0].startswith('11.1') and s11[1].startswith('11.1')
    assert s11[2].startswith('11.2') and s11[3].startswith('11.3')

def test_numeric_property_value_is_not_a_numbered_label(tmp_path):
    p=tmp_path/'numeric.docx'; d=Document(); t=d.add_table(rows=1, cols=1)
    t.rows[0].cells[0].text='7.0-9.0'; d.save(p)
    assert collect(Document(p)) == []

def test_structured_endpoint_numbers_may_have_source_defined_gaps():
    d=Document(); t=d.add_table(rows=3, cols=1)
    for row, value in zip(t.rows, ['11.4  致敏性：', '11.6  致癌性：', '11.7  生殖毒性：']):
        row.cells[0].text=value
    assert not audit(collect(d))
