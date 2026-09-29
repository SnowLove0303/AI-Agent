from __future__ import annotations

import sys
from pathlib import Path

from docx import Document

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "scripts"))

from recognition_benchmark import (  # noqa: E402
    compare_docx,
    digest_json,
    discover,
    raw_docx_evidence,
    recognition_projection,
    sample_paths,
    update_regression_ledger,
)


def _make_docx(path: Path) -> None:
    document = Document()
    document.add_paragraph("1 Identification")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Product name:"
    table.cell(0, 1).text = "Sample resin"
    table.cell(1, 0).text = "Revision date:"
    table.cell(1, 1).text = "2026-09-29"
    document.save(path)


def test_raw_docx_evidence_and_compare(tmp_path):
    source = tmp_path / "sample.docx"
    _make_docx(source)
    document = __import__("msds_table_search").read_file(source)
    evidence = raw_docx_evidence(source)
    comparison = compare_docx(source, document)
    assert evidence["package_hash"]
    assert evidence["tables"]
    assert comparison["dimensions"]["content"]["status"] in {"PASS", "PARTIAL", "MISMATCH"}
    assert comparison["dimensions"]["images"]["status"] == "PASS"
    assert comparison["overall"] != "PASS" or comparison["dimensions"]["render"]["status"] == "UNAVAILABLE"


def test_dimension_hash_is_order_sensitive():
    assert digest_json([{"text": "a"}, {"text": "b"}]) != digest_json([{"text": "b"}, {"text": "a"}])


def test_sampling_is_repeatable_and_stratified(tmp_path):
    paths = []
    for index in range(10):
        suffix = ".docx" if index % 2 else ".pdf"
        path = tmp_path / f"PU-{index:04d}{suffix}"
        path.write_bytes(f"{index}".encode("ascii"))
        paths.append(path)
    first = sample_paths(paths, 6, 1234)
    second = sample_paths(paths, 6, 1234)
    assert first == second
    assert len(first) == 6


def test_discover_excludes_word_lock_files_and_samples_extensions(tmp_path):
    (tmp_path / "~$locked.docx").write_bytes(b"lock")
    (tmp_path / "PU-1001.docx").write_bytes(b"docx")
    (tmp_path / "PU-1002.doc").write_bytes(b"doc")
    (tmp_path / "PU-1003.pdf").write_bytes(b"pdf")
    paths = discover(tmp_path, {".docx", ".doc", ".pdf"})
    assert all(not path.name.startswith("~$") for path in paths)
    sampled = sample_paths(paths, 3, 7)
    assert {path.suffix for path in sampled} == {".docx", ".doc", ".pdf"}


def test_regression_ledger_classifies_fixed_and_regressed(tmp_path):
    ledger = tmp_path / "ledger.json"
    base = {
        "manifest": {"seed": 7},
        "comparator_version": "cmp",
        "results": [{"source": {"sha256": "abc"}, "dimensions": {"content": "MISMATCH", "images": "PASS"}}],
    }
    update_regression_ledger(ledger, base)
    later = {
        "manifest": {"seed": 7},
        "comparator_version": "cmp",
        "results": [{"source": {"sha256": "abc"}, "dimensions": {"content": "PASS", "images": "MISMATCH"}}],
    }
    changes = update_regression_ledger(ledger, later)
    assert {item["classification"] for item in changes} == {"fixed", "regressed"}


def test_projection_keeps_images_from_text_records():
    projection = recognition_projection({
        "records": [{
            "kind": "text",
            "label": "第 0 部分",
            "text": "Header",
            "content": [{"type": "image", "image": {"data": "aGVsbG8=", "search_text": "VML"}}],
        }]
    })
    assert projection["images"][0]["sha256"]
