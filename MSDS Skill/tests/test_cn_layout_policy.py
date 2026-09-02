from pathlib import Path
import sys

from docx import Document
from lxml import etree
from docx.oxml.ns import qn


ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "scripts"))
from compact_cn_layout import compact_cn_document
from compact_cn_layout import normalize_footer


def test_cn_compact_pass_preserves_locked_label_properties_and_normalizes_body():
    document = Document(ROOT / "examples" / "template_reference.docx")
    label = document.tables[0].rows[1].cells[0].paragraphs[0]
    label_before = etree.tostring(label._p.pPr, method="c14n")
    document.tables[0].rows[1].cells[1].paragraphs[0].add_run("value")

    compact_cn_document(document)

    assert etree.tostring(label._p.pPr, method="c14n") == label_before
    value = document.tables[0].rows[1].cells[1].paragraphs[0]
    assert value.paragraph_format.line_spacing == 1.0
    assert value.paragraph_format.space_before.pt == 0
    assert value.paragraph_format.space_after.pt == 0
    assert value.runs[0].font.size.pt == 10


def test_cn_compact_pass_clears_footer_indent_and_keeps_section_16_together():
    document = Document(ROOT / "examples" / "template_reference.docx")
    compact_cn_document(document)

    footer_cell = document.sections[0].footer.tables[0].rows[0].cells[1]
    paragraph = footer_cell.paragraphs[0]
    assert paragraph.alignment == 2
    assert paragraph.paragraph_format.left_indent.pt == 0
    assert paragraph.paragraph_format.right_indent.pt == 0
    assert paragraph.paragraph_format.first_line_indent is None or paragraph.paragraph_format.first_line_indent.pt == 0
    ind = paragraph._p.pPr.find(qn("w:ind"))
    assert ind is None or qn("w:firstLineChars") not in ind.attrib
    assert ind is None or qn("w:firstLine") not in ind.attrib

    tr_pr = document.tables[15].rows[1]._tr.get_or_add_trPr()
    assert tr_pr.find(qn("w:cantSplit")) is not None


def test_footer_normalization_is_language_neutral():
    document = Document(ROOT / "examples" / "template_reference.docx")
    normalize_footer(document)
    paragraph = document.sections[0].footer.tables[0].rows[0].cells[1].paragraphs[0]
    ind = paragraph._p.pPr.find(qn("w:ind"))
    assert ind is None or qn("w:firstLineChars") not in ind.attrib
    assert ind is None or qn("w:firstLine") not in ind.attrib
