"""Controlled mutation boundary for the unified MSDS templates.

The template is an executable document skeleton.  This module centralizes the
small set of mutations that generation is allowed to make:

* ordinary fields write only to value cells;
* S3 data rows write name/CAS/concentration to the three data cells;
* S8.2 top-level data rows write substance/basis/type/value to the four data
  cells, with the template example rows cleared per the data/placeholder rule;
* one-cell note slots may be replaced as a whole slot;
* S2/S9 may remove an explicitly missing row and change only the visible
  sequence prefix afterwards;
* Section 2.8 route prefixes and Section 11 middle sublabels are template
  content, not writable values;
* no operation may rewrite a sequence/label cell as a generic value write.

The functions intentionally operate on the existing OOXML nodes.  They do not
rebuild tables or normalize the formatting of locked cells.
"""

from __future__ import annotations

import copy
import re
from dataclasses import dataclass
from typing import Callable, Iterable, Sequence

from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from diagnostics import format_diagnostic


S82_TOP_HEADERS = {
    "zh": ("物质", "依据", "类型", "数值"),
    "en": ("Substance", "Basis", "Type", "Value"),
}

S82_CHILD_HEADERS = S82_TOP_HEADERS

S82_MISSING = {
    "zh": "无数据",
    "en": "No data available",
}

VALUE_FONT_BY_LANGUAGE = {"zh": "宋体", "en": "Times New Roman"}
VALUE_SIZE_PT = 12.0

# Blank cells are not automatically writable.  These are explicit semantic
# input exceptions; the S8 recommendation is source-gated and is written only
# when the source provides a non-empty recommendation.  A missing
# recommendation remains empty and is handled by the source-presence policy.
BLANK_VALUE_SLOTS = {(0, 1), (1, 1), (7, 8)}

# These rows are structural headings in the single-column Section 15 table,
# not ordinary note slots.  Their text and bold/template formatting are owned
# by the active template in both languages.
S15_LOCKED_HEADINGS = frozenset({
    "其它的规定:",
    "符合下列法规要求:",
    "Other provisions:",
    "Complies with the following regulations:",
})

# The maintained CN and EN templates currently use the Chinese route prefixes
# in Section 2.8.  English aliases are included because older customer
# templates may carry translated prefixes while retaining the same topology.
# The actual prefix text is always taken from the cloned template; these
# values are recognition keys only and are never written into the document.
S28_ROUTE_PREFIXES = frozenset({
    "吸入：", "食入：", "皮肤：", "眼睛：", "症状和体征：",
    "Inhalation:", "Ingestion:", "Skin:", "Eyes:",
    "Signs and symptoms:",
})
# Section 2.8 is a semantic health-hazard row, not a permanently literal
# number.  Authorized omission may renumber it to 2.7 (or another 2.x slot),
# so detection must follow the maintained label text after renumbering.
S28_LABEL_RE = re.compile(r"^\s*2\.\d+\s*(?:健康危害|health\s+hazards)\b", re.I)


class MutationViolation(RuntimeError):
    """Raised when a generator attempts to mutate outside the whitelist."""


@dataclass(frozen=True)
class TemplateSlot:
    table_index: int
    row_index: int
    cell_indices: tuple[int, ...]
    role: str
    authorized: bool
    # These fields make the registry an executable contract rather than a
    # positional convenience list.  They are deterministic for a fixed
    # template and intentionally do not contain process-local XML identities.
    locked_cell_indices: tuple[int, ...] = ()
    merge_identity: tuple[str, ...] = ()
    field_key: str = ""
    special_policy: str = ""


class TemplateSlotRegistry:
    """Runtime map of template-owned writable cells.

    It is built from the cloned template before clearing values.  Formatting,
    merges and labels remain in the DOCX; the registry only answers where data
    may go.
    """

    def __init__(self, slots: dict[tuple[int, int], TemplateSlot]):
        self.slots = slots

    @classmethod
    def from_document(cls, document) -> "TemplateSlotRegistry":
        slots = {}
        for table_index, table in enumerate(document.tables):
            for row_index, row in enumerate(table.rows):
                cells = unique_cells(row)
                if not cells or row_index == 0:
                    continue
                if table_index == 14 and is_s15_locked_heading_row(row):
                    continue
                if table_index == 2 and row_index >= 4:
                    slots[table_index, row_index] = TemplateSlot(
                        table_index, row_index, tuple(range(len(cells))), "s3_data", True,
                        locked_cell_indices=(),
                        merge_identity=_merge_identity(cells),
                        field_key=_slot_field_key(table_index, row_index, cells, "s3_data"),
                        special_policy="three_column_data",
                    )
                    continue
                if table_index == 2 and row_index in {2, 3}:
                    continue
                if table_index == 7 and row_index >= 14:
                    slots[table_index, row_index] = TemplateSlot(
                        table_index, row_index, tuple(range(len(cells))), "s82_data", True,
                        locked_cell_indices=(),
                        merge_identity=_merge_identity(cells),
                        field_key=_slot_field_key(table_index, row_index, cells, "s82_data"),
                        special_policy="four_column_data",
                    )
                    continue
                if table_index == 7 and row_index == 1:
                    continue
                if table_index == 7 and row_index in {12, 13}:
                    continue

                # The formal Section 2.8 rows are a composite value slot in
                # the current two-column template: the second physical cell
                # already owns the route prefix (吸入：/食入：/...).  In an
                # alternate three-column layout the middle cell owns that
                # prefix and only the final cell is writable.
                s28 = _s28_parts(cells) if table_index == 1 else None
                if s28 is not None:
                    _, value_index, _ = s28
                    locked = tuple(index for index in range(len(cells))
                                   if index != value_index)
                    slots[table_index, row_index] = TemplateSlot(
                        table_index, row_index, (value_index,),
                        "s2_composite_value", True,
                        locked_cell_indices=locked,
                        merge_identity=_merge_identity(cells),
                        field_key=_slot_field_key(
                            table_index, row_index, cells, "s2_composite_value"
                        ),
                        special_policy="s2_route_prefix",
                    )
                    continue

                if (
                    table_index == 10
                    and len(cells) >= 3
                    and re.match(r"^\s*11\.(?:1|2|7)\b", cells[0].text)
                ):
                    slots[table_index, row_index] = TemplateSlot(
                        table_index, row_index, (len(cells) - 1,), "endpoint_value", True,
                        locked_cell_indices=tuple(range(len(cells) - 1)),
                        merge_identity=_merge_identity(cells),
                        field_key=_slot_field_key(table_index, row_index, cells, "endpoint_value"),
                        special_policy="structured_endpoint",
                    )
                    continue
                if len(cells) == 1:
                    slots[table_index, row_index] = TemplateSlot(
                        table_index, row_index, (0,), "note", True,
                        locked_cell_indices=(),
                        merge_identity=_merge_identity(cells),
                        field_key=_slot_field_key(table_index, row_index, cells, "note"),
                        special_policy="whole_note_slot",
                    )
                    continue
                payload = tuple(index for index, cell in enumerate(cells[1:], 1)
                                if cell.text.strip())
                if not payload and (table_index, row_index) in BLANK_VALUE_SLOTS:
                    payload = tuple(range(1, len(cells)))
                slots[table_index, row_index] = TemplateSlot(
                    table_index, row_index, payload, "field", bool(payload),
                    locked_cell_indices=(0,),
                    merge_identity=_merge_identity(cells),
                    field_key=_slot_field_key(table_index, row_index, cells, "field"),
                    special_policy="ordinary_value" if payload else "blank_template_slot",
                )
        return cls(slots)

    def writable_cells(self, table_index: int, row_index: int, row) -> list:
        slot = self.slots.get((table_index, row_index))
        if not slot or not slot.authorized:
            return []
        cells = unique_cells(row)
        return [cells[index] for index in slot.cell_indices if index < len(cells)]

    def is_authorized(self, table_index: int, row_index: int) -> bool:
        slot = self.slots.get((table_index, row_index))
        return bool(slot and slot.authorized and slot.cell_indices)


def unique_cells(row) -> list:
    """Return physical cells once, including stable handling of merged cells."""
    seen: set[int] = set()
    result = []
    for cell in row.cells:
        key = hash(cell._tc)
        if key not in seen:
            seen.add(key)
            result.append(cell)
    return result


def _merge_identity(cells: Sequence) -> tuple[str, ...]:
    """Return deterministic physical-cell merge/topology metadata."""
    identities = []
    for cell in cells:
        tc_pr = cell._tc.tcPr
        if tc_pr is None:
            identities.append("")
            continue
        parts = []
        for tag in ("gridSpan", "vMerge", "hMerge"):
            node = tc_pr.find(qn(f"w:{tag}"))
            if node is not None:
                parts.append(f"{tag}={node.get(qn('w:val')) or 'continue'}")
        identities.append(";".join(parts))
    return tuple(identities)


def _slot_field_key(table_index: int, row_index: int, cells: Sequence,
                    role: str) -> str:
    """Build a stable registry key from the template's semantic label."""
    label = re.sub(r"\s+", " ", cells[0].text.strip()) if cells else ""
    return f"t{table_index + 1}.r{row_index + 1}:{role}:{label}"


def _is_s15_locked_heading_cells(cells: Sequence) -> bool:
    if len(cells) != 1:
        return False
    text = re.sub(r"\s+", " ", cells[0].text.strip())
    text = text.replace("：", ":")
    return text.casefold() in {value.casefold() for value in S15_LOCKED_HEADINGS}


def is_s15_locked_heading_row(row) -> bool:
    """Return whether a Section 15 row is a template-owned bold heading."""
    return _is_s15_locked_heading_cells(unique_cells(row))


def _route_prefix_in_text(text: str) -> str:
    """Return the exact template-owned route prefix at the start of text."""
    text = str(text or "")
    leading = len(text) - len(text.lstrip())
    remainder = text[leading:]
    for prefix in sorted(S28_ROUTE_PREFIXES, key=len, reverse=True):
        if remainder.startswith(prefix):
            return text[:leading] + prefix
    return ""


def _s28_parts(cells: Sequence) -> tuple[int | None, int, str] | None:
    """Return ``(prefix_cell, value_cell, exact_prefix)`` for an S2.8 row."""
    if not cells or not S28_LABEL_RE.match(cells[0].text or ""):
        return None
    if len(cells) >= 3:
        prefix = _route_prefix_in_text(cells[1].text)
        if prefix and cells[1].text.strip() == prefix.strip():
            return 1, len(cells) - 1, prefix
    if len(cells) >= 2:
        prefix = _route_prefix_in_text(cells[1].text)
        if prefix:
            return None, 1, prefix
    return None


def is_s28_row(table_index: int, row) -> bool:
    """Return whether a physical row is the maintained Section 2.8 shape."""
    return table_index == 1 and _s28_parts(unique_cells(row)) is not None


def composite_value_text(row) -> str:
    """Return only the writable value tail of a Section 2.8 row.

    The route prefix is template-owned.  This helper is shared by suppression
    and audits so a prefix-only row is treated as an empty value row.
    """
    cells = unique_cells(row)
    parts = _s28_parts(cells)
    if parts is None:
        return ""
    prefix_cell_index, value_index, prefix = parts
    value = cells[value_index].text or ""
    if prefix_cell_index is None:
        if value.startswith(prefix):
            value = value[len(prefix):]
    return value.lstrip(" \t\r\n")


def composite_prefix_text(row) -> str:
    """Return the exact template-owned prefix for a Section 2.8 row."""
    parts = _s28_parts(unique_cells(row))
    return parts[2] if parts is not None else ""


def composite_value_cell_index(row) -> int | None:
    """Return the physical value-cell index for a Section 2.8 row."""
    parts = _s28_parts(unique_cells(row))
    return parts[1] if parts is not None else None


def english_body_cells(table_index: int, row_index: int, row) -> list:
    """Return the EN cells whose non-bold value runs use the body exemplar.

    The first cell is normally a locked sequence/label cell.  S3 and S8.2
    data rows are explicit all-value subtables; Section 11 keeps its middle
    sublabel locked and its final cell as the value cell.
    """
    cells = unique_cells(row)
    if not cells or row_index == 0:
        return []
    if table_index == 14 and is_s15_locked_heading_row(row):
        return []
    if table_index == 2:
        if row_index in {2, 3}:
            return []
        if row_index >= 4:
            return cells
    if table_index == 7:
        if row_index in {1, 12, 13}:
            return []
        if row_index >= 14:
            return cells
    if table_index == 1 and _s28_parts(cells) is not None:
        # The current EN template also carries the route prefix in the
        # composite second cell.  Its value tail is written with the cell's
        # own approved run anchor; treating the entire cell as an ordinary EN
        # body cell would reformat the locked prefix run.
        return []
    if len(cells) == 1:
        return cells
    if table_index == 10 and len(cells) >= 3:
        return cells[-1:]
    return cells[1:]


def _clear_paragraph_content(paragraph) -> None:
    for child in list(paragraph._p):
        if child.tag != qn("w:pPr"):
            paragraph._p.remove(child)


def normalize_value_text(text: str) -> str:
    """Remove synthetic slash separators before writing a value cell.

    `` / `` is a legacy joiner, not source content.  It can wrap as a lone
    slash in Word.  Preserve compact source slashes such as ``通风/排气`` and
    ``有/无``; only the synthetic spaced separator becomes a semantic line
    break.  Empty lines and a slash-only line are never customer-facing.
    """
    normalized = re.sub(r"[ \t\u3000]+[/／][ \t\u3000]+", "\n", str(text or ""))
    normalized = re.sub(r"(?m)^[ \t]*[/／][ \t]*", "", normalized)
    return "\n".join(
        line.rstrip() for line in normalized.splitlines()
        if line.strip() and line.strip() not in {"/", "／"}
    )


def _set_value_rpr(r_pr, language: str | None) -> None:
    """Force approved value font/size without touching template labels."""
    if language not in VALUE_FONT_BY_LANGUAGE:
        return
    font = VALUE_FONT_BY_LANGUAGE[language]
    r_fonts = r_pr.find(qn("w:rFonts"))
    if r_fonts is None:
        r_fonts = OxmlElement("w:rFonts")
        r_pr.insert(0, r_fonts)
    for key in ("ascii", "hAnsi", "eastAsia", "cs"):
        r_fonts.set(qn(f"w:{key}"), font)
    for tag in ("w:sz", "w:szCs"):
        node = r_pr.find(qn(tag))
        if node is None:
            node = OxmlElement(tag)
            r_pr.append(node)
        node.set(qn("w:val"), str(int(VALUE_SIZE_PT * 2)))


def _set_value_cell_layout(cell, language: str | None, *, alignment: int = 0) -> None:
    """Apply value-only alignment/vertical rules; labels never call this."""
    if language not in VALUE_FONT_BY_LANGUAGE:
        return
    tc_pr = cell._tc.get_or_add_tcPr()
    v_align = tc_pr.find(qn("w:vAlign"))
    if v_align is None:
        v_align = OxmlElement("w:vAlign")
        tc_pr.append(v_align)
    v_align.set(qn("w:val"), "center")
    for paragraph in cell.paragraphs:
        paragraph.alignment = alignment
        for run in paragraph.runs:
            if not run.bold:
                _set_value_rpr(run._r.get_or_add_rPr(), language)


def enforce_value_typography(document, language: str,
                             registry: TemplateSlotRegistry | None = None) -> None:
    """Final S1-S16 value-only typography gate; locked bold runs are skipped."""
    if language not in VALUE_FONT_BY_LANGUAGE:
        raise MutationViolation(f"unsupported value typography language: {language}")
    registry = registry or TemplateSlotRegistry.from_document(document)
    for table_index, table in enumerate(document.tables[:16]):
        for row_index, row in enumerate(table.rows):
            for _cell_index, cell, locked_prefix in _writable_value_cells(
                table_index, row_index, row
            ):
                # Section 3's three data columns are explicitly centered by
                # both maintained templates; other value slots are left-aligned.
                alignment = 1 if table_index == 2 and row_index >= 4 else 0
                _set_value_cell_layout(cell, language, alignment=alignment)


def _write_paragraph_content(paragraph, text: str, *, force_nonbold: bool = True,
                            language: str | None = None) -> None:
    """Replace a writable value while inheriting template style without bold.

    The maintained templates use bold for labels and fixed structure.  A
    writable value may inherit the template value anchor's font, size, color,
    language and other character properties, but it can never inherit bold.
    ``force_nonbold=False`` is reserved for the one controlled insertion path
    that seeds a newly cloned S9 label; ordinary value callers must use the
    default.
    """
    text = normalize_value_text(text)
    old_rpr = None
    for run in paragraph.runs:
        if run._r.rPr is not None:
            old_rpr = copy.deepcopy(run._r.rPr)
            break
    if old_rpr is None and paragraph._p.pPr is not None:
        template_rpr = paragraph._p.pPr.find(qn("w:rPr"))
        if template_rpr is not None:
            old_rpr = copy.deepcopy(template_rpr)
    if force_nonbold:
        paragraph_rpr = (
            paragraph._p.pPr.find(qn("w:rPr"))
            if paragraph._p.pPr is not None else None
        )
        inherited_bold = _element_has_bold(paragraph_rpr)
        if old_rpr is None and inherited_bold:
            old_rpr = paragraph._p.makeelement(qn("w:rPr"), {})
        if old_rpr is not None:
            # Bold is a template-owned label/structure property, never a
            # value property.  Remove direct run bold from the newly created
            # value run.  If the value paragraph itself inherits bold, add an
            # explicit off override to the value run without touching shared
            # paragraph/label properties.
            _set_rpr_nonbold(old_rpr, override_inherited=inherited_bold)
    _clear_paragraph_content(paragraph)
    run = paragraph._p.makeelement(qn("w:r"), {})
    if old_rpr is not None:
        run.append(old_rpr)
    _set_value_rpr(run.get_or_add_rPr(), language)
    for index, line in enumerate(str(text).split("\n")):
        if index:
            run.append(paragraph._p.makeelement(qn("w:br"), {}))
        text_node = paragraph._p.makeelement(qn("w:t"), {})
        text_node.text = line
        if line and (line[0].isspace() or line[-1].isspace()):
            text_node.set(qn("xml:space"), "preserve")
        run.append(text_node)
    paragraph._p.append(run)


def _replace_leading_pattern_in_runs(paragraph, pattern: str, replacement: str) -> str:
    """Replace a leading pattern without collapsing the existing run tree."""
    current = paragraph.text
    match = re.match(pattern, current)
    if not match:
        return current
    old_prefix = match.group(0)
    new_prefix = re.sub(pattern, replacement, old_prefix, count=1)
    start, end = match.span()
    spans = []
    cursor = 0
    for run in paragraph.runs:
        spans.append((run, cursor, cursor + len(run.text)))
        cursor += len(run.text)
    overlapping = [item for item in spans if item[2] > start and item[1] < end]
    if not overlapping:
        return current
    first_run, first_start, _ = overlapping[0]
    last_run, _, last_end = overlapping[-1]
    if first_run is last_run:
        local_start = start - first_start
        local_end = end - first_start
        first_run.text = first_run.text[:local_start] + new_prefix + first_run.text[local_end:]
    else:
        first_local_start = start - first_start
        first_run.text = first_run.text[:first_local_start] + new_prefix
        for run, _, _ in overlapping[1:-1]:
            run.text = ""
        last_local_end = end - (last_end - len(last_run.text))
        last_run.text = last_run.text[last_local_end:]
    return paragraph.text


def set_value_cell_text(cell, text: str, *, force_nonbold: bool = True,
                        language: str | None = None, alignment: int = 0) -> None:
    """Write only a pre-authorized value/note cell as non-bold text.

    All writable value slots use the template's existing character anchor
    except that ``w:b`` and ``w:bCs`` are always removed.  This keeps value
    typography inherited and prevents a bold placeholder or paragraph mark
    from promoting source content into a label.
    """
    if not cell.paragraphs:
        raise MutationViolation("value cell has no template paragraph")
    _write_paragraph_content(
        cell.paragraphs[0], text, force_nonbold=force_nonbold, language=language,
    )
    for paragraph in cell.paragraphs[1:]:
        paragraph._element.getparent().remove(paragraph._element)
    _set_value_cell_layout(cell, language, alignment=alignment)


def _append_text_nodes(parent, text: str) -> None:
    """Append Word text/break children without changing the parent properties."""
    for index, line in enumerate(str(text).split("\n")):
        if index:
            parent.append(parent.makeelement(qn("w:br"), {}))
        text_node = parent.makeelement(qn("w:t"), {})
        text_node.text = line
        if line and (line[0].isspace() or line[-1].isspace()):
            text_node.set(qn("xml:space"), "preserve")
        parent.append(text_node)


def _run_with_text(paragraph, text: str, r_pr=None):
    run = paragraph._p.makeelement(qn("w:r"), {})
    if r_pr is not None:
        run.append(copy.deepcopy(r_pr))
    _append_text_nodes(run, text)
    return run


def _run_with_break(paragraph, r_pr=None):
    run = paragraph._p.makeelement(qn("w:r"), {})
    if r_pr is not None:
        run.append(copy.deepcopy(r_pr))
    run.append(run.makeelement(qn("w:br"), {}))
    return run


def _prefix_run_clones(paragraph, prefix: str) -> list:
    """Clone only the original run fragments that contain a locked prefix."""
    clones = []
    cursor = 0
    for source_run in paragraph.runs:
        source_text = source_run.text or ""
        if not source_text or cursor >= len(prefix):
            continue
        take = prefix[cursor:cursor + len(source_text)]
        if not take:
            continue
        clone = copy.deepcopy(source_run._r)
        for child in list(clone):
            if child.tag != qn("w:rPr"):
                clone.remove(child)
        _append_text_nodes(clone, take)
        clones.append(clone)
        cursor += len(take)
    if cursor != len(prefix):
        raise MutationViolation(
            "Section 2.8 route prefix is not represented by the template run tree"
        )
    return clones


def _value_rpr_for_composite(paragraph, prefix_clones: Sequence):
    """Use the composite cell's template anchor for a non-bold value tail."""
    r_pr = None
    for clone in prefix_clones:
        r_pr = clone.find(qn("w:rPr"))
        if r_pr is not None:
            break
    if r_pr is None and paragraph._p.pPr is not None:
        r_pr = paragraph._p.pPr.find(qn("w:rPr"))
    if r_pr is None:
        return None
    r_pr = copy.deepcopy(r_pr)
    # A route prefix may be bold in a customer-maintained template.  The
    # prefix remains bold/unchanged, but the source value tail is not promoted
    # to a label merely because it shares a physical cell.
    paragraph_rpr = (
        paragraph._p.pPr.find(qn("w:rPr"))
        if paragraph._p.pPr is not None else None
    )
    _set_rpr_nonbold(r_pr, override_inherited=_element_has_bold(paragraph_rpr))
    return r_pr


def set_s28_composite_value_cell(cell, text: str, language: str | None = None) -> None:
    """Write only the value tail while preserving the Section 2.8 prefix.

    This is intentionally not implemented with ``cell.text`` or the generic
    value writer: those operations erase the template-owned route prefix and
    its run formatting.  The prefix runs are cloned verbatim; only a new,
    non-bold value run is appended after a semantic line break.
    """
    if not cell.paragraphs:
        raise MutationViolation("Section 2.8 composite cell has no template paragraph")
    paragraph = cell.paragraphs[0]
    prefix = _route_prefix_in_text(paragraph.text)
    if not prefix:
        raise MutationViolation("Section 2.8 composite cell has no recognized route prefix")
    prefix_clones = _prefix_run_clones(paragraph, prefix)
    value = normalize_value_text(text)
    _clear_paragraph_content(paragraph)
    for clone in prefix_clones:
        paragraph._p.append(clone)
    if value:
        separator_rpr = _value_rpr_for_composite(paragraph, prefix_clones)
        _set_value_rpr(separator_rpr, language)
        paragraph._p.append(_run_with_break(paragraph, separator_rpr))
        paragraph._p.append(_run_with_text(paragraph, value, separator_rpr))
    for extra in cell.paragraphs[1:]:
        extra._element.getparent().remove(extra._element)
    _set_value_cell_layout(cell, language)


def _s28_value_without_prefix(value: object, prefix: str) -> str:
    """Accept legacy ``prefix + value`` payloads without duplicating prefix."""
    value = str(value or "")
    leading = len(value) - len(value.lstrip())
    candidate = value[leading:]
    bare_prefix = prefix.strip()
    if candidate.startswith(bare_prefix):
        return candidate[len(bare_prefix):].lstrip(" \t\r\n")
    return value


def set_sequence_prefix(cell, section: int, item: int,
                        *, prefix_width: int | None = None) -> str:
    """Change only an approved visible sequence prefix in an existing cell.

    ``prefix_width`` is used by Section 9, whose maintained template reserves
    five character positions for ``9.n`` plus separator spaces.  When a
    two-digit source label is moved into a one-digit visible slot, its original
    single separator must become two separators; otherwise the label column
    develops a visible one-space zig-zag.  No non-prefix label text is touched.
    """
    if not cell.paragraphs:
        return ""
    paragraph = cell.paragraphs[0]
    current = paragraph.text
    pattern = rf"^([ \t]*){section}\.\d+([ \t]*)"
    if prefix_width is None:
        replacement = rf"\g<1>{section}.{item}\g<2>"
    else:
        separator_width = max(1, prefix_width - len(f"{section}.{item}"))
        replacement = rf"\g<1>{section}.{item}{' ' * separator_width}"
    return _replace_leading_pattern_in_runs(
        paragraph, pattern, replacement
    )


def _payload_for_field_row(cells: list, values: Sequence[object]) -> list[str]:
    values = [str(value) for value in values]
    if len(cells) == 2 and len(values) > 2:
        return ["\n".join(value for value in values[1:] if value.strip())]
    return values[1:]


def _single_column_payload(values: Sequence[object]) -> str:
    """Return only the value part for S15/S16 whole-cell rows.

    Older semantic drafts represented a one-cell legal/disclaimer row as
    ``[label, value]`` even though the formal template has no separate value
    column. Joining both fields produced ``label label value`` in output. The
    template still owns the cell topology; this helper only normalizes the
    intermediate payload before the authorized whole-slot write.
    """
    items = [str(value or "").strip() for value in values if str(value or "").strip()]
    if not items:
        return ""
    if len(items) == 1:
        return items[0]
    label = items[0].rstrip("：:").strip()
    tail = "\n".join(items[1:]).strip()
    if not tail:
        return label
    normalized_label = re.sub(r"\s+", "", label).casefold()
    normalized_tail = re.sub(r"\s+", "", tail).casefold()
    if normalized_label and normalized_tail.startswith(normalized_label):
        tail = tail[len(label):].lstrip(" \t\r\n：:;-–—")
    return tail or label


def write_row_values(row, values: Sequence[object], *, table_index: int | None = None,
                     row_index: int | None = None,
                     registry: TemplateSlotRegistry | None = None,
                     inserted_data_row: bool = False,
                     language: str | None = None) -> dict | None:
    """Write a semantic row through the mutation whitelist.

    ``values`` keeps the historical source-fact shape, where the first item is
    a source label.  For normal field rows that first item is deliberately
    ignored: the label already belongs to the template.  Subtable data rows
    are the explicit exception because every cell is data, not a label.
    """
    cells = unique_cells(row)
    if not cells:
        raise MutationViolation("attempted to write a row without cells")
    values = [str(value) for value in values]

    if table_index == 14 and is_s15_locked_heading_row(row):
        # Single-cell Section 15 headings are structural labels, not note
        # payloads.  A generic writer must never clear, replace, or de-bold
        # them; the locked-skeleton audit protects the persisted result.
        if any(value.strip() for value in values[1:]):
            raise MutationViolation("Section 15 structural heading is not writable")
        return None

    if inserted_data_row:
        if table_index not in {8, 14}:
            raise MutationViolation(
                "inserted source rows are restricted to S9 physical/chemical properties or S15 regulations"
            )
        if not values or not values[0].strip():
            raise MutationViolation("an inserted source row requires source-backed row content")
        if len(cells) > 1:
            # A newly cloned S9 row has no pre-existing locked label. Seed it
            # from the verified source row while retaining the cloned label
            # formatting. Existing template labels never use this path.
            set_value_cell_text(cells[0], values[0], force_nonbold=False)

    if table_index == 2 and row_index in {2, 3}:
        # S3 parent row and the three-column table header are template-owned.
        return

    if table_index == 7 and row_index == 1:
        # Feishu S8.1 is a section parent node.  The supplied template uses a
        # single merged cell here, so it must not be mistaken for a writable
        # one-cell note slot.
        return

    if table_index == 7 and row_index is not None and row_index >= 12:
        # The formal S8.2 parent/header/data rows are owned by
        # write_s82_top_rows; generic row writes must not touch them.
        raise MutationViolation("S8.2 top-level rows must be written through write_s82_top_rows")

    if table_index == 7 and row_index == 8:
        # Section 8 recommendation is an approved source-gated value slot.
        # Its template label remains locked/bold, while its value must use
        # ordinary body weight even though the blank template paragraph
        # inherits bold paragraph-mark properties.
        targets = registry.writable_cells(table_index, row_index, row) \
            if registry else cells[1:]
        if not targets:
            if values and values[-1].strip():
                return {"status": "skipped_blank_template_slot",
                        "table_index": table_index, "row_index": row_index}
            return None
        set_value_cell_text(
            targets[0], values[-1] if values else "", force_nonbold=True,
            language=language,
        )
        return None

    # Section 2.8 is a composite template slot.  The route prefix is owned
    # by the template and the source payload is only the value tail.
    s28 = _s28_parts(cells) if table_index == 1 else None
    if s28 is not None:
        prefix_cell_index, value_index, prefix = s28
        targets = registry.writable_cells(table_index, row_index, row) \
            if registry else [cells[value_index]]
        if not targets:
            if values and values[-1].strip():
                return {"status": "skipped_blank_template_slot",
                        "table_index": table_index, "row_index": row_index}
            return None
        value = values[-1] if values else ""
        value = _s28_value_without_prefix(value, prefix)
        target = targets[0]
        if prefix_cell_index is None:
            set_s28_composite_value_cell(target, value, language=language)
        else:
            # Alternate three-column layouts have a dedicated locked middle
            # cell and an ordinary final value cell.
            set_value_cell_text(target, value, language=language)
        return None

    # S3 component data rows: name, CAS, concentration are all writable.
    if table_index == 2 and row_index is not None and row_index >= 4:
        if len(cells) != 3 or len(values) < 3:
            raise MutationViolation("S3 data row must contain exactly name/CAS/concentration")
        for cell, value in zip(cells, values[:3]):
            set_value_cell_text(cell, value, language=language, alignment=1)
        return

    # One-cell fixed semantic note slots are expressly writable as a slot.
    if len(cells) == 1:
        targets = registry.writable_cells(table_index, row_index, row) if registry else cells
        if targets:
            if table_index in {14, 15}:
                content = _single_column_payload(values)
            else:
                content = "\n".join(values[:-1]) if len(values) > 1 and not values[-1].strip() else "\n".join(values)
            set_value_cell_text(targets[0], content, language=language)
        return

    payload = _payload_for_field_row(cells, values)
    if (
        table_index == 10
        and len(cells) >= 3
        and re.match(r"^\s*11\.(?:1|2|7)\b", cells[0].text)
    ):
        targets = registry.writable_cells(table_index, row_index, row) if registry else [cells[-1]]
        if targets:
            set_value_cell_text(targets[-1], values[-1] if values else "", language=language)
        return None
    if registry is not None:
        targets = registry.writable_cells(table_index, row_index, row)
        if not targets:
            if any(value.strip() for value in payload):
                return {"status": "skipped_blank_template_slot",
                        "table_index": table_index, "row_index": row_index}
            return None
        for cell, value in zip(targets, payload):
            set_value_cell_text(cell, value, language=language)
        return None
    for offset, value in enumerate(payload, start=1):
        if offset >= len(cells):
            break
        set_value_cell_text(cells[offset], value, language=language)


def write_s82_top_rows(table, records: Iterable[Sequence[object]], language: str) -> dict:
    """Write verified S8.2 control-parameter records into the formal top-level rows.

    The one-cell parent row (physical row 12) and the four-column header row
    (physical row 13) belong to the template.  Data rows (physical row 14 on)
    are writable: extra rows are cloned from the existing styled data row, and
    the template example rows are cleared.  With no verified records the
    entire workplace-component block is removed; an empty source slot is not
    customer-facing content and must not become a synthetic placeholder row.
    """
    if language not in S82_TOP_HEADERS:
        raise MutationViolation(f"unsupported S8.2 language: {language}")
    normalized = [tuple(str(value) for value in record[:4]) for record in records]
    if any(len(record) != 4 for record in normalized):
        raise MutationViolation("each S8.2 record must contain substance/basis/type/value")
    if not normalized:
        while len(table.rows) > 12:
            table._tbl.remove(table.rows[-1]._tr)
        return {
            "record_count": 0,
            "row_count": 0,
            "placeholder": False,
            "hidden": True,
        }
    if len(table.rows) < 16:
        raise MutationViolation("S8.2 formal layout requires 16 Section-8 rows")
    parent = unique_cells(table.rows[12])
    if len(parent) != 1:
        raise MutationViolation("S8.2 parent must be a one-cell row")
    header = tuple(cell.text.strip() for cell in unique_cells(table.rows[13]))
    if header != S82_TOP_HEADERS[language]:
        raise MutationViolation(
            f"S8.2 header mismatch: expected {S82_TOP_HEADERS[language]}, found {header}"
        )
    while len(table.rows) < 14 + len(normalized):
        table._tbl.append(copy.deepcopy(table.rows[14]._tr))
    while len(table.rows) > 14 + len(normalized):
        table._tbl.remove(table.rows[-1]._tr)
    for position, row in enumerate(list(table.rows)[14:14 + len(normalized)]):
        cells = unique_cells(row)
        if len(cells) != 4:
            raise MutationViolation("S8.2 data row must have four cells")
        for target, value in zip(cells, normalized[position]):
            set_value_cell_text(target, value, language=language)
    return {
        "record_count": len(records),
        "row_count": len(normalized),
        "placeholder": False,
    }


def clear_value_cells(document, registry: TemplateSlotRegistry | None = None,
                      language: str | None = None) -> None:
    """Clear only writable value cells, leaving all labels and headings intact."""
    for table_index, table in enumerate(document.tables):
        for row_index, row in enumerate(table.rows):
            if row_index == 0:
                continue
            cells = unique_cells(row)
            if not cells:
                continue
            if table_index == 14 and is_s15_locked_heading_row(row):
                continue
            if registry is not None:
                slot = registry.slots.get((table_index, row_index))
                targets = registry.writable_cells(table_index, row_index, row)
                if slot and slot.special_policy == "s2_route_prefix" and targets:
                    if len(cells) == 2:
                        set_s28_composite_value_cell(targets[0], "", language=language)
                    else:
                        set_value_cell_text(targets[0], "", language=language)
                else:
                    for cell in targets:
                        set_value_cell_text(cell, "", language=language)
                continue
            if table_index == 2 and row_index in {2, 3}:
                # S3 parent row and three-column table header are locked.
                continue
            if table_index == 7 and row_index == 1:
                # S8.1 is a locked parent node, not a value slot.
                continue
            if table_index == 7 and row_index in {12, 13}:
                # S8.2 parent row and four-column header are locked.
                continue
            if table_index == 7 and row_index >= 14:
                for cell in cells:
                    set_value_cell_text(cell, "")
                continue
            if table_index == 2 and row_index >= 4:
                for cell in cells:
                    set_value_cell_text(cell, "")
                continue
            if table_index == 1 and _s28_parts(cells) is not None:
                parts = _s28_parts(cells)
                target = cells[parts[1]]
                if parts[0] is None:
                    set_s28_composite_value_cell(target, "")
                else:
                    set_value_cell_text(target, "")
                continue
            if len(cells) == 1:
                set_value_cell_text(cells[0], "")
            else:
                for cell in cells[1:]:
                    set_value_cell_text(cell, "")


@dataclass(frozen=True)
class LockedCellSnapshot:
    table_index: int
    row_key: tuple[str, int]
    cell_role: str
    tc_pr: str
    paragraph_props: tuple[str, ...]
    run_props: tuple[tuple[str, ...], ...]
    text: str


def _without_text(element) -> str:
    if element is None:
        return ""
    clone = copy.deepcopy(element)
    for text_node in clone.xpath(".//w:t"):
        text_node.text = ""
    return clone.xml


def _cell_style_snapshot(cell) -> tuple[str, tuple[str, ...], tuple[tuple[str, ...], ...]]:
    tc_pr = _without_text(cell._tc.tcPr)
    p_props = tuple(_without_text(paragraph._p.pPr) for paragraph in cell.paragraphs)
    r_props = tuple(
        tuple(_without_text(run._r.rPr) for run in paragraph.runs)
        for paragraph in cell.paragraphs
    )
    return tc_pr, p_props, r_props


def _format_anchor(cell) -> tuple[str, str, str]:
    """Return the immutable formatting anchor for a writable cell.

    Value text may change and may use semantic line breaks, but the cell,
    first paragraph and first run must retain the fresh template's format.
    """
    tc_pr = _without_text(cell._tc.tcPr)
    paragraph = cell.paragraphs[0] if cell.paragraphs else None
    p_pr = _without_text(paragraph._p.pPr) if paragraph is not None else ""
    run = paragraph.runs[0] if paragraph is not None and paragraph.runs else None
    if run is not None and run._r.rPr is not None:
        r_pr = _without_text(run._r.rPr)
    elif paragraph is not None and paragraph._p.pPr is not None:
        r_pr = _without_text(paragraph._p.pPr.find(qn("w:rPr")))
    else:
        r_pr = ""
    return tc_pr, p_pr, r_pr


def _bold_node_is_on(node) -> bool:
    """Interpret the OOXML boolean used by ``w:b`` and ``w:bCs``."""
    value = node.get(qn("w:val"))
    return value is None or str(value).casefold() not in {"0", "false", "off", "no"}


def _element_has_bold(element) -> bool:
    if element is None:
        return False
    return any(
        _bold_node_is_on(node)
        for tag in (qn("w:b"), qn("w:bCs"))
        for node in element.iter(tag)
    )


def _set_rpr_nonbold(rpr, *, override_inherited: bool = False) -> None:
    """Remove direct bold and optionally override bold inherited by a run."""
    for tag in (qn("w:b"), qn("w:bCs")):
        for node in list(rpr.iter(tag)):
            parent = node.getparent()
            if parent is not None:
                parent.remove(node)
    if override_inherited:
        for tag in (qn("w:b"), qn("w:bCs")):
            node = rpr.makeelement(tag, {qn("w:val"): "0"})
            rpr.append(node)


def _value_run_is_bold(run, paragraph) -> bool:
    """Detect direct run or paragraph-mark bold on a writable value."""
    direct_rpr = run._r.rPr
    direct_nodes = [
        node
        for tag in (qn("w:b"), qn("w:bCs"))
        for node in (list(direct_rpr.iter(tag)) if direct_rpr is not None else [])
    ]
    if direct_nodes:
        return any(_bold_node_is_on(node) for node in direct_nodes)
    if run.bold is True:
        return True
    paragraph_rpr = (
        paragraph._p.pPr.find(qn("w:rPr"))
        if paragraph._p.pPr is not None else None
    )
    return _element_has_bold(paragraph_rpr)


def _without_bold(element) -> str:
    """Return an XML formatting anchor with only bold removed."""
    if element is None:
        return ""
    clone = copy.deepcopy(element)
    for tag in (qn("w:b"), qn("w:bCs")):
        for node in list(clone.iter(tag)):
            parent = node.getparent()
            if parent is not None:
                parent.remove(node)
    return _without_text(clone)


def _format_anchor_without_bold(cell) -> tuple[str, str, str]:
    """Return a value anchor while treating bold removal as intentional."""
    tc_pr = _without_bold(cell._tc.tcPr)
    paragraph = cell.paragraphs[0] if cell.paragraphs else None
    p_pr = _without_bold(paragraph._p.pPr) if paragraph is not None else ""
    run = paragraph.runs[0] if paragraph is not None and paragraph.runs else None
    if run is not None and run._r.rPr is not None:
        r_pr = _without_bold(run._r.rPr)
    elif paragraph is not None and paragraph._p.pPr is not None:
        r_pr = _without_bold(paragraph._p.pPr.find(qn("w:rPr")))
    else:
        r_pr = ""
    return tc_pr, p_pr, r_pr


def _format_anchor_value_layout(cell) -> tuple[str, str, str]:
    """Return the immutable layout anchor for a writable value cell.

    Value typography is a global contract: CN values are Songti small-four
    and EN values are Times New Roman small-four.  Therefore value-cell run
    font/size/lang properties are intentionally excluded here; cell and
    paragraph layout remain locked and are still compared byte-for-byte.
    Dedicated typography audits validate the required value font separately.
    """
    tc_pr = _without_bold(cell._tc.tcPr)
    paragraph = cell.paragraphs[0] if cell.paragraphs else None
    p_pr = _without_bold(paragraph._p.pPr) if paragraph is not None else ""
    return tc_pr, p_pr, ""


def _iter_value_text_runs(cell, locked_prefix: str | None = None):
    """Yield non-empty value portions while skipping a shared-cell prefix."""
    prefix_remaining = len(locked_prefix or "")
    for paragraph in cell.paragraphs:
        for run in paragraph.runs:
            text = run.text or ""
            if not text:
                continue
            if prefix_remaining:
                consumed = min(prefix_remaining, len(text))
                value_text = text[consumed:]
                prefix_remaining -= consumed
            else:
                value_text = text
            if value_text.strip():
                yield run, paragraph, value_text


def audit_values_nonbold(document, language: str = "cn", context=None) -> list[dict]:
    """Fail closed when any writable value text is bold.

    ``language`` is retained in the evidence API so CN/EN audits use the same
    contract.  The resolver deliberately excludes all template-owned labels,
    sublabels, table headers and Section 2.8 route-prefix text.
    """
    if context is not None:
        return context.value_typography(language=language)
    problems: list[dict] = []
    for table_index, table in enumerate(document.tables):
        for row_index, row in enumerate(table.rows):
            for cell_index, cell, locked_prefix in _writable_value_cells(
                table_index, row_index, row
            ):
                for run, paragraph, value_text in _iter_value_text_runs(
                    cell, locked_prefix
                ):
                    if _value_run_is_bold(run, paragraph):
                        problems.append({
                            "type": "bold_writable_value",
                            "language": language,
                            "table": table_index,
                            "row": row_index,
                            "cell": cell_index,
                            "text": value_text,
                        })
    return problems


def audit_value_typography_contract(document, language: str = "cn") -> list[str]:
    """Fail closed on explicit font, size, alignment and boldness for values."""
    normalized_language = "zh" if language in {"zh", "cn"} else language
    expected_font = VALUE_FONT_BY_LANGUAGE.get(normalized_language)
    if expected_font is None:
        return [f"unsupported value typography language: {language}"]
    expected_size = str(int(VALUE_SIZE_PT * 2))
    errors: list[str] = []
    for table_index, table in enumerate(document.tables[:16]):
        for row_index, row in enumerate(table.rows):
            expected_alignment = 1 if table_index == 2 and row_index >= 4 else 0
            for cell_index, cell, locked_prefix in _writable_value_cells(
                table_index, row_index, row
            ):
                tc_pr = cell._tc.tcPr
                v_align = tc_pr.find(qn("w:vAlign")) if tc_pr is not None else None
                if v_align is None or v_align.get(qn("w:val")) != "center":
                    errors.append(
                        f"value cell vertical alignment must be center: table={table_index + 1} "
                        f"row={row_index + 1} cell={cell_index + 1}"
                    )
                for paragraph in cell.paragraphs:
                    paragraph_alignment = getattr(paragraph.alignment, "value", paragraph.alignment)
                    if paragraph_alignment != expected_alignment and (paragraph.text.strip() or table_index == 2):
                        errors.append(
                            f"value paragraph alignment mismatch: table={table_index + 1} row={row_index + 1} "
                            f"cell={cell_index + 1} expected={expected_alignment} actual={paragraph_alignment}"
                        )
                for run, paragraph, value_text in _iter_value_text_runs(cell, locked_prefix):
                    r_pr = run._r.rPr
                    if r_pr is None:
                        errors.append(
                            f"value run has no explicit typography: table={table_index + 1} "
                            f"row={row_index + 1} cell={cell_index + 1} text={value_text[:24]!r}"
                        )
                        continue
                    r_fonts = r_pr.find(qn("w:rFonts"))
                    actual_fonts = {
                        key: r_fonts.get(qn(f"w:{key}")) if r_fonts is not None else None
                        for key in ("ascii", "hAnsi", "eastAsia", "cs")
                    }
                    if any(font != expected_font for font in actual_fonts.values()):
                        errors.append(
                            f"value run font mismatch: table={table_index + 1} row={row_index + 1} "
                            f"cell={cell_index + 1} expected={expected_font!r} actual={actual_fonts}"
                        )
                    for tag in ("w:sz", "w:szCs"):
                        size = r_pr.find(qn(tag))
                        if size is None or size.get(qn("w:val")) != expected_size:
                            errors.append(
                                f"value run explicit size mismatch: table={table_index + 1} "
                                f"row={row_index + 1} cell={cell_index + 1} text={value_text[:24]!r} "
                                f"expected={expected_size} actual={size.get(qn('w:val')) if size is not None else None}"
                            )
                    if _value_run_is_bold(run, paragraph):
                        errors.append(
                            f"writable value is bold: table={table_index + 1} row={row_index + 1} "
                            f"cell={cell_index + 1} text={value_text[:24]!r}"
                        )
    return errors


def _is_s8_recommendation_value_cell(table_index: int, row_index: int,
                                     row, cell_index: int) -> bool:
    """Return true only for the source-gated Section 8 value cell."""
    # The row may move after an earlier source-absent row is suppressed.  The
    # semantic label is the stable identity; physical row position is not.
    if table_index != 7:
        return False
    cells = unique_cells(row)
    return (
        len(cells) > 1
        and cell_index == len(cells) - 1
        and bool(re.match(r"^\s*(?:建议|recommendation)\b", cells[0].text, re.I))
    )


def _is_s82_parent_row_from_cells(cells: Sequence) -> bool:
    """Identify the Section 8.2 parent row after earlier rows move."""
    if len(cells) != 1:
        return False
    return bool(re.match(
        r"^\s*(?:工作场所组分控制参数|control parameters for workplace components)\s*$",
        cells[0].text,
        re.I,
    ))


def _is_s82_parent_row(row) -> bool:
    return _is_s82_parent_row_from_cells(unique_cells(row))


def _is_s82_header_row_from_cells(cells: Sequence) -> bool:
    """Identify the locked four-column Section 8.2 header semantically."""
    header = tuple(cell.text.strip() for cell in cells)
    return header in set(S82_TOP_HEADERS.values())


def _is_s82_header_row(row) -> bool:
    return _is_s82_header_row_from_cells(unique_cells(row))


def _is_s82_data_row_from_cells(cells: Sequence) -> bool:
    """Identify a four-column Section 8.2 data row independent of its index."""
    return len(cells) == 4 and not _is_s82_header_row_from_cells(cells)


def _is_s82_data_row(row) -> bool:
    return _is_s82_data_row_from_cells(unique_cells(row))


def _writable_value_cells(table_index: int, row_index: int, row) -> list[tuple[int, object, str | None]]:
    """Return semantic value cells as ``(index, cell, locked_prefix)``.

    This resolver is intentionally independent of the original row number so
    it remains correct after an approved empty-row suppression.  It is shared
    by the format comparison and the non-bold value gate; locked labels,
    sublabels, headers and route prefixes are never classified as values.
    """
    cells = unique_cells(row)
    if not cells or row_index == 0:
        return []
    if table_index == 14 and is_s15_locked_heading_row(row):
        # Section 15 has two template-owned structural headings that happen to
        # be one-cell rows.  They are not value slots and must never be
        # cleared, reformatted, or treated as non-bold output values.
        return []

    if table_index == 2:
        if row_index in {2, 3}:
            return []
        if len(cells) == 3:
            first = re.sub(r"\s+", "", cells[0].text).casefold()
            if first in {"成分", "化学品名称", "chemicalname", "cas编号", "casno."}:
                return []
            return [(index, cell, None) for index, cell in enumerate(cells)]

    if table_index == 7:
        if _is_s82_parent_row_from_cells(cells) or _is_s82_header_row_from_cells(cells):
            return []
        if _is_s82_data_row_from_cells(cells):
            return [(index, cell, None) for index, cell in enumerate(cells)]
        if len(cells) == 1 and re.match(
            r"^\s*(?:8\.1\s*)?(?:暴露控制|exposure\s+controls?)\b",
            cells[0].text,
            re.I,
        ):
            return []

    s28 = _s28_parts(cells) if table_index == 1 else None
    if s28 is not None:
        prefix_cell_index, value_index, prefix = s28
        return [(value_index, cells[value_index], prefix if prefix_cell_index is None else None)]

    if (
        table_index == 10
        and len(cells) >= 3
        and re.match(r"^\s*11\.(?:1|2|7)\b", cells[0].text, re.I)
    ):
        return [(len(cells) - 1, cells[-1], None)]

    if len(cells) == 1:
        return [(0, cells[0], None)]
    return [(index, cell, None) for index, cell in enumerate(cells[1:], start=1)]


def _row_label(row) -> str:
    cells = unique_cells(row)
    return _canonical_locked_text(cells[0].text) if cells else ""


def _canonical_locked_text(text: str) -> str:
    """Compare locked labels apart from approved sequence digits/spaces."""
    text = str(text or "").strip()
    # Sequence renumbering is the only approved label-cell mutation.  The
    # separator spaces belong to that prefix too: S9 changes 9.12 + one space
    # into 9.5 + two spaces to keep the label body aligned.
    return re.sub(r"^\s*\d+\.\d+[ \t]*", "<SEQ> ", text)


def _row_candidates(template_table, output_table, table_index, output_index, output_row):
    output_cells = unique_cells(output_row)
    if table_index == 2 and output_index >= 4:
        reference_index = min(output_index, len(template_table.rows) - 1)
        return [template_table.rows[reference_index]] if len(template_table.rows) > 4 else []
    if table_index == 7:
        if _is_s82_parent_row(output_row):
            return [template_table.rows[12]] if len(template_table.rows) > 12 else []
        if _is_s82_header_row(output_row):
            return [template_table.rows[13]] if len(template_table.rows) > 13 else []
        if _is_s82_data_row(output_row):
            return [
                row for row in template_table.rows[14:]
                if _is_s82_data_row(row)
            ]
    if table_index in {8, 14} and output_index >= len(template_table.rows):
        # S9/S15 may append source-backed rows by cloning the last styled row.
        # The caller must have used the explicit inserted-data-row path.
        return [template_table.rows[-1]] if template_table.rows else []
    if not output_cells:
        return []
    label = _row_label(output_row)
    candidates = [
        row for row in template_table.rows
        if len(unique_cells(row)) == len(output_cells)
        and (len(output_cells) == 1 or _row_label(row) == label)
    ]
    if not candidates and output_index < len(template_table.rows):
        candidate = template_table.rows[output_index]
        if len(unique_cells(candidate)) == len(output_cells):
            candidates = [candidate]
    return candidates


def _row_paging_signature(row) -> tuple[bool, bool]:
    """Return the Word row settings that control page crossing behavior."""
    tr_pr = row._tr.trPr
    if tr_pr is None:
        return False, False
    return (
        tr_pr.find(qn("w:cantSplit")) is not None,
        tr_pr.find(qn("w:tblHeader")) is not None,
    )


def audit_cross_page_contract(template, output) -> dict:
    """Verify table spanning and surviving row page-crossing settings.

    Word permits a table to continue on a later page by default.  ``w:cantSplit``
    is a row-level instruction: it prevents that row from splitting internally
    but does not prohibit the table from spanning pages.  The active template
    owns this setting; output rows must preserve the fresh baseline after
    approved row omission.
    """
    errors: list[str] = []
    tables: list[dict] = []
    if len(template.tables) != len(output.tables):
        return {"errors": [
            f"table count changed for cross-page contract: {len(template.tables)} -> {len(output.tables)}"
        ], "tables": []}
    for table_index, (template_table, output_table) in enumerate(
        zip(template.tables, output.tables), start=1
    ):
        template_pr = template_table._tbl.tblPr
        output_pr = output_table._tbl.tblPr
        template_layout = template_pr.find(qn("w:tblLayout")) if template_pr is not None else None
        output_layout = output_pr.find(qn("w:tblLayout")) if output_pr is not None else None
        template_layout_type = template_layout.get(qn("w:type")) if template_layout is not None else None
        output_layout_type = output_layout.get(qn("w:type")) if output_layout is not None else None
        template_table_blocked = bool(
            template_pr is not None and template_pr.find(qn("w:cantSplit")) is not None
        )
        output_table_blocked = bool(
            output_pr is not None and output_pr.find(qn("w:cantSplit")) is not None
        )
        if template_layout_type != output_layout_type:
            errors.append(
                f"table {table_index} layout changed: {template_layout_type} -> {output_layout_type}"
            )
        if template_table_blocked != output_table_blocked:
            errors.append(f"table {table_index} cross-page permission changed")

        template_rows = [_row_paging_signature(row) for row in template_table.rows]
        output_rows = [_row_paging_signature(row) for row in output_table.rows]
        checked = 0
        for output_index, output_row in enumerate(output_table.rows):
            output_cells = unique_cells(output_row)
            # Several formal tables use consecutive one-cell rows.  Their
            # labels are intentionally writable note slots, so label matching
            # cannot distinguish them; approved omission keeps their physical
            # positions stable and the row index is the safer anchor.
            if (
                len(output_cells) == 1
                and output_index < len(template_table.rows)
                and len(unique_cells(template_table.rows[output_index])) == 1
            ):
                candidates = [template_table.rows[output_index]]
            else:
                candidates = _row_candidates(
                    template_table, output_table, table_index - 1, output_index, output_row
                )
            if not candidates:
                errors.append(
                    f"table {table_index} row {output_index} has no cross-page template anchor"
                )
                continue
            expected_signature = _row_paging_signature(candidates[0])
            actual_signature = _row_paging_signature(output_row)
            if expected_signature != actual_signature:
                errors.append(
                    f"table {table_index} row {output_index} cross-page row settings changed: "
                    f"{expected_signature} -> {actual_signature}"
                )
            checked += 1
        tables.append({
            "index": table_index,
            "template_layout": template_layout_type,
            "output_layout": output_layout_type,
            "template_allows_cross_page": not template_table_blocked,
            "output_allows_cross_page": not output_table_blocked,
            "template_row_breakable": sum(not cant for cant, _ in template_rows),
            "output_row_breakable": sum(not cant for cant, _ in output_rows),
            "template_row_cant_split": sum(cant for cant, _ in template_rows),
            "output_row_cant_split": sum(cant for cant, _ in output_rows),
            "template_all_rows_breakable": not any(cant for cant, _ in template_rows),
            "output_all_rows_breakable": not any(cant for cant, _ in output_rows),
            "surviving_rows_checked": checked,
        })
    return {"errors": errors, "tables": tables}


def _style_signature(element):
    """Compare OOXML formatting without document-part namespace noise."""
    if element is None:
        return ()
    return (
        element.tag,
        tuple(sorted((key, value) for key, value in element.attrib.items())),
        tuple(_style_signature(child) for child in element),
    )


def _format_layout_anchor(cell) -> tuple[str, str]:
    tc_pr = _without_text(cell._tc.tcPr)
    paragraph = cell.paragraphs[0] if cell.paragraphs else None
    p_pr = _without_text(paragraph._p.pPr) if paragraph is not None else ""
    return tc_pr, p_pr


def _value_run_rpr(run, paragraph):
    if run._r.rPr is not None:
        return run._r.rPr
    if paragraph._p.pPr is not None:
        return paragraph._p.pPr.find(qn("w:rPr"))
    return None


def compare_format_anchors(template, output, *, language: str = "cn",
                           approved_en_body_rpr=None) -> list[str]:
    """Audit all surviving cells against the fresh template's format anchors.

    This is deliberately separate from the locked-label audit: the latter
    protects labels, while this audit prevents a writable value cell from
    becoming a silently re-formatted paragraph. Approved row omissions and S3
    or S8.2 styled-row cloning are handled by candidate matching.
    """
    errors: list[str] = []
    if len(template.tables) != len(output.tables):
        return [f"table count changed: {len(template.tables)} -> {len(output.tables)}"]
    for table_index, (template_table, output_table) in enumerate(
        zip(template.tables, output.tables)
    ):
        if _without_text(template_table._tbl.tblPr) != _without_text(output_table._tbl.tblPr):
            errors.append(f"table properties changed: table {table_index}")
        if _without_text(template_table._tbl.tblGrid) != _without_text(output_table._tbl.tblGrid):
            errors.append(f"table grid changed: table {table_index}")
    for table_index, output_table in enumerate(output.tables):
        template_table = template.tables[table_index]
        for output_index, output_row in enumerate(output_table.rows):
            candidates = _row_candidates(template_table, output_table, table_index,
                                          output_index, output_row)
            if not candidates:
                actual_anchor = tuple(_format_anchor(cell) for cell in unique_cells(output_row))
                errors.append(
                    f"no template format anchor: table {table_index} row {output_index}; "
                    + format_diagnostic(
                        "FORMAT_ANCHOR_MISSING",
                        location=f"table={table_index + 1} row={output_index + 1}",
                        expected="a matching fresh-template row anchor",
                        actual=actual_anchor,
                        hint="preserve the cloned row and use only an authorized row policy",
                    )
                )
                continue
            output_cells = unique_cells(output_row)
            matched = False
            body_cells = english_body_cells(table_index, output_index, output_row) \
                if language == "en" else []
            body_tc_ids = {hash(cell._tc) for cell in body_cells}
            value_cells = {
                cell_index: locked_prefix
                for cell_index, _cell, locked_prefix in _writable_value_cells(
                    table_index, output_index, output_row
                )
            }
            composite_parts = _s28_parts(output_cells) if table_index == 1 else None
            for candidate in candidates:
                candidate_cells = unique_cells(candidate)
                if len(candidate_cells) != len(output_cells):
                    continue
                if _without_text(output_row._tr.trPr) != _without_text(candidate._tr.trPr):
                    continue
                cell_formats_match = True
                for cell_index, (expected_cell, actual_cell) in enumerate(
                    zip(candidate_cells, output_cells)
                ):
                    if hash(actual_cell._tc) in body_tc_ids:
                        if _format_layout_anchor(expected_cell) != _format_layout_anchor(actual_cell):
                            cell_formats_match = False
                            break
                    elif _is_s8_recommendation_value_cell(
                        table_index, output_index, output_row, cell_index
                    ):
                        # The formal blank recommendation paragraph inherits
                        # the label's bold paragraph-mark properties.  Its
                        # source-gated value is deliberately written with a
                        # non-bold run, so compare the cell/paragraph layout
                        # here and validate the value run below.
                        if _format_layout_anchor(expected_cell) != _format_layout_anchor(actual_cell):
                            cell_formats_match = False
                            break
                    elif (
                        cell_index in value_cells
                        # In the maintained two-column S2.8 layout the first
                        # run is the locked route prefix.  Its value tail is
                        # checked by the dedicated composite writer/audit;
                        # do not relax the prefix anchor here.
                        and not (
                            table_index == 1
                            and composite_parts is not None
                            and len(output_cells) == 2
                            and cell_index == composite_parts[1]
                        )
                    ):
                        # The only intentional character-format difference in
                        # a writable value is removal of bold.  Font, size,
                        # color, language, spacing and all cell/paragraph
                        # layout must still inherit the fresh template.
                        expected_layout = _format_anchor_value_layout(expected_cell)
                        actual_layout = _format_anchor_value_layout(actual_cell)
                        # S2.5 deliberately expands one value into heading and
                        # detail paragraphs; its cell geometry remains locked,
                        # while paragraph indentation is the local rule.
                        if table_index == 1 and (
                            "防范说明" in output_cells[0].text
                            or "Precautionary" in output_cells[0].text
                        ):
                            expected_layout = (expected_layout[0], "", expected_layout[2])
                            actual_layout = (actual_layout[0], "", actual_layout[2])
                        if expected_layout != actual_layout:
                            cell_formats_match = False
                            break
                    elif _format_anchor(expected_cell) != _format_anchor(actual_cell):
                        cell_formats_match = False
                        break
                if cell_formats_match:
                    matched = True
                    break
            if not matched:
                expected_anchor = tuple(
                    tuple(_format_anchor(cell) for cell in unique_cells(candidate))
                    for candidate in candidates[:1]
                )
                actual_anchor = tuple(_format_anchor(cell) for cell in output_cells)
                errors.append(
                    f"template format anchor changed: table {table_index} row {output_index}; "
                    + format_diagnostic(
                        "FORMAT_MISMATCH",
                        location=f"table={table_index + 1} row={output_index + 1}",
                        expected=expected_anchor,
                        actual=actual_anchor,
                        hint="compare tcPr/pPr/rPr with the fresh template; do not reformat the output",
                    )
                )
                continue
            if _is_s8_recommendation_value_cell(
                table_index, output_index, output_row, len(output_cells) - 1
            ):
                value_cell = output_cells[-1]
                if any(
                    run.text.strip() and run.bold is True
                    for paragraph in value_cell.paragraphs
                    for run in paragraph.runs
                ):
                    errors.append(
                        f"Section 8 recommendation value must use non-bold body text: "
                        f"table {table_index} row {output_index}"
                    )
            if language == "en" and approved_en_body_rpr is not None:
                for cell in body_cells:
                    for paragraph in cell.paragraphs:
                        for run in paragraph.runs:
                            if run.text.strip() and not run.bold:
                                actual_rpr = _value_run_rpr(run, paragraph)
                                if (_style_signature(actual_rpr)
                                        != _style_signature(approved_en_body_rpr)):
                                    errors.append(
                                        f"EN body value format is not the approved exemplar: "
                                        f"table {table_index} row {output_index}; "
                                        + format_diagnostic(
                                            "EN_BODY_FORMAT_MISMATCH",
                                            location=f"table={table_index + 1} row={output_index + 1}",
                                            expected=_style_signature(approved_en_body_rpr),
                                            actual=_style_signature(actual_rpr),
                                            hint="retain the approved EN body exemplar run properties",
                                        )
                                    )
                                    break

    for section_index, (template_section, output_section) in enumerate(
        zip(template.sections, output.sections)
    ):
        if _without_text(template_section._sectPr) != _without_text(output_section._sectPr):
            errors.append(f"section properties changed: section {section_index}")
        for role in ("header", "footer"):
            template_tables = getattr(template_section, role).tables
            output_tables = getattr(output_section, role).tables
            if len(template_tables) != len(output_tables):
                errors.append(f"{role} table count changed: section {section_index}")
                continue
            for table_index, output_table in enumerate(output_tables):
                template_table = template_tables[table_index]
                for row_index, output_row in enumerate(output_table.rows):
                    if row_index >= len(template_table.rows):
                        errors.append(f"{role} row added: section {section_index} table {table_index}")
                        continue
                    expected_cells = unique_cells(template_table.rows[row_index])
                    actual_cells = unique_cells(output_row)
                    if len(expected_cells) != len(actual_cells) or any(
                        _format_anchor(expected) != _format_anchor(actual)
                        for expected, actual in zip(expected_cells, actual_cells)
                    ):
                        errors.append(f"{role} format changed: section {section_index} table {table_index} row {row_index}")
    return errors


def _locked_cell_indices(table_index: int, row_index: int, cells: Sequence) -> tuple[int, ...]:
    """Return physical cells whose content/format is template-owned."""
    if row_index == 0:
        return tuple(range(len(cells)))
    if table_index == 14 and _is_s15_locked_heading_cells(cells):
        return tuple(range(len(cells)))
    if table_index == 2 and row_index in {2, 3}:
        return tuple(range(len(cells)))
    if table_index == 2 and row_index >= 4:
        return ()
    if table_index == 7:
        # S8.2 can move upward when preceding source-absent PPE rows are
        # removed.  Use the semantic parent/header/data shapes first; retain
        # the physical-index fallback for a pristine template clone.
        if _is_s82_parent_row_from_cells(cells) or _is_s82_header_row_from_cells(cells):
            return tuple(range(len(cells)))
        if _is_s82_data_row_from_cells(cells):
            return ()
        if row_index in {12, 13}:
            return tuple(range(len(cells)))
        if row_index >= 14:
            return ()
    s28 = _s28_parts(cells) if table_index == 1 else None
    if s28 is not None:
        _, value_index, _ = s28
        # A current two-cell composite keeps the route prefix inside the
        # writable value cell, so its prefix is checked by the dedicated
        # composite-prefix audit rather than as a whole locked cell.
        return tuple(index for index in range(len(cells)) if index != value_index) \
            if len(cells) >= 3 else (0,)
    if table_index == 10 and len(cells) >= 3 \
            and re.match(r"^\s*11\.(?:1|7)\b", cells[0].text):
        return tuple(range(len(cells) - 1))
    if len(cells) == 1:
        return ()
    return (0,)


def _composite_prefix_style_snapshot(cell, prefix: str) -> tuple:
    """Capture only the immutable prefix run formatting, not the value tail."""
    paragraph = cell.paragraphs[0] if cell.paragraphs else None
    if paragraph is None:
        return ("", (), ())
    cursor = 0
    run_props = []
    for run in paragraph.runs:
        text = run.text or ""
        if not text or cursor >= len(prefix):
            continue
        take = prefix[cursor:cursor + len(text)]
        if not take:
            continue
        run_props.append(_without_text(run._r.rPr))
        cursor += len(take)
    return (
        _without_text(cell._tc.tcPr),
        (_without_text(paragraph._p.pPr),),
        tuple(run_props),
    )


def _compare_s28_composite_prefixes(template, output) -> list[str]:
    """Fail closed when a template-owned Section 2.8 route prefix drifts."""
    if len(template.tables) <= 1 or len(output.tables) <= 1:
        return []
    expected_rows = []
    actual_rows = []
    for row in template.tables[1].rows[1:]:
        cells = unique_cells(row)
        if cells and S28_LABEL_RE.match(cells[0].text or ""):
            expected_rows.append((cells, _s28_parts(cells)))
    for row in output.tables[1].rows[1:]:
        cells = unique_cells(row)
        if cells and S28_LABEL_RE.match(cells[0].text or ""):
            actual_rows.append((cells, _s28_parts(cells)))
    if not expected_rows:
        return []
    # Some maintained EN baselines use ordinary value cells for S2.8 rather
    # than embedding route prefixes in the value cell.  In that topology the
    # generic value-cell/locked-label audits are authoritative; this composite
    # audit has no prefix boundary to validate.
    if all(parts is None for _cells, parts in expected_rows):
        return []

    errors = []
    expected_by_label = {}
    actual_by_label = {}
    for cells, parts in expected_rows:
        label = _canonical_locked_text(cells[0].text)
        expected_by_label.setdefault(label, []).append((cells, parts))
    for cells, parts in actual_rows:
        label = _canonical_locked_text(cells[0].text)
        actual_by_label.setdefault(label, []).append((cells, parts))

    # Approved omission may remove any complete S2.8 child row.  Compare only
    # surviving route prefixes, while still rejecting a changed/duplicated
    # prefix or a value that no longer follows its locked prefix boundary.
    for label, candidates in actual_by_label.items():
        expected = expected_by_label.get(label)
        if not expected:
            errors.append(f"Section 2.8 route label is not in the template skeleton: {label}")
            continue
        expected_by_prefix = {parts[2]: (cells, parts) for cells, parts in expected if parts}
        seen_prefixes = set()
        for actual_cells, actual_parts in candidates:
            if actual_parts is None:
                errors.append(
                    f"Section 2.8 route prefix/value shape is not recognized "
                    f"for locked label {label}"
                )
                continue
            prefix = actual_parts[2]
            expected_entry = expected_by_prefix.get(prefix)
            if expected_entry is None:
                errors.append(
                    f"Section 2.8 route prefix changed for locked label {label}: "
                    f"found {prefix!r}"
                )
                continue
            if prefix in seen_prefixes:
                errors.append(
                    f"Section 2.8 route prefix duplicated for locked label {label}: {prefix!r}"
                )
                continue
            seen_prefixes.add(prefix)
            expected_cells, expected_parts = expected_entry
            expected_target = expected_cells[expected_parts[0] or expected_parts[1]]
            actual_target = actual_cells[actual_parts[0] or actual_parts[1]]
            actual_prefix = _route_prefix_in_text(actual_target.text)
            if actual_target.text != actual_prefix \
                    and not actual_target.text.startswith(actual_prefix + "\n"):
                errors.append(
                    f"Section 2.8 route prefix/value boundary changed for {label}: "
                    "the locked prefix must be followed only by a template line break and value"
                )
            if _composite_prefix_style_snapshot(expected_target, prefix) \
                    != _composite_prefix_style_snapshot(actual_target, prefix):
                errors.append(
                    f"Section 2.8 route prefix formatting changed for locked label {label}: {prefix!r}"
                )
    return errors


def locked_cell_snapshots(document, *, allow_added_data_rows: bool = False) -> list[LockedCellSnapshot]:
    """Snapshot sequence/label cells and all locked table headers.

    Normal field rows lock the first physical cell.  S3/S8.2 table headers
    lock every header cell; S3 data rows are explicitly writable data rows.
    The row key uses label text with the numeric prefix removed plus an
    occurrence number, allowing S2/S9 approved row omission without confusing
    later rows.
    """
    snapshots: list[LockedCellSnapshot] = []
    for table_index, table in enumerate(document.tables):
        occurrence: dict[str, int] = {}
        for row_index, row in enumerate(table.rows):
            cells = unique_cells(row)
            if not cells:
                continue
            label = cells[0].text.strip()
            logical = _canonical_locked_text(label)
            logical = re.sub(r"\s+", " ", logical)
            occurrence[logical] = occurrence.get(logical, 0) + 1
            row_key = (logical, occurrence[logical])
            added_data_row = allow_added_data_rows and (
                (table_index == 8 and row_index >= 24)
                or (table_index == 14 and row_index >= 9)
            )
            if added_data_row:
                # The first cell of a newly inserted S9 source row is seeded
                # from the source field name; it is not an existing locked
                # template label. Its style is still checked by the format
                # and cross-page audits against the cloned row anchor.
                protected_indices = ()
            else:
                protected_indices = _locked_cell_indices(table_index, row_index, cells)
            protected = [cells[index] for index in protected_indices]
            for cell_index, cell in zip(protected_indices, protected):
                tc_pr, p_props, r_props = _cell_style_snapshot(cell)
                role = "sequence_label" if cell_index == 0 else "sub_label_or_header"
                snapshots.append(LockedCellSnapshot(
                    table_index, row_key, role, tc_pr, p_props, r_props, cell.text,
                ))
    return snapshots


def compare_locked_skeleton(template, output) -> list[str]:
    """Return release-blocking differences in locked cells."""
    expected = locked_cell_snapshots(template)
    actual = locked_cell_snapshots(output, allow_added_data_rows=True)
    # Match on the stable label body, not occurrence index.  S2/S9 are
    # allowed to remove missing rows, and repeated children (for example
    # multiple 2.8 rows) must not make later rows look like format drift.
    expected_map: dict[tuple[int, str, str], list[LockedCellSnapshot]] = {}
    for snapshot in expected:
        expected_map.setdefault(
            (snapshot.table_index, snapshot.row_key[0], snapshot.cell_role),
            [],
        ).append(snapshot)
    errors: list[str] = []
    consumed: dict[tuple[int, str, str], set[int]] = {}
    for actual_item in actual:
        key = (actual_item.table_index, actual_item.row_key[0], actual_item.cell_role)
        candidates = expected_map.get(key)
        if not candidates:
            if actual_item.cell_role == "sequence_label":
                errors.append(f"locked label text changed or is not in template skeleton: {key}")
            else:
                errors.append(f"locked cell is not in template skeleton: {key}")
            continue
        used = consumed.setdefault(key, set())
        available = [
            (index, candidate) for index, candidate in enumerate(candidates)
            if index not in used
        ]
        if not available:
            available = list(enumerate(candidates))
        matching = next(
            ((index, candidate) for index, candidate in available
             if candidate.tc_pr == actual_item.tc_pr
             and candidate.paragraph_props == actual_item.paragraph_props
             and candidate.run_props == actual_item.run_props),
            available[0],
        )
        used.add(matching[0])
        expected_item = matching[1]
        if expected_item.tc_pr != actual_item.tc_pr:
            errors.append(
                f"locked cell tcPr changed: {key}; "
                + format_diagnostic(
                    "LOCKED_TCPR_MISMATCH", location=f"table={key[0] + 1} key={key[1:]}",
                    expected=expected_item.tc_pr, actual=actual_item.tc_pr,
                    hint="restore the locked cell properties from a fresh template clone",
                )
            )
        if expected_item.paragraph_props != actual_item.paragraph_props:
            errors.append(
                f"locked cell paragraph properties changed: {key}; "
                + format_diagnostic(
                    "LOCKED_PPR_MISMATCH", location=f"table={key[0] + 1} key={key[1:]}",
                    expected=expected_item.paragraph_props, actual=actual_item.paragraph_props,
                    hint="restore locked label paragraph properties; only value text may change",
                )
            )
        if expected_item.run_props != actual_item.run_props:
            errors.append(
                f"locked cell run properties changed: {key}; "
                + format_diagnostic(
                    "LOCKED_RPR_MISMATCH", location=f"table={key[0] + 1} key={key[1:]}",
                    expected=expected_item.run_props, actual=actual_item.run_props,
                    hint="restore bold/font/run properties from the fresh template",
                )
            )
        if expected_item.cell_role in {"sequence_label", "sub_label_or_header"}:
            expected_text = _canonical_locked_text(expected_item.text)
            actual_text = _canonical_locked_text(actual_item.text)
            if expected_text != actual_text:
                errors.append(
                    f"locked {'label' if expected_item.cell_role == 'sequence_label' else 'sub-label/header'} text changed: {key}; "
                    + format_diagnostic(
                        "LOCKED_LABEL_MISMATCH" if expected_item.cell_role == "sequence_label" else "LOCKED_SUBLABEL_MISMATCH",
                        location=f"table={key[0] + 1} key={key[1:]}",
                        expected=expected_text, actual=actual_text,
                        hint="write only the final value cell; never rewrite a template-owned label or sub-label",
                    )
                )
    errors.extend(_compare_s28_composite_prefixes(template, output))
    # The formal S8.2 header row is a locked top-level structure.  Its text is
    # checked explicitly here because data-row writes share the same table;
    # formatting is covered by locked_cell_snapshots above.  Nested child
    # tables, when present in older baselines, are still checked below for
    # rollback/audit compatibility.
    if len(template.tables) > 7 and len(output.tables) > 7:
        template_t7, output_t7 = template.tables[7], output.tables[7]
        template_header = next(
            (row for row in template_t7.rows if _is_s82_header_row(row)), None
        )
        output_header = next(
            (row for row in output_t7.rows if _is_s82_header_row(row)), None
        )
        if output_header is None and template_header is not None:
            # A changed header no longer matches the semantic header key, but
            # the row immediately after a surviving S8.2 parent is still the
            # authoritative header candidate and must fail closed on text.
            output_parent = next(
                (row for row in output_t7.rows if _is_s82_parent_row(row)), None
            )
            if output_parent is not None:
                parent_index = next(
                    index for index, row in enumerate(output_t7.rows)
                    if _is_s82_parent_row(row)
                )
                candidate_index = parent_index + 1
                if candidate_index < len(output_t7.rows):
                    candidate = output_t7.rows[candidate_index]
                    if len(unique_cells(candidate)) == len(unique_cells(template_header)):
                        output_header = candidate
        if template_header is not None and output_header is None:
            if any(_is_s82_parent_row(row) for row in output_t7.rows):
                errors.append("locked S8.2 header missing: table 7")
        elif template_header is not None and output_header is not None:
            expected_cells = unique_cells(template_header)
            actual_cells = unique_cells(output_header)
            if len(expected_cells) != len(actual_cells):
                errors.append("locked S8.2 header width changed: table 7")
            else:
                for expected_cell, actual_cell in zip(expected_cells, actual_cells):
                    if expected_cell.text != actual_cell.text:
                        errors.append("locked S8.2 header text changed: table 7")
                    elif _cell_style_snapshot(expected_cell) != _cell_style_snapshot(actual_cell):
                        errors.append("locked S8.2 header formatting changed: table 7")
    for table_index, template_table in enumerate(template.tables):
        if table_index >= len(output.tables):
            continue
        output_table = output.tables[table_index]
        for row_index, template_row in enumerate(template_table.rows):
            if row_index >= len(output_table.rows):
                continue
            for template_cell, output_cell in zip(unique_cells(template_row), unique_cells(output_table.rows[row_index])):
                template_children = list(template_cell.tables)
                output_children = list(output_cell.tables)
                if not template_children:
                    continue
                if len(template_children) != len(output_children):
                    errors.append(f"locked child-table count changed: table {table_index} row {row_index}")
                    continue
                for child_index, (expected_child, actual_child) in enumerate(zip(template_children, output_children)):
                    if len(expected_child.rows) == 0 or len(actual_child.rows) == 0:
                        errors.append(f"locked child-table header missing: table {table_index} row {row_index}")
                        continue
                    expected_header = unique_cells(expected_child.rows[0])
                    actual_header = unique_cells(actual_child.rows[0])
                    if len(expected_header) != len(actual_header):
                        errors.append(f"locked child-table header width changed: table {table_index} row {row_index}")
                        continue
                    for header_cell, actual_header_cell in zip(expected_header, actual_header):
                        if header_cell.text != actual_header_cell.text:
                            errors.append(f"locked child-table header text changed: table {table_index} row {row_index}")
                        expected_style = _cell_style_snapshot(header_cell)
                        actual_style = _cell_style_snapshot(actual_header_cell)
                        if expected_style != actual_style:
                            errors.append(f"locked child-table header formatting changed: table {table_index} row {row_index}")
    return errors


__all__ = [
    "MutationViolation",
    "TemplateSlot",
    "TemplateSlotRegistry",
    "LockedCellSnapshot",
    "unique_cells",
    "is_s28_row",
    "composite_value_text",
    "composite_prefix_text",
    "composite_value_cell_index",
    "is_s15_locked_heading_row",
    "english_body_cells",
    "set_value_cell_text",
    "set_s28_composite_value_cell",
    "set_sequence_prefix",
    "write_row_values",
    "enforce_value_typography",
    "write_s82_top_rows",
    "S82_TOP_HEADERS",
    "S82_CHILD_HEADERS",
    "S82_MISSING",
    "clear_value_cells",
    "audit_values_nonbold",
    "audit_value_typography_contract",
    "locked_cell_snapshots",
    "compare_locked_skeleton",
    "compare_format_anchors",
    "audit_cross_page_contract",
]
