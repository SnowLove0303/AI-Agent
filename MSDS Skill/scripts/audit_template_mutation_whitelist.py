#!/usr/bin/env python3
"""Release-blocking audit for the template mutation whitelist."""

from __future__ import annotations

import argparse
import json
from pathlib import Path

from docx import Document

from template_mutation_whitelist import (
    audit_cross_page_contract,
    compare_format_anchors,
    compare_locked_skeleton,
)


def audit(template_path: Path, output_path: Path, *, template=None, output=None,
          language: str = "cn") -> dict:
    template = template or Document(str(template_path))
    output = output or Document(str(output_path))
    errors = compare_locked_skeleton(template, output)
    approved_en_body_rpr = None
    if language == "en":
        from normalize_en_layout import _approved_body_rpr
        approved_en_body_rpr = _approved_body_rpr(template)
    errors.extend(compare_format_anchors(
        template, output, language=language,
        approved_en_body_rpr=approved_en_body_rpr,
    ))
    cross_page = audit_cross_page_contract(template, output)
    errors.extend(cross_page["errors"])
    result = {
        "template": str(template_path),
        "output": str(output_path),
        "table_count": len(output.tables),
        "errors": errors,
        "cross_page": cross_page,
        "status": "passed" if not errors else "failed",
    }
    return result


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--template", required=True, type=Path)
    parser.add_argument("--output", required=True, type=Path, nargs="+")
    parser.add_argument("--language", choices=("cn", "en"), default="cn")
    parser.add_argument("--json", type=Path)
    args = parser.parse_args()
    results = [audit(args.template, output, language=args.language) for output in args.output]
    rendered = json.dumps(results, ensure_ascii=False, indent=2)
    print(rendered)
    if args.json:
        args.json.parent.mkdir(parents=True, exist_ok=True)
        args.json.write_text(rendered, encoding="utf-8")
    return 0 if all(item["status"] == "passed" for item in results) else 1


if __name__ == "__main__":
    raise SystemExit(main())
