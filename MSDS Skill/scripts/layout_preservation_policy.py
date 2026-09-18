"""Output-only audits for the maintained S8.2 and S11.4 layouts.

The formal templates are the authority.  These checks do not repair or
standardize the formal files; they reject an output clone when a writer,
row-removal policy or row-insertion policy changes a template-owned layout.
"""
from __future__ import annotations

import copy
import re

from docx.oxml.ns import qn

from template_mutation_whitelist import unique_cells


S82_HEADERS = {
    "cn": ("物质", "依据", "类型", "数值"),
    "en": ("Substance", "Basis", "Type", "Value"),
}


def _without_text(element) -> str:
    if element is None:
        return ""
    clone = copy.deepcopy(element)
    for node in clone.xpath(".//w:t"):
        node.text = ""
    return clone.xml


def _without_bold(element) -> str:
    if element is None:
        return ""
    clone = copy.deepcopy(element)
    for tag in (qn("w:b"), qn("w:bCs")):
        for node in list(clone.iter(tag)):
            parent = node.getparent()
            if parent is not None:
                parent.remove(node)
    return _without_text(clone)


def _cell_layout(cell, *, allow_value_bold: bool = False) -> tuple:
    tc_pr = _without_text(cell._tc.tcPr)
    paragraphs = []
    for paragraph in cell.paragraphs:
        p_pr = _without_text(paragraph._p.pPr)
        if allow_value_bold:
            # Value writing is allowed to collapse a sample cell's illustrative
            # multi-run text into one run.  The immutable check is the first
            # value-run anchor (with bold removed), plus the cell/paragraph
            # layout above; the non-bold value audit covers the actual runs.
            first_run = paragraph.runs[0] if paragraph.runs else None
            runs = (
                _without_bold(first_run._r.rPr) if first_run is not None else "",
            )
        else:
            runs = tuple(_without_text(run._r.rPr) for run in paragraph.runs)
        paragraphs.append((p_pr, runs))
    return tc_pr, tuple(paragraphs)


def _row_layout(row, *, allow_value_bold: bool = False) -> tuple:
    return (
        _without_text(row._tr.trPr),
        tuple(
            _cell_layout(cell, allow_value_bold=allow_value_bold)
            for cell in unique_cells(row)
        ),
    )


def _grid(table) -> tuple[str, ...]:
    grid = table._tbl.tblGrid
    return tuple(_without_text(col) for col in grid)


def _row_text(row) -> str:
    return " ".join(cell.text.strip() for cell in unique_cells(row)).strip()


def _is_s82_parent(row) -> bool:
    text = _row_text(row).casefold()
    return "工作场所组分控制参数" in text or "control parameters for workplace components" in text


def _is_s82_header(row, language: str) -> bool:
    return tuple(cell.text.strip() for cell in unique_cells(row)) == S82_HEADERS[language]


def _find_s82_block(table, language: str) -> tuple[int | None, int | None, list]:
    parent_index = next((i for i, row in enumerate(table.rows) if _is_s82_parent(row)), None)
    header_index = next(
        (i for i, row in enumerate(table.rows)
         if parent_index is not None and i > parent_index and _is_s82_header(row, language)),
        None,
    )
    data_rows = []
    if header_index is not None:
        for row in list(table.rows)[header_index + 1:]:
            cells = unique_cells(row)
            if len(cells) != 4:
                break
            data_rows.append(row)
    return parent_index, header_index, data_rows


def audit_s82(template, output, language: str = "cn", expected_present: bool | None = None) -> dict:
    """Audit the formal S8.2 grid and its surviving block.

    ``expected_present`` is supplied by the source-controlled writer.  When it
    is false, the complete block may be hidden, but a partial block is still a
    failure.  With ``None`` the audit infers presence from the output and is
    suitable for standalone persisted-file audits.
    """
    errors: list[str] = []
    language = "en" if language == "en" else "cn"
    if len(template.tables) <= 7 or len(output.tables) <= 7:
        return {"status": "failed", "errors": ["S8.2 table 8 is missing"], "state": "missing"}
    expected_table, actual_table = template.tables[7], output.tables[7]
    if _grid(expected_table) != _grid(actual_table):
        errors.append("S8.2 table grid changed from the formal template")
    if len(_grid(expected_table)) != 5:
        errors.append(f"formal S8.2 grid must contain five columns, found {len(_grid(expected_table))}")

    expected_parent, expected_header, expected_data = _find_s82_block(expected_table, language)
    actual_parent, actual_header, actual_data = _find_s82_block(actual_table, language)
    output_present = actual_parent is not None or actual_header is not None or bool(actual_data)
    should_be_present = expected_present if expected_present is not None else output_present
    if should_be_present and actual_parent is None:
        errors.append("S8.2 parent row is missing while source control records are present")
    if should_be_present and actual_header is None:
        errors.append("S8.2 header row is missing while source control records are present")
    if not should_be_present and output_present:
        errors.append("S8.2 block is only partially hidden; parent/header/data must be removed together")

    if expected_parent is not None and actual_parent is not None:
        expected_row, actual_row = expected_table.rows[expected_parent], actual_table.rows[actual_parent]
        if len(unique_cells(actual_row)) != 1:
            errors.append("S8.2 parent row is not a one-cell row")
        if _row_layout(expected_row) != _row_layout(actual_row):
            errors.append("S8.2 parent row topology or format changed")
    if expected_header is not None and actual_header is not None:
        expected_row, actual_row = expected_table.rows[expected_header], actual_table.rows[actual_header]
        if tuple(cell.text.strip() for cell in unique_cells(actual_row)) != S82_HEADERS[language]:
            errors.append("S8.2 header text or column order changed")
        if len(unique_cells(actual_row)) != 4:
            errors.append("S8.2 header must contain four logical cells")
        if _row_layout(expected_row) != _row_layout(actual_row):
            errors.append("S8.2 header topology or format changed")

    if should_be_present and not actual_data:
        errors.append("S8.2 data rows are missing while source control records are present")
    for position, row in enumerate(actual_data, start=1):
        cells = unique_cells(row)
        if len(cells) != 4:
            errors.append(f"S8.2 data row {position} must contain four logical cells")
            continue
        if len(_grid(actual_table)) != 5:
            errors.append(f"S8.2 data row {position} is not attached to the five-column grid")
        # The formal template can intentionally carry more than one sample
        # data-row style. Compare each surviving row with its corresponding
        # template anchor; only rows beyond the baseline use the first data
        # row as the approved clone anchor.
        expected_data_anchor = (
            expected_data[position - 1]
            if position <= len(expected_data)
            else (expected_data[0] if expected_data else None)
        )
        if expected_data_anchor is not None and _row_layout(expected_data_anchor, allow_value_bold=True) != _row_layout(row, allow_value_bold=True):
            errors.append(f"S8.2 data row {position} topology or inherited format changed")
        for cell_index, cell in enumerate(cells):
            tc_pr = cell._tc.tcPr
            grid_span = tc_pr.find(qn("w:gridSpan")) if tc_pr is not None else None
            if cell_index == 1 and (grid_span is None or grid_span.get(qn("w:val")) != "2"):
                errors.append(f"S8.2 data row {position} basis cell must retain gridSpan=2")
    return {
        "status": "passed" if not errors else "failed",
        "errors": errors,
        "state": "present" if output_present else "hidden",
        "expected_present": should_be_present,
        "data_row_count": len(actual_data),
        "grid_column_count": len(_grid(actual_table)),
    }


def _find_s114_rows(table) -> list:
    rows = []
    for row in table.rows:
        label = unique_cells(row)[0].text.strip() if unique_cells(row) else ""
        if re.match(r"^\s*11\.4\b", label) or re.search(r"sensitization|致敏性", label, re.I):
            rows.append(row)
    return rows


def _v_align(cell) -> str | None:
    tc_pr = cell._tc.tcPr
    node = tc_pr.find(qn("w:vAlign")) if tc_pr is not None else None
    return node.get(qn("w:val")) if node is not None else None


def audit_s114_vertical_alignment(template, output) -> dict:
    """Compare S11.4 cell vertical alignment with the fresh template clone."""
    errors: list[str] = []
    if len(template.tables) <= 10 or len(output.tables) <= 10:
        return {"status": "failed", "errors": ["S11.4 table 11 is missing"]}
    expected_rows = _find_s114_rows(template.tables[10])
    actual_rows = _find_s114_rows(output.tables[10])
    if not actual_rows:
        return {
            "status": "passed",
            "errors": [],
            "state": "hidden",
            "expected_rows": len(expected_rows),
        }
    if len(actual_rows) != len(expected_rows):
        errors.append(
            f"S11.4 row count changed: template {len(expected_rows)} -> output {len(actual_rows)}"
        )
    for index, actual_row in enumerate(actual_rows):
        if index >= len(expected_rows):
            break
        expected_cells = unique_cells(expected_rows[index])
        actual_cells = unique_cells(actual_row)
        if len(expected_cells) != len(actual_cells):
            errors.append(f"S11.4 row {index + 1} physical cell count changed")
            continue
        for cell_index, (expected, actual) in enumerate(zip(expected_cells, actual_cells), start=1):
            expected_align = _v_align(expected)
            actual_align = _v_align(actual)
            if expected_align != actual_align:
                errors.append(
                    f"S11.4 row {index + 1} cell {cell_index} vertical alignment changed: "
                    f"{expected_align!r} -> {actual_align!r}"
                )
    return {
        "status": "passed" if not errors else "failed",
        "errors": errors,
        "state": "present",
        "expected_rows": len(expected_rows),
        "actual_rows": len(actual_rows),
    }


__all__ = ["S82_HEADERS", "audit_s82", "audit_s114_vertical_alignment"]
