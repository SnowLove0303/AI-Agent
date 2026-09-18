"""Read-only indexes shared by compatible MSDS release audits.

The context intentionally stores references to an already-loaded DOCX and
precomputed observations only.  It has no mutation methods and never replaces
the persisted-file checks required by the release contract.
"""
from __future__ import annotations

import hashlib
import re
from dataclasses import dataclass

from docx.oxml.ns import qn

from renumber_visible_items import collect as collect_numbered_items
from template_mutation_whitelist import (
    _iter_value_text_runs,
    _value_run_is_bold,
    _writable_value_cells,
    unique_cells,
)


def _xml_signature(element) -> str:
    return hashlib.sha256(element.xml.encode("utf-8")).hexdigest()


@dataclass(frozen=True)
class RowRecord:
    table_index: int
    row_index: int
    row: object
    cells: tuple
    writable_cells: tuple


class AuditContext:
    """Immutable-by-contract observations for one in-memory DOCX."""

    def __init__(self, document):
        self.document = document
        records = []
        by_table: dict[int, list[RowRecord]] = {}
        for table_index, table in enumerate(document.tables):
            table_records = []
            for row_index, row in enumerate(table.rows):
                cells = tuple(unique_cells(row))
                writable = tuple(
                    (cell_index, cell, locked_prefix)
                    for cell_index, cell, locked_prefix in _writable_value_cells(
                        table_index, row_index, row
                    )
                )
                record = RowRecord(table_index, row_index, row, cells, writable)
                records.append(record)
                table_records.append(record)
            by_table[table_index] = table_records
        self._records = tuple(records)
        self._by_table = {key: tuple(value) for key, value in by_table.items()}
        self.table_signatures = tuple(_xml_signature(table._tbl) for table in document.tables)
        self.document_signature = hashlib.sha256(
            "|".join(self.table_signatures).encode("ascii")
        ).hexdigest()

    def rows(self, table_index: int | None = None, *, include_header: bool = True):
        """Return indexed rows; callers must treat row/cell references as read-only."""
        records = self._records if table_index is None else self._by_table.get(table_index, ())
        if include_header:
            return records
        return tuple(record for record in records if record.row_index > 0)

    def labels(self) -> tuple[str, ...]:
        return tuple(
            record.cells[0].text.strip()
            for record in self._records
            if record.cells
        )

    def writable_boundaries(self) -> tuple[tuple[int, int, tuple[int, ...]], ...]:
        return tuple(
            (record.table_index, record.row_index,
             tuple(cell_index for cell_index, _cell, _prefix in record.writable_cells))
            for record in self._records
            if record.writable_cells
        )

    def value_typography(self, language: str = "cn") -> list[dict]:
        problems: list[dict] = []
        for record in self._records:
            for cell_index, cell, locked_prefix in record.writable_cells:
                for run, paragraph, value_text in _iter_value_text_runs(cell, locked_prefix):
                    if _value_run_is_bold(run, paragraph):
                        problems.append({
                            "type": "bold_writable_value",
                            "language": language,
                            "table": record.table_index,
                            "row": record.row_index,
                            "cell": cell_index,
                            "text": value_text,
                        })
        return problems

    def whitespace_report(self) -> dict:
        issues: list[dict] = []
        intentional_breaks = 0
        intentional_tabs = 0
        for record in self._records:
            row = record.row
            tr_pr = row._tr.trPr
            if tr_pr is not None:
                for height in tr_pr.findall(qn("w:trHeight")):
                    rule = height.get(qn("w:hRule"))
                    value = height.get(qn("w:val"))
                    if rule == "exact":
                        issues.append({
                            "type": "exact_row_height_constraint",
                            "table": record.table_index,
                            "row": record.row_index,
                            "rule": rule,
                            "val": value,
                        })
            for cell_index, cell in enumerate(record.cells):
                paragraphs = list(cell.paragraphs)
                nonempty = [paragraph for paragraph in paragraphs if paragraph.text.strip()]
                if nonempty and len(paragraphs) > 1:
                    for paragraph_index, paragraph in enumerate(paragraphs):
                        if not paragraph.text.strip():
                            issues.append({
                                "type": "empty_paragraph_in_populated_cell",
                                "table": record.table_index,
                                "row": record.row_index,
                                "cell": cell_index,
                                "paragraph": paragraph_index,
                            })
                for paragraph_index, paragraph in enumerate(paragraphs):
                    for line_number, line in enumerate(paragraph.text.splitlines(), 1):
                        if line.strip() in {"/", "锛?"}:
                            issues.append({
                                "type": "slash_only_line",
                                "table": record.table_index,
                                "row": record.row_index,
                                "cell": cell_index,
                                "paragraph": paragraph_index,
                                "line": line_number,
                            })
                    locked = any(run.bold and run.text.strip() for run in paragraph.runs)
                    for run_index, run in enumerate(paragraph.runs):
                        text = run.text
                        if not locked and not run.bold and text == "" and len(paragraph.runs) > 1:
                            issues.append({
                                "type": "empty_nonbold_run",
                                "table": record.table_index,
                                "row": record.row_index,
                                "cell": cell_index,
                                "paragraph": paragraph_index,
                                "run": run_index,
                            })
                        if not locked and "\t" in text:
                            issues.append({
                                "type": "tab_in_text",
                                "table": record.table_index,
                                "row": record.row_index,
                                "cell": cell_index,
                                "paragraph": paragraph_index,
                                "run": run_index,
                            })
                        if not locked and re.search(r"[ \u3000]+$", text):
                            issues.append({
                                "type": "trailing_space",
                                "table": record.table_index,
                                "row": record.row_index,
                                "cell": cell_index,
                                "paragraph": paragraph_index,
                                "run": run_index,
                            })
                    paragraph_xml = paragraph._p.xml
                    if not locked and "<w:br" in paragraph_xml:
                        intentional_breaks += 1
                    if not locked and "<w:tab" in paragraph_xml:
                        intentional_tabs += 1
        return {
            "pass": not issues,
            "issue_count": len(issues),
            "intentional_break_count": intentional_breaks,
            "intentional_tab_count": intentional_tabs,
            "issues": issues,
        }

    def value_whitespace(self) -> list[dict]:
        """Return the OpenSpec value-cell whitespace observations."""
        problems: list[dict] = []
        for record in self._records:
            table_index = record.table_index
            row_index = record.row_index
            cells = record.cells
            if len(cells) == 1:
                value_cells = list(enumerate(cells))
            elif table_index == 2 and row_index >= 4:
                value_cells = list(enumerate(cells))
            elif table_index == 7 and row_index >= 14:
                value_cells = list(enumerate(cells))
            elif table_index == 10 and len(cells) >= 3:
                value_cells = [(len(cells) - 1, cells[-1])]
            else:
                value_cells = list(enumerate(cells[1:], start=1))
            for cell_index, cell in value_cells:
                for paragraph_index, paragraph in enumerate(cell.paragraphs):
                    text = paragraph.text
                    if "\n\n" in text or "\r\n\r\n" in text:
                        problems.append({
                            "type": "blank_line_in_value",
                            "table": table_index,
                            "row": row_index,
                            "cell": cell_index,
                            "paragraph": paragraph_index,
                        })
                    if re.search(r" {3,}|\u3000{2,}", text):
                        problems.append({
                            "type": "artificial_spacing_in_value",
                            "table": table_index,
                            "row": row_index,
                            "cell": cell_index,
                            "paragraph": paragraph_index,
                        })
        return problems

    def numbering_items(self):
        if not hasattr(self, "_numbering_items"):
            self._numbering_items = tuple(collect_numbered_items(self.document))
        return self._numbering_items

    def summary(self) -> dict:
        return {
            "table_count": len(self.document.tables),
            "row_count": len(self._records),
            "writable_boundary_count": len(self.writable_boundaries()),
            "document_signature": self.document_signature,
            "table_signatures": list(self.table_signatures),
        }


__all__ = ["AuditContext", "RowRecord"]
