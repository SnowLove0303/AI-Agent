#!/usr/bin/env python3
import argparse, json
from docx import Document
from docx.oxml.ns import qn

def unique_cells(row):
    out=[]; seen=set()
    for c in row.cells:
        key=id(c._tc)
        if key not in seen:
            seen.add(key); out.append(c)
    return out

def main():
    ap=argparse.ArgumentParser()
    ap.add_argument('docx')
    ap.add_argument('--out')
    a=ap.parse_args()
    d=Document(a.docx); result=[]
    for ti,t in enumerate(d.tables):
        td={'table':ti,'rows':[]}
        for ri,row in enumerate(t.rows):
            trPr=row._tr.trPr
            heights=[]
            if trPr is not None:
                for h in trPr.findall(qn('w:trHeight')):
                    heights.append({'val':h.get(qn('w:val')),'rule':h.get(qn('w:hRule'))})
            rd={'row':ri,'heights':heights,'cells':[]}
            for ci,c in enumerate(unique_cells(row)):
                cd={'cell':ci,'paragraphs':[]}
                for pi,p in enumerate(c.paragraphs):
                    cd['paragraphs'].append({'p':pi,'text':p.text,'runs':[{'text':r.text,'bold':bool(r.bold)} for r in p.runs]})
                rd['cells'].append(cd)
            td['rows'].append(rd)
        result.append(td)
    text=json.dumps(result,ensure_ascii=False,indent=2)
    if a.out: open(a.out,'w',encoding='utf-8').write(text)
    else: print(text)
if __name__=='__main__': main()
