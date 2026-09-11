#!/usr/bin/env python3
"""Source-pictogram extraction and insertion helpers.

The source DOCX remains the authority for a supplied pictogram.  This helper
copies the first embedded raster image into the cloned template's existing
Section 2 pictogram cell without rebuilding tables or changing cell geometry.
"""
from __future__ import annotations

import io
import zipfile
from copy import deepcopy
from pathlib import Path, PurePosixPath
from xml.etree import ElementTree

from docx.shared import Inches


DEFAULT_PICTOGRAM_WIDTH_INCHES = 0.9
_NS = {
    "a": "http://schemas.openxmlformats.org/drawingml/2006/main",
    "r": "http://schemas.openxmlformats.org/officeDocument/2006/relationships",
    "pr": "http://schemas.openxmlformats.org/package/2006/relationships",
    "wp": "http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing",
}


def extract_first_embedded_image(source_docx: str | Path) -> tuple[str, bytes]:
    source = Path(source_docx)
    with zipfile.ZipFile(source) as archive:
        names = sorted(name for name in archive.namelist() if name.startswith("word/media/") and not name.endswith("/"))
        if not names:
            raise ValueError(f"source DOCX contains no embedded image: {source}")
        name = names[0]
        return Path(name).name, archive.read(name)


def _source_image_width_inches(source_docx: str | Path, image_name: str) -> float | None:
    """Read the source drawing's physical width, when it is declared.

    Image bytes alone do not reliably preserve the Word display size.  The
    source drawing extent is the authoritative size and prevents a small
    single pictogram from being enlarged to the old 3.25-inch composite-image
    default.  A missing or malformed extent is handled by the compact fallback
    width used for ordinary single-pictogram cells.
    """
    source = Path(source_docx)
    try:
        with zipfile.ZipFile(source) as archive:
            for part in sorted(name for name in archive.namelist()
                               if name.startswith("word/") and name.endswith(".xml")
                               and "/_rels/" not in name):
                part_posix = PurePosixPath(part)
                rels_name = str(part_posix.parent / "_rels" / (part_posix.name + ".rels"))
                if rels_name not in archive.namelist():
                    continue
                root = ElementTree.fromstring(archive.read(part))
                rel_root = ElementTree.fromstring(archive.read(rels_name))
                targets = {
                    rel.get("Id"): Path(rel.get("Target", "")).name
                    for rel in rel_root.findall("pr:Relationship", _NS)
                }
                for drawing in root.iter():
                    if drawing.tag not in {
                        f"{{{_NS['wp']}}}inline",
                        f"{{{_NS['wp']}}}anchor",
                    }:
                        continue
                    blip = drawing.find(".//a:blip", _NS)
                    extent = drawing.find("wp:extent", _NS)
                    if blip is None or extent is None:
                        continue
                    target = targets.get(blip.get(f"{{{_NS['r']}}}embed"))
                    if target != image_name:
                        continue
                    cx = int(extent.get("cx", "0"))
                    if cx > 0:
                        return cx / 914400.0
    except (OSError, ValueError, ElementTree.ParseError, KeyError):
        return None
    return None


def insert_source_pictogram(document, source_docx: str | Path,
                            width_inches: float | None = None) -> dict:
    """Insert the source's first embedded image into the cloned S2 pictogram row."""
    name, payload = extract_first_embedded_image(source_docx)
    source_width = _source_image_width_inches(source_docx, name)
    width_inches = width_inches or source_width or DEFAULT_PICTOGRAM_WIDTH_INCHES
    if width_inches <= 0:
        raise ValueError("pictogram width must be positive")
    row = document.tables[1].rows[4]
    cell = row.cells[-1]
    if not cell.paragraphs:
        paragraph = cell.add_paragraph()
    else:
        paragraph = cell.paragraphs[0]
    saved_rpr = None
    for existing in paragraph.runs:
        if existing._r.rPr is not None:
            saved_rpr = deepcopy(existing._r.rPr)
            break
    for child in list(paragraph._p):
        if child.tag != "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}pPr":
            paragraph._p.remove(child)
    run = paragraph.add_run()
    if saved_rpr is not None:
        run._r.append(saved_rpr)
    run.add_picture(io.BytesIO(payload), width=Inches(width_inches))
    return {
        "source_image_name": name,
        "source_image_bytes": len(payload),
        "source_width_inches": source_width,
        "target_width_inches": width_inches,
        "target_table": 1,
        "target_row": 4,
    }


__all__ = [
    "DEFAULT_PICTOGRAM_WIDTH_INCHES",
    "extract_first_embedded_image",
    "insert_source_pictogram",
]
