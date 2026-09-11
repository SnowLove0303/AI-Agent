"""Stable template/runtime primitives for the parameterized MSDS pipeline."""
from __future__ import annotations

import copy
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn

from template_mutation_whitelist import unique_cells as whitelist_unique_cells
from template_mutation_whitelist import normalize_value_text
from template_mutation_whitelist import write_row_values


CN_GUANZHI = "广州冠志新材料科技有限公司"
CN_GUANZHI_ADDR = "广州市萝岗区科学城掬泉路3号广州国际企业孵化器A区1106室"
CN_GUOCAI = "英德市国彩精细化工有限公司"
CN_GUOCAI_ADDR = "广东省英德市白沙镇太平村更古坑凯迪工业园区"
EN_GUANZHI = "Guangzhou Guanzhi New Materials Technology Co., Ltd."
EN_GUANZHI_ADDR = "Room 1106, Area A, Guangzhou International Enterprise Incubator, No. 3 Juquan Road, Science City, Luogang District, Guangzhou"
EN_GUOCAI = "Yingde Guocai Fine Chemical Co., Ltd."
EN_GUOCAI_ADDR = "Kaidi Industrial Park, Genggukeng, Taiping Village, Baisha Town, Yingde, Guangdong, China"


def unique_cells(row):
    return whitelist_unique_cells(row)


def template_geometry(language: str) -> dict:
    rows = ([10, 16, 6, 6, 5, 4, 3, 16, 24, 6, 18, 6, 3, 5, 9, 2]
            if language == "zh" else
            [9, 16, 6, 6, 5, 4, 3, 16, 24, 6, 18, 6, 3, 5, 9, 2])
    if language not in {"zh", "en"}:
        raise ValueError(f"unsupported language: {language}")
    return {"table_count": 16, "rows": rows}


def validate_template_capacity(doc: Document, language: str) -> None:
    expected = template_geometry(language)
    actual = {"table_count": len(doc.tables),
              "rows": [len(table.rows) for table in doc.tables]}
    if actual != expected:
        raise RuntimeError(f"{language} template geometry mismatch: expected {expected}, found {actual}")


def project_rows_to_template(values, language: str, section: int, table: object) -> list:
    rows = list(values)
    if language == "en" and section == 1 and len(table.rows) == 9 and len(rows) == 9:
        return [rows[0], *rows[2:]]
    return rows


def _set_paragraph_text(paragraph, text: str) -> None:
    text = normalize_value_text(text)
    old_rpr = None
    for run in paragraph.runs:
        if run._r.rPr is not None:
            old_rpr = copy.deepcopy(run._r.rPr)
            break
    if old_rpr is None and paragraph._p.pPr is not None:
        template_rpr = paragraph._p.pPr.find(qn("w:rPr"))
        if template_rpr is not None:
            old_rpr = copy.deepcopy(template_rpr)
    for child in list(paragraph._p):
        if child.tag != qn("w:pPr"):
            paragraph._p.remove(child)
    run = paragraph._p.makeelement(qn("w:r"), {})
    if old_rpr is not None:
        run.append(old_rpr)
    for index, line in enumerate(str(text).split("\n")):
        if index:
            run.append(paragraph._p.makeelement(qn("w:br"), {}))
        text_node = paragraph._p.makeelement(qn("w:t"), {})
        text_node.text = line
        if line and (line[0].isspace() or line[-1].isspace()):
            text_node.set(qn("xml:space"), "preserve")
        run.append(text_node)
    paragraph._p.append(run)


def set_paragraph_text(paragraph, text: str) -> None:
    _set_paragraph_text(paragraph, text)


def set_cell_text(cell, text: str) -> None:
    if not cell.paragraphs:
        cell.add_paragraph()
    _set_paragraph_text(cell.paragraphs[0], text)
    for paragraph in cell.paragraphs[1:]:
        paragraph._element.getparent().remove(paragraph._element)


def sanitize_template_artifacts(doc: Document) -> None:
    """Deprecated no-op: cloned template labels are immutable."""
    return None


def set_row(row, values, *, table_index=None, row_index=None, registry=None,
            inserted_data_row=False):
    return write_row_values(row, values, table_index=table_index,
                            row_index=row_index, registry=registry,
                            inserted_data_row=inserted_data_row)


def ensure_s3_component_rows(doc: Document, component_count: int) -> None:
    table = doc.tables[2]
    required_rows = 4 + component_count
    while len(table.rows) < required_rows:
        table._tbl.append(copy.deepcopy(table.rows[-1]._tr))


def ensure_source_data_rows(doc: Document, section: int, source_row_count: int) -> None:
    """Clone only the last styled data row for approved expandable sections."""
    if section not in {9, 15}:
        raise ValueError(f"source-row insertion is not allowed for S{section}")
    table = doc.tables[section - 1]
    required_rows = 1 + source_row_count
    while len(table.rows) < required_rows:
        table._tbl.append(copy.deepcopy(table.rows[-1]._tr))


__all__ = [
    "CN_GUANZHI", "CN_GUANZHI_ADDR", "CN_GUOCAI", "CN_GUOCAI_ADDR",
    "EN_GUANZHI", "EN_GUANZHI_ADDR", "EN_GUOCAI", "EN_GUOCAI_ADDR",
    "ensure_s3_component_rows", "ensure_source_data_rows", "project_rows_to_template", "set_cell_text",
    "set_paragraph_text", "set_row", "template_geometry", "unique_cells",
    "sanitize_template_artifacts", "validate_template_capacity",
]
