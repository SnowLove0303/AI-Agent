#!/usr/bin/env python3
"""Create a deterministic DOCX structural/geometry snapshot for template pinning."""
import argparse
import hashlib
import json
from pathlib import Path

from docx import Document
from lxml import etree

W_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
NS = {"w": W_NS}


def sha(value):
    if value is None:
        return None
    if isinstance(value, bytes):
        data = value
    else:
        data = etree.tostring(value, method="c14n")
    return hashlib.sha256(data).hexdigest()


def attr(el, name):
    return None if el is None else el.get(f"{{{W_NS}}}{name}")


def unique_cells(row):
    cells = []
    seen = set()
    for cell in row.cells:
        key = id(cell._tc)
        if key not in seen:
            seen.add(key)
            cells.append(cell)
    return cells


def cell_ref(cell):
    return cell._tc.getroottree().getpath(cell._tc)


def paragraph_snapshot(paragraph):
    return {
        "text": paragraph.text,
        "pPr_hash": sha(paragraph._p.pPr),
        "runs": [
            {
                "text": run.text,
                "rPr_hash": sha(run._r.rPr),
                "bold": run.bold,
                "italic": run.italic,
                "underline": str(run.underline) if run.underline is not None else None,
                "font": {
                    "name": run.font.name,
                    "size_pt": float(run.font.size.pt) if run.font.size else None,
                    "color": run.font.color.rgb if run.font.color and run.font.color.rgb else None,
                },
            }
            for run in paragraph.runs
        ],
    }


def table_snapshot(table, index):
    grid = table._tbl.tblGrid
    grid_widths = [attr(col, "w") for col in grid.findall("w:gridCol", NS)]
    rows = []
    for row_index, row in enumerate(table.rows):
        cells = []
        for cell_index, cell in enumerate(unique_cells(row)):
            tc_pr = cell._tc.tcPr
            tc_w = tc_pr.find("w:tcW", NS) if tc_pr is not None else None
            v_merge = tc_pr.find("w:vMerge", NS) if tc_pr is not None else None
            grid_span = tc_pr.find("w:gridSpan", NS) if tc_pr is not None else None
            cells.append(
                {
                    "cell": cell_index,
                    "xml_path": cell_ref(cell),
                    "tcPr_hash": sha(tc_pr),
                    "tcW": attr(tc_w, "w"),
                    "tcW_type": attr(tc_w, "type"),
                    "gridSpan": attr(grid_span, "val"),
                    "vMerge": attr(v_merge, "val") if v_merge is not None else None,
                    "text": cell.text,
                    "paragraphs": [paragraph_snapshot(p) for p in cell.paragraphs],
                }
            )
        rows.append(
            {
                "row": row_index,
                "trPr_hash": sha(row._tr.trPr),
                "cell_count": len(cells),
                "cells": cells,
            }
        )
    return {
        "table": index,
        "row_count": len(table.rows),
        "column_count": len(grid_widths),
        "unique_cell_counts": [row["cell_count"] for row in rows],
        "tblPr_hash": sha(table._tbl.tblPr),
        "tblGrid_hash": sha(grid),
        "grid_widths_dxa": grid_widths,
        "rows": rows,
    }


def header_footer_snapshot(section):
    result = {}
    for role, container in (("header", section.header), ("footer", section.footer)):
        result[role] = {
            "paragraphs": [paragraph_snapshot(p) for p in container.paragraphs],
            "tables": [table_snapshot(t, i) for i, t in enumerate(container.tables)],
            "part_xml_hash": sha(container._element),
        }
    return result


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("docx", type=Path)
    parser.add_argument("json_out", type=Path)
    args = parser.parse_args()
    document = Document(args.docx)
    data = {
        "source_filename": args.docx.name,
        "source_sha256": hashlib.sha256(args.docx.read_bytes()).hexdigest(),
        "tables": [table_snapshot(table, i) for i, table in enumerate(document.tables)],
        "sections": [header_footer_snapshot(section) for section in document.sections],
    }
    args.json_out.parent.mkdir(parents=True, exist_ok=True)
    args.json_out.write_text(json.dumps(data, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps({
        "source_sha256": data["source_sha256"],
        "table_count": len(data["tables"]),
        "row_counts": [t["row_count"] for t in data["tables"]],
        "column_counts": [t["column_count"] for t in data["tables"]],
        "section_count": len(data["sections"]),
        "snapshot": str(args.json_out),
    }, ensure_ascii=False))


if __name__ == "__main__":
    main()
