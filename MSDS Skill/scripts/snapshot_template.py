#!/usr/bin/env python3
import argparse, json, hashlib
from docx import Document
from lxml import etree
from docx.oxml.ns import qn

def unique_cells(row):
    out=[]; seen=set()
    for c in row.cells:
        k=id(c._tc)
        if k not in seen: seen.add(k); out.append(c)
    return out

def hx(el):
    return None if el is None else hashlib.sha256(etree.tostring(el, method='c14n')).hexdigest()

def main():
    ap=argparse.ArgumentParser(); ap.add_argument('docx'); ap.add_argument('json_out'); a=ap.parse_args()
    d=Document(a.docx); data={'tables':[]}
    for ti,t in enumerate(d.tables):
        td={'table':ti,'tblPr':hx(t._tbl.tblPr),'grid':hx(t._tbl.tblGrid),'rows':[]}
        for ri,row in enumerate(t.rows):
            rd={'row':ri,'trPr':hx(row._tr.trPr),'cells':[]}
            for ci,c in enumerate(unique_cells(row)):
                cd={'cell':ci,'tcPr':hx(c._tc.tcPr),'paragraphs':[]}
                for pi,p in enumerate(c.paragraphs):
                    cd['paragraphs'].append({'p':pi,'text':p.text,'pPr':hx(p._p.pPr),'runs':[{'text':r.text,'bold':bool(r.bold),'rPr':hx(r._r.rPr)} for r in p.runs]})
                rd['cells'].append(cd)
            td['rows'].append(rd)
        data['tables'].append(td)
    open(a.json_out,'w',encoding='utf-8').write(json.dumps(data,ensure_ascii=False,indent=2))
if __name__=='__main__': main()
