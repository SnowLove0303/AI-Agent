#!/usr/bin/env python3
"""Customer-facing Section 2 policy for the unified MSDS skill.

This module keeps Section 2 source-grounded and readable:

* GHS label elements are written as explicit, line-separated tips.  A reader
  must never be sent to another row with a phrase such as ``See 2.4-2.6``.
* A source pictogram image is treated as an asset, not as text or a missing
  value.  The caller inserts it into the cloned template pictogram cell.
* Whole Section 2 rows whose value is only missing data are removed before
  visible numbering is recalculated.  A value such as ``Eyes: No data`` is
  also a missing row; substantive negatives such as ``No irritation`` are not.

The module does not infer hazard classes, H/P codes, signal words or label
ingredients.  Those inputs must already be present in the verified semantic
model or in the source asset.
"""
from __future__ import annotations

import re
from copy import deepcopy

from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from docx.shared import Pt

from section2_hp_policy import is_missing_data_value, render_precautionary_groups
from template_mutation_whitelist import composite_value_text, is_s28_row, unique_cells

try:
    from ghs_code_resolver import format_precautionary_text, resolve_precautionary_statements
    from jev_dispatcher import resolve_signal_word, route_s3_to_s2
except ImportError:
    try:
        from .ghs_code_resolver import format_precautionary_text, resolve_precautionary_statements
        from .jev_dispatcher import resolve_signal_word, route_s3_to_s2
    except ImportError:
        format_precautionary_text = None
        resolve_precautionary_statements = None
        resolve_signal_word = None
        route_s3_to_s2 = None


def set_cell_value_unified(
    cell,
    text_or_lines: str | list[str] | tuple[str, ...],
    lang: str = "zh",
    bold: bool = False,
    size_pt: float = 12.0,
    indent_sub_items: bool = False,
    sub_item_indent_pt: float = 18.0,
):
    """Set cell value with guaranteed 12.0 pt typography and consistent font family.

    Splits multi-line content into distinct paragraphs, cleans extraneous empty paragraphs,
    and injects complete w:rPr on every run.
    """
    if isinstance(text_or_lines, str):
        lines = [line.strip() for line in text_or_lines.split("\n") if line.strip()]
    elif isinstance(text_or_lines, (list, tuple)):
        lines = [str(line).strip() for line in text_or_lines if str(line).strip()]
    else:
        s = str(text_or_lines).strip() if text_or_lines is not None else ""
        lines = [s] if s else []

    if not lines:
        lines = [""]

    half_pts = str(int(round(size_pt * 2)))

    # Ensure exact number of paragraphs
    while len(cell.paragraphs) < len(lines):
        cell.add_paragraph()
    while len(cell.paragraphs) > len(lines):
        p_elem = cell.paragraphs[-1]._p
        p_elem.getparent().remove(p_elem)

    heading_pattern = re.compile(
        r"^(?:预防措施|事故响应|安全储存|安全存储|废弃处置|Prevention|Response|Storage|Disposal)[：:]?$",
        re.IGNORECASE,
    )

    for p, line in zip(cell.paragraphs, lines):
        p.text = line
        if indent_sub_items:
            if heading_pattern.match(line):
                p.paragraph_format.left_indent = Pt(0)
            else:
                p.paragraph_format.left_indent = Pt(sub_item_indent_pt)
        run = p.runs[0]
        rPr = run._r.get_or_add_rPr()

        # 1. Fonts
        rFonts = rPr.find(qn("w:rFonts"))
        if rFonts is None:
            rFonts = OxmlElement("w:rFonts")
            rPr.append(rFonts)
        rFonts.set(qn("w:ascii"), "Arial")
        rFonts.set(qn("w:hAnsi"), "Arial")
        if lang == "zh":
            rFonts.set(qn("w:eastAsia"), "宋体")
            rFonts.set(qn("w:hint"), "eastAsia")

        # 2. Size
        sz = rPr.find(qn("w:sz"))
        if sz is None:
            sz = OxmlElement("w:sz")
            rPr.append(sz)
        sz.set(qn("w:val"), half_pts)

        szCs = rPr.find(qn("w:szCs"))
        if szCs is None:
            szCs = OxmlElement("w:szCs")
            rPr.append(szCs)
        szCs.set(qn("w:val"), half_pts)

        # 3. Bold
        run.font.bold = bold

def format_label_elements(
    language: str,
    hazardous_ingredients: list[str] | tuple[str, ...],
    extra_note: str | None = None,
) -> str:
    """Return the explicit, line-separated GHS label-ingredient tip.

    With no verified ingredients the value stays empty so the existing
    missing-row suppression removes the whole label-elements row instead of
    leaving a bare heading in customer-facing output.
    """
    ingredients = [str(x).strip() for x in hazardous_ingredients if str(x).strip()]
    if not ingredients and not extra_note:
        if language not in ("en", "zh"):
            raise ValueError("language must be zh or en")
        return ""
    if language == "en":
        heading = "Hazardous ingredients required to be listed on the label:"
    elif language == "zh":
        heading = "必须列在标签上的有害成分："
    else:
        raise ValueError("language must be zh or en")
    parts = []
    if ingredients:
        parts.append(heading)
        parts.extend(ingredients)
    if extra_note and str(extra_note).strip():
        parts.append(str(extra_note).strip())
    return "\n".join(parts)


def row_has_visual_content(row) -> bool:
    """Return true when a row contains an embedded drawing/picture."""
    for cell in row.cells:
        if cell._tc.xpath(".//w:drawing") or cell._tc.xpath(".//w:pict"):
            return True
    return False


def _value_text(row) -> str:
    if is_s28_row(1, row):
        return composite_value_text(row).strip()
    cells = []
    seen = set()
    for cell in row.cells:
        key = hash(cell._tc)
        if key in seen:
            continue
        seen.add(key)
        # Keep empty value cells in the shape calculation.  Dropping them
        # would make a two-cell row look like a merged label-only row.
        cells.append(cell.text.strip())
    # Two-cell rows have a label cell followed by a value cell.  Exclude the
    # label even when the value cell is empty; this is what lets an empty
    # pictogram slot be removed.  A genuinely merged one-cell row contains
    # both label and value and therefore remains self-contained.
    return " ".join(cells[1:]) if len(cells) > 1 else " ".join(cells)


def is_missing_section2_value(value: str) -> bool:
    """Recognize only whole-item missing values, including ``Label: No data``."""
    if is_missing_data_value(value):
        return True
    stripped = re.sub(r"\s+", " ", (value or "").strip())
    if ":" in stripped or "：" in stripped:
        remainder = re.split(r"[:：]", stripped, maxsplit=1)[1].strip()
        return is_missing_data_value(remainder)
    return False


_LABEL_INGREDIENT_VALUE_RE = re.compile(
    r"(?:必须列在标签上的有害成分|标签上(?:列出的|要求列出的)?有害成分|"
    r"hazardous ingredients required to be listed on the label|"
    r"ingredients required to be listed on the label)",
    re.I,
)
_SIGNAL_WORD_RE = re.compile(
    r"^(?:危险|警告|无信号词|无|danger|warning|no signal word|none|not applicable)$",
    re.I,
)


def is_signal_word_value(value: str) -> bool:
    """Return whether a semantic S2 value is in the controlled vocabulary."""
    candidate = re.sub(r"[\s。；;.!！]+$", "", (value or "").strip())
    return bool(_SIGNAL_WORD_RE.fullmatch(candidate))


def looks_like_label_ingredient_value(value: str) -> bool:
    """Recognize the explicit label-ingredient explanation marker."""
    return bool(_LABEL_INGREDIENT_VALUE_RE.search((value or "").strip()))


def validate_s2_semantics(rows, language: str = "zh") -> list[str]:
    """Block the two high-risk S2 slot swaps before any template is cloned.

    The maintained template's source-semantic mapping is:
    source 2.2 label elements -> template 2.3 GHS Label Elements, while the
    template's 2.4 Signal Word is a separate slot.  A label-ingredient
    explanation must never be written into the signal-word slot, and a signal
    word alone is not a valid label-elements explanation.
    """
    errors: list[str] = []
    if not isinstance(rows, list) or len(rows) < 5:
        return errors

    def value_at(index: int) -> str:
        row = rows[index]
        if not isinstance(row, (list, tuple)) or len(row) < 2:
            return ""
        return str(row[1] or "").strip()

    label_value = value_at(2)
    signal_value = value_at(4)
    if is_signal_word_value(label_value):
        errors.append(
            f"{language} Section 2 label-elements slot contains only a signal word; "
            "write the source label-ingredient explanation instead"
        )
    if looks_like_label_ingredient_value(signal_value):
        errors.append(
            f"{language} Section 2 signal-word slot contains label-ingredient prose; "
            "move it to the label-elements slot"
        )
    if signal_value and not is_signal_word_value(signal_value):
        errors.append(
            f"{language} Section 2 signal-word slot must use the controlled signal-word vocabulary"
        )
    return errors


def is_explicit_other_hazards_row(row) -> bool:
    """Keep only a source-backed ``Other hazards`` conclusion visible.

    The template label itself is not evidence.  Older callers treated the
    presence of ``2.10 其他危害`` as an exception and consequently retained a
    blank row forever.  The exception now requires a non-empty value after the
    label, including the valid explicit conclusion ``无适用资料。``.
    """
    cells = []
    seen = set()
    for cell in row.cells:
        key = hash(cell._tc)
        if key not in seen:
            seen.add(key)
            cells.append(cell)
    text = " ".join(cell.text.strip() for cell in cells if cell.text.strip())
    if not re.search(
        r"(?:^|\s)(?:2\.\d+\s*)?(?:其他危险|其他危害|other hazards)\b",
        text,
        re.I,
    ):
        return False
    if len(cells) > 1:
        value = " ".join(cell.text.strip() for cell in cells[1:] if cell.text.strip())
        # Be defensive when a source value was written into the label cell by
        # a legacy intermediate model: recover only text after the label.
        if not value:
            match = re.search(
                r"(?:其他危险|其他危害|other hazards)\s*[:：]\s*(.+)$",
                cells[0].text.strip(), re.I,
            )
            value = match.group(1).strip() if match else ""
    else:
        match = re.search(
            r"(?:其他危险|其他危害|other hazards)\s*[:：]?\s*(.*)$",
            cells[0].text.strip(), re.I,
        )
        value = match.group(1).strip() if match else ""
    return bool(value)


def project_source_cn_headings(document, language: str = "zh") -> list[str]:
    """Reject the retired label-rewrite path.

    Labels are template-owned. The source heading is used only to select the
    existing semantic slot; it must never be copied into the locked label
    cell. Keeping this function as an explicit guard prevents legacy callers
    from silently reintroducing label mutations.
    """
    if language == "zh":
        from template_mutation_whitelist import MutationViolation
        raise MutationViolation(
            "locked Section 2 labels are immutable; write the value cell only"
        )
    return []


def project_source_cn_facts(s2: dict) -> tuple[list[list[str]], dict[int, int]]:
    """Project source S2 semantics into the fixed CN template slots.

    The source numbers are semantic, not positional: source 2.1/2.2/2.3
    correspond to template slots 2.2/2.3/2.10.  Blank template slots are
    intentionally retained here so the shared suppression pass can remove
    them and renumber only the surviving source-backed items.
    """
    def first(name: str) -> str:
        values = s2.get(name) or []
        if isinstance(values, str):
            return values.strip()
        return str(values[0]).strip() if values else ""

    def health_route(route: str) -> str:
        health = s2.get("health_hazards") or {}
        if isinstance(health, dict):
            values = health.get(route) or []
            if isinstance(values, str):
                return values.strip()
            return str(values[0]).strip() if values else ""
        return ""

    def route_label(route: str) -> str:
        # This semantic marker exists only in the reviewed facts layer.  The
        # writer ignores source labels and keeps the cloned template prefix;
        # the marker lets source mapping/trace audits identify repeated 2.8
        # rows after empty-row suppression without using physical position.
        return f"2.8 健康危害 [route={route}]"

    label_elements = first("label_elements")
    if not label_elements:
        label_elements = format_label_elements("zh", s2.get("label_ingredients") or [])

    precautionary_value = render_precautionary_groups(
        s2.get("precautionary_groups"), "zh"
    )
    if not precautionary_value:
        precautionary_value = "\n".join(
            str(x).strip() for x in (s2.get("p_statements") or []) if str(x).strip()
        )

    return ([
        ["2.1 紧急情况概述", ""],
        ["2.2 GHS危险性类别：", first("ghs_classes")],
        ["2.3 GHS标签要素：", label_elements],
        [" GHS象形图", ""],
        ["2.4 信号词：", str(s2.get("signal") or "").strip()],
        ["2.5 危险性说明：", "\n".join(str(x).strip() for x in (s2.get("h_statements") or []) if str(x).strip())],
        ["2.6 防范说明：", precautionary_value],
        ["2.7 物理和化学危险：", ""],
        [route_label("inhalation"), health_route("inhalation")],
        [route_label("ingestion"), health_route("ingestion")],
        [route_label("skin"), health_route("skin")],
        [route_label("eyes"), health_route("eyes")],
        [route_label("symptoms_signs"), health_route("symptoms_signs")],
        ["2.9 环境危害", ""],
        ["2.10 其他危害：", str(s2.get("other_hazards") or "").strip()],
    ], {2: 1, 3: 2, 10: 3})


def _replace_prefix(paragraph, section: int, item: int, set_paragraph_text) -> None:
    current = paragraph.text
    updated = re.sub(rf"^(\s*){section}\.\d+(\b)", rf"\g<1>{section}.{item}\g<2>", current, count=1)
    if updated != current:
        # This is the sole approved label-cell content mutation.  Preserve
        # the existing run tree even when Word split ``2.2`` across runs.
        from template_mutation_whitelist import _replace_leading_pattern_in_runs
        _replace_leading_pattern_in_runs(
            paragraph,
            rf"^(\s*){section}\.\d+(\b)",
            rf"\g<1>{section}.{item}\g<2>",
        )


_S28_ROUTE_PREFIX_TEXTS = (
    "吸入：", "食入：", "皮肤：", "眼睛：", "症状和体征：",
    "Inhalation:", "Ingestion:", "Skin:", "Eyes:",
    "Signs and symptoms:",
)


def _xml_cell_text(cell) -> str:
    return "".join(node.text or "" for node in cell.iter(qn("w:t"))).strip()


def _raw_s28_route_row(row) -> bool:
    """Detect a route row from raw XML even if its merged label is blank."""
    cells = row._tr.findall(qn("w:tc"))
    if len(cells) < 2:
        return False
    value_text = _xml_cell_text(cells[1])
    return any(value_text.startswith(prefix) for prefix in _S28_ROUTE_PREFIX_TEXTS)


def _promote_s28_label_cell(target_cell, source_cell) -> None:
    """Repair a vertical S2.8 merge after its restart row was removed.

    The maintained template stores the repeated ``健康危害`` label in a
    vertically merged restart cell; continuation cells are physically blank.
    Removing the restart row without promoting the next continuation cell
    makes python-docx expose the previous row's label for every later route.
    Promote the next raw XML cell and copy only the deleted label's paragraph
    content/format.  The target cell's width, borders and other geometry stay
    intact; changing vMerge from continuation to restart is the smallest
    authorized merge repair.
    """
    target_pr = target_cell.get_or_add_tcPr()
    for child in list(target_cell):
        if child is not target_pr:
            target_cell.remove(child)
    for child in list(source_cell):
        if child.tag != qn("w:tcPr"):
            target_cell.append(deepcopy(child))
    vmerge = target_pr.find(qn("w:vMerge"))
    if vmerge is None:
        vmerge = target_pr.makeelement(qn("w:vMerge"), {})
        target_pr.append(vmerge)
    vmerge.set(qn("w:val"), "restart")


def suppress_missing_section2_rows_and_renumber(document, set_paragraph_text,
                                                number_map: dict[int, int] | None = None):
    """Remove missing Section 2 rows and renumber unique visible items.

    Repeated child rows such as multiple ``2.8 Health hazards`` rows retain a
    shared number.  The unnumbered GHS pictogram row is retained when it has an
    embedded image, even though its text value is empty.
    """
    table = document.tables[1]
    removed_labels = []
    removed_s28_label_cell = None
    # Remove from the bottom up.  Section 2.8 contains vertically merged
    # continuation cells; forward XML removal invalidates later row proxies
    # and can cause an empty route to survive or inherit the prior label.
    for row in reversed(list(table.rows)[1:]):
        cells = []
        seen = set()
        for cell in row.cells:
            key = hash(cell._tc)
            if key not in seen:
                seen.add(key)
                cells.append(cell)
        label = cells[0].text.strip() if cells else ""
        if (
            not is_explicit_other_hazards_row(row)
            and not row_has_visual_content(row)
            and (not _value_text(row) or is_missing_section2_value(_value_text(row)))
        ):
            if removed_s28_label_cell is None and is_s28_row(1, row) and cells:
                # Preserve the deleted merged restart cell as the style/text
                # source for the first surviving continuation row.
                removed_s28_label_cell = deepcopy(cells[0]._tc)
            # Keep an unnumbered non-data row only when it is the pictogram slot
            # and the picture is present.  Other empty rows are not customer-facing.
            removed_labels.append(label)
            table._tbl.remove(row._tr)
        else:
            pass

    # Re-read rows after XML removal.  Retaining row proxy objects across
    # deletion can make python-docx resolve a later repeated row to the old
    # physical position, which is unsafe for repeated 2.8 child rows.
    visible_rows = list(table.rows)[1:]
    if removed_s28_label_cell is not None:
        for row in visible_rows:
            if not _raw_s28_route_row(row):
                continue
            raw_cells = row._tr.findall(qn("w:tc"))
            if raw_cells and not _xml_cell_text(raw_cells[0]):
                _promote_s28_label_cell(raw_cells[0], removed_s28_label_cell)
            break

    # Snapshot the original item number before changing any label text.  A
    # vertically merged label cell is returned by python-docx for every child
    # row, so reading its text after the first rewrite would make the next
    # child appear to have a different original number.
    row_info = []
    for row in visible_rows:
        cells = []
        seen = set()
        for cell in row.cells:
            key = hash(cell._tc)
            if key not in seen:
                seen.add(key)
                cells.append(cell)
        if not cells or not cells[0].paragraphs:
            continue
        paragraph = cells[0].paragraphs[0]
        match = re.match(r"^\s*2\.(\d+)\b", paragraph.text)
        if match:
            row_info.append((row, cells, hash(cells[0]._tc), int(match.group(1))))

    old_to_new = {}
    next_item = 1
    labels = []
    rewritten_cells = set()
    for row, cells, label_cell_key, old in row_info:
        if old not in old_to_new:
            old_to_new[old] = number_map.get(old, next_item) if number_map else next_item
            next_item = max(next_item, old_to_new[old] + 1)
        if label_cell_key not in rewritten_cells:
            _replace_prefix(cells[0].paragraphs[0], 2, old_to_new[old], set_paragraph_text)
            rewritten_cells.add(label_cell_key)
        labels.append(cells[0].paragraphs[0].text.strip())

    return {
        "removed_count": len(removed_labels),
        "removed_labels": removed_labels,
        "visible_count": len(labels),
        "visible_labels": labels,
        "number_map": {f"2.{old}": f"2.{new}" for old, new in old_to_new.items()},
    }


__all__ = [
    "format_label_elements",
    "project_source_cn_facts",
    "row_has_visual_content",
    "is_missing_section2_value",
    "is_explicit_other_hazards_row",
    "is_signal_word_value",
    "looks_like_label_ingredient_value",
    "validate_s2_semantics",
    "suppress_missing_section2_rows_and_renumber",
]
